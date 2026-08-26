"""
Postoperative CT Registration Panel for CT-MRI Coregistration
Author: Sunil Mathew (with AI assistance)
Date: December 2025

This panel provides an interactive interface for CT-MRI coregistration with:
- Before alignment visualization
- Progress tracking during registration
- After alignment visualization with comparison
"""

import os
import sys
import warnings
import csv
import shutil
import time
from datetime import datetime
from pathlib import Path
import re
import numpy as np
import nibabel as nib
import scipy.ndimage as ndi
from qtpy import QtCore, QtWidgets, QtGui
from qtpy.QtWidgets import (QDialog, QVBoxLayout, QHBoxLayout, QPushButton, 
                             QLabel, QProgressBar, QTextEdit, QTabWidget, QWidget, QSlider,
                             QCheckBox, QDoubleSpinBox, QGroupBox, QGridLayout, QListWidget,
                             QInputDialog, QListWidgetItem, QStyle, QMessageBox)
from qtpy.QtCore import Qt, QThread, Signal
import pyqtgraph as pg
import matplotlib
matplotlib.use('Qt5Agg')
from matplotlib.backends.backend_qt5agg import FigureCanvasQTAgg
from matplotlib.figure import Figure
from matplotlib.widgets import RectangleSelector
import mne
from utils.electrode_map import find_electrode_map, read_electrode_map
from utils.planning_worksheet import (
    count_embedded_images,
    find_planning_worksheet,
    read_planning_worksheet,
)
from utils.planning_schema import assign_schema_prior
from utils.shaft_utils import (
    cluster_shaft_contacts,
    fit_shaft_axis_pca,
    project_point_to_line,
    split_labels_by_group,
    split_line_clusters_by_gap,
)
from viewer3d.electrode import Electrode

DEFAULT_GROUP_MIN = 5
DEFAULT_BLOB_MIN = 3
DEFAULT_MASK_DILATE = 3
DEFAULT_AXIAL_LINE_WIGGLE = 5.0
DEFAULT_AXIAL_MAX_GAP = 25.0
DEFAULT_MARKER_SIZE = 6
DETECTED_SYNC_DELAY_MS = 50
DEFAULT_SHAFT_LABEL_OFFSET = 14.0
ELECTRODE_MAT_CLOSE_DISTANCE = 3.0
ELECTRODE_MAT_OFF_DISTANCE = 10.0
DEFAULT_CONTROLS_HBOX_STRETCH = {
    "controls": 5,
    "shafts": 3,
    "electrodes": 2
}
PROFILE_MERGE = bool(int(os.getenv("SEEG_LOCALIZER_PROFILE_MERGE", "0")))


class ListItemWidget(QWidget):
    """Custom widget for list items with embedded action icons."""
    
    edit_clicked = Signal(object)  # Emits the item
    hide_clicked = Signal(object)  # Emits the item  
    delete_clicked = Signal(object)  # Emits the item
    
    def __init__(self, text, color=None, parent=None):
        super().__init__(parent)
        layout = QHBoxLayout()
        layout.setContentsMargins(2, 2, 2, 2)
        layout.setSpacing(4)
        
        # Text label
        self.text_label = QLabel(text)
        if color:
            self.text_label.setStyleSheet(f"color: {color}; font-weight: bold;")
        layout.addWidget(self.text_label, 1)  # Stretch factor 1
        
        # Icon buttons
        self.edit_btn = QPushButton("✏️")
        self.edit_btn.setMaximumSize(20, 20)
        self.edit_btn.setStyleSheet("QPushButton { font-size: 12px; padding: 0px; border: none; background: transparent; }")
        self.edit_btn.setToolTip("Rename")
        self.edit_btn.clicked.connect(lambda: self.edit_clicked.emit(self))
        layout.addWidget(self.edit_btn)
        
        self.hide_btn = QPushButton("👁")
        self.hide_btn.setMaximumSize(20, 20)
        self.hide_btn.setCheckable(True)
        self.hide_btn.setStyleSheet("QPushButton { font-size: 12px; padding: 0px; border: none; background: transparent; }")
        self.hide_btn.setToolTip("Hide/Show")
        self.hide_btn.clicked.connect(lambda: self.hide_clicked.emit(self))
        layout.addWidget(self.hide_btn)
        
        self.delete_btn = QPushButton("🗑")
        self.delete_btn.setMaximumSize(20, 20)
        self.delete_btn.setStyleSheet("QPushButton { font-size: 12px; padding: 0px; border: none; background: transparent; }")
        self.delete_btn.setToolTip("Delete")
        self.delete_btn.clicked.connect(lambda: self.delete_clicked.emit(self))
        layout.addWidget(self.delete_btn)
        
        self.setLayout(layout)
    
    def sizeHint(self):
        """Provide a size hint for the widget."""
        from qtpy.QtCore import QSize
        return QSize(200, 24)
    
    def set_text(self, text):
        """Update the text label."""
        self.text_label.setText(text)
    
    def set_hidden_state(self, is_hidden):
        """Update the hide button state."""
        self.hide_btn.setChecked(is_hidden)


 


class RegistrationWorker(QThread):
    """Worker thread for performing registration without blocking the UI."""
    
    progress = Signal(int, str)  # progress percentage, status message
    finished = Signal(bool, str, object, str)  # success, message, transform, registered_ct_path
    
    def __init__(self, mri_path, ct_path, output_dir, method="mne_rigid"):
        super().__init__()
        self.mri_path = mri_path
        self.ct_path = ct_path
        self.output_dir = output_dir
        self.method = method
        
    def run(self):
        """Perform the registration."""
        try:
            self.progress.emit(10, "Loading images...")

            if self.method.startswith("ants"):
                try:
                    import ants
                except Exception as exc:
                    raise RuntimeError(
                        "ANTs registration requested but antspyx is not available."
                    ) from exc

                mri_img = ants.image_read(self.mri_path)
                ct_img = ants.image_read(self.ct_path)

                self.progress.emit(25, "Initializing ANTs registration...")
                method = "Rigid" if self.method == "ants_rigid" else "Affine"
                self.progress.emit(35, f"Computing ANTs {method.lower()} registration...")

                reg = ants.registration(
                    fixed=mri_img,
                    moving=ct_img,
                    type_of_transform=method,
                    verbose=True,
                )

                self.progress.emit(75, "Registration complete! Applying transformation to CT...")
                ct_aligned = reg.get("warpedmovout", None)
                if ct_aligned is None:
                    raise RuntimeError("ANTs registration did not return warped output.")

                self.progress.emit(85, "Saving registered CT...")
                os.makedirs(self.output_dir, exist_ok=True)
                suffix = "_ants" if self.method.startswith("ants") else ""
                ct_registered_path = os.path.join(
                    self.output_dir, f"ct_registered_to_mri{suffix}.nii.gz"
                )
                ants.image_write(ct_aligned, ct_registered_path)

                ants_transforms = []
                for path in reg.get("fwdtransforms", []):
                    if not path:
                        continue
                    dest = os.path.join(self.output_dir, os.path.basename(path))
                    try:
                        if os.path.abspath(path) != os.path.abspath(dest):
                            shutil.copy(path, dest)
                        ants_transforms.append(dest)
                    except Exception:
                        ants_transforms.append(path)

                trans = {
                    "trans": None,
                    "ants_fwdtransforms": ants_transforms,
                }

                self.progress.emit(100, "Registration complete!")
                self.finished.emit(
                    True,
                    f"ANTs registration successful! Saved to: {ct_registered_path}",
                    trans,
                    ct_registered_path,
                )
                return

            # Default: MNE rigid registration
            mri_img = nib.load(self.mri_path)
            ct_img = nib.load(self.ct_path)

            self.progress.emit(20, "Initializing registration algorithm...")

            # Use faster zooms for quicker registration (can be adjusted for accuracy)
            # Higher values = faster but less accurate. Default=None is most accurate but slowest
            zooms = {
                'translation': 5,  # coarse initial alignment
                'rigid': 3,         # medium refinement
                'affine': None,     # skip affine (we only want rigid)
                'sdr': None         # skip SDR (we only want rigid)
            }

            self.progress.emit(30, "Computing rigid registration (translation phase)...")
            self.progress.emit(35, "This will take 3-7 minutes. Please wait...")

            # Use MNE's compute_volume_registration with faster settings
            reg_affine, sdr_morph = mne.transforms.compute_volume_registration(
                ct_img,
                mri_img,
                pipeline='rigids',  # Use rigid registration (rotation + translation only)
                zooms=zooms,        # Faster settings
                verbose=True
            )

            self.progress.emit(70, "Registration complete! Applying transformation to CT...")

            # Apply registration to CT
            ct_aligned = mne.transforms.apply_volume_registration(
                ct_img, mri_img, reg_affine, cval='1%'
            )

            self.progress.emit(85, "Saving registered CT...")

            # Save registered CT
            os.makedirs(self.output_dir, exist_ok=True)
            ct_registered_path = os.path.join(self.output_dir, "ct_registered_to_mri.nii.gz")
            nib.save(ct_aligned, ct_registered_path)

            self.progress.emit(95, "Creating transformation matrix...")

            # Create MNE transform object
            trans = mne.transforms.Transform(fro="head", to="mri", trans=reg_affine)

            self.progress.emit(100, "Registration complete!")

            self.finished.emit(
                True,
                f"Registration successful! Saved to: {ct_registered_path}",
                trans,
                ct_registered_path,
            )
            
        except Exception as e:
            import traceback
            error_msg = f"Registration failed: {str(e)}\n\n{traceback.format_exc()}"
            self.finished.emit(False, error_msg, None, None)


class MplCanvas(FigureCanvasQTAgg):
    """Matplotlib canvas for embedding in Qt."""
    
    def __init__(self, parent=None, width=5, height=4, dpi=100):
        fig = Figure(figsize=(width, height), dpi=dpi)
        self.axes = fig.add_subplot(111)
        super().__init__(fig)


class SliceViewBox(pg.ViewBox):
    """ViewBox with drag-rectangle selection and hover/click signals."""

    sigDragSelect = Signal(str, object)
    sigMouseHover = Signal(str, object)
    sigMouseClick = Signal(str, object)

    def __init__(self, view_key, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.view_key = view_key
        self.setMouseEnabled(x=True, y=True)
        self._drag_start = None
        self._drag_rect_item = None
        self._drag_electrode_idx = None

    def mouseDragEvent(self, ev, axis=None):
        if ev.button() != Qt.LeftButton:
            return super().mouseDragEvent(ev, axis)
        if hasattr(self, "marking_check") and self.marking_check():
            ev.ignore()
            return
        ev.accept()
        pos = self.mapSceneToView(ev.scenePos())
        if ev.isStart():
            self._drag_electrode_idx = None
            if hasattr(self, "pick_electrode") and callable(self.pick_electrode):
                try:
                    self._drag_electrode_idx = self.pick_electrode(self.view_key, pos)
                except Exception:
                    self._drag_electrode_idx = None
            if self._drag_electrode_idx is None:
                self._drag_start = pos
                if self._drag_rect_item is None:
                    pen = pg.mkPen('#00ffff', width=1, style=Qt.DashLine)
                    self._drag_rect_item = QtWidgets.QGraphicsRectItem()
                    self._drag_rect_item.setPen(pen)
                    self._drag_rect_item.setBrush(QtGui.QBrush(QtGui.QColor(0, 255, 255, 30)))
                    self.addItem(self._drag_rect_item, ignoreBounds=True)

        if self._drag_electrode_idx is not None:
            if hasattr(self, "drag_electrode") and callable(self.drag_electrode):
                self.drag_electrode(self.view_key, self._drag_electrode_idx, pos, ev.isFinish())
            if ev.isFinish():
                self._drag_electrode_idx = None
            return

        if self._drag_start is None:
            return
        if ev.isFinish():
            rect = QtCore.QRectF(self._drag_start, pos).normalized()
            self.sigDragSelect.emit(self.view_key, rect)
            if self._drag_rect_item is not None:
                try:
                    self.removeItem(self._drag_rect_item)
                except Exception:
                    pass
                self._drag_rect_item = None
            self._drag_start = None
        else:
            rect = QtCore.QRectF(self._drag_start, pos).normalized()
            if self._drag_rect_item is not None:
                self._drag_rect_item.setRect(rect)

    def mouseClickEvent(self, ev):
        if ev.button() == Qt.LeftButton:
            pos = self.mapSceneToView(ev.scenePos())
            self.sigMouseClick.emit(self.view_key, pos)
            ev.accept()
        else:
            ev.ignore()

    def mouseMoveEvent(self, ev):
        pos = self.mapSceneToView(ev.scenePos())
        self.sigMouseHover.emit(self.view_key, pos)
        ev.accept()


class RegistrationDialog(QDialog):
    """Panel for CT-MRI coregistration with a single visualization view."""

    applied = Signal(object)
    panel_hidden = Signal()
    
    def __init__(self, mri_path, ct_path, output_dir, parent=None, as_panel=False, controller=None):
        super().__init__(parent)
        self.as_panel = as_panel
        self.controller = controller
        self.mri_path = mri_path
        self.ct_path = ct_path
        self.output_dir = output_dir
        self.trans = None
        self.ct_registered_path = None
        
        # Store volume data for slice browsing
        self.before_mri_img = None
        self.before_ct_img = None
        self.before_mri_data = None
        self.before_ct_data = None
        
        # Current slice indices
        self.before_slice_x = None
        self.before_slice_y = None
        self.before_slice_z = None
        
        # Visualization controls
        self.before_show_mri = True
        self.before_show_ct = True
        self.before_ct_alpha = 0.5
        self.before_show_shaft_labels = True
        self.before_show_axial = True
        self.before_parcellation_checkbox = None
        self.before_parcellation_overlay = False
        self.before_parcellation_volume = None
        self.before_parcellation_available = False
        self.before_parcellation_mismatch_warned = False
        self.before_parcellation_lut = {}
        self.before_selected_parcellation_label = None
        self.before_selected_parcellation_name = None
        self.before_show_coronal = True
        self.before_show_sagittal = True
        
        # Electrode marking
        self.before_electrodes = []  # List of Electrode objects
        self.electrode_marking_enabled = False
        self.before_hover_elec_idx = None
        self.electrode_space = "unregistered"
        self.electrode_display_space = self.electrode_space
        self.before_zoom_ax = None
        self.before_pg_views = {}
        self.before_pg_slices = {}
        self.before_pg_widgets = {}
        self.before_pg_containers = {}
        self.before_view_order = ["axial", "coronal", "sagittal"]
        self.before_view_buttons = {}
        self.before_minimized_rail = None
        self.before_minimized_layout = None
        self.before_zoom_items = {}
        self.before_zoom_active_view = None
        self._pending_detected_sync = None
        self._detected_sync_timer = None
        self.registration_running = False
        self.status_mri_chip = None
        self.status_ct_chip = None
        self.status_fs_chip = None
        self.status_reg_chip = None
        self.apply_btn = None
        self.cancel_btn = None
        self.save_transform_btn = None
        self.transform_mni_btn = None
        
        # Electrode shaft management
        self.before_shafts = {}  # Dict: shaft_name -> list of electrode indices
        self.shaft_counter = 0
        self.before_detected_positions = None
        self.before_axial_mip_positions = None
        self.fs_subject_path = None
        self.planning_worksheet_path = None
        self.planning_worksheet_mtime = None
        self.planning_trajectories = []
        self.planning_shaft_order = {}
        self.planning_schema_image_count = 0
        self.electrode_map_path = None
        self.map_shafts = {}
        self.map_trajectory_assignments = {}
        # Micro contacts from the .map (1.C / 1.D rows). Keyed by bundle name
        # (e.g. "mLAMY"). Populated by _refresh_electrode_map so Map Trajectories
        # can manually assign a micro bundle when target-name matching is missing.
        self.micro_map_shafts = {}
        # User-edited overrides: bundle name -> planning trajectory primary key.
        self.micro_map_trajectory_assignments = {}
        
        # Shaft visibility
        self.before_shaft_visibility = {}  # Dict: shaft_name -> bool (visible or not)
        self._before_shaft_toggle_off = False
        self._before_shaft_action_click = False
        
        # Maximum Intensity Projection
        self.before_show_mip = False
        # Visibility toggles
        self.before_detected_visible = True
        self.before_mark_cid = None
        self._suppress_transform_load = False
        self._before_shaft_last_clicked = None
        
        if self.as_panel:
            self.setWindowFlags(Qt.Widget)
        self.setWindowTitle("CT-MRI Coregistration")
        self.resize(1400, 900)
        
        self.setup_ui()
        if hasattr(self, "mri_path_display"):
            self.mri_path_display.setText(self.mri_path or "")
        if hasattr(self, "ct_path_display"):
            self.ct_path_display.setText(self.ct_path or "")
        if hasattr(self, "fs_subject_path_display") and self.controller:
            subject_id = getattr(self.controller, "custom_subject", None)
            subjects_dir = getattr(self.controller, "custom_subjects_dir", None)
            if subject_id and subjects_dir:
                self.fs_subject_path = os.path.join(subjects_dir, subject_id)
                self.fs_subject_path_display.setText(self.fs_subject_path)
        self._refresh_subject_view_combo()
        if self.fs_subject_path:
            self._load_before_parcellation_volume()
        self._sync_output_dir()
        self._refresh_planning_trajectories()
        self._refresh_electrode_map()
        self.load_and_display_before()

    def _current_electrode_space(self):
        """Return electrode space label based on registration state."""
        return "registered" if (self.trans is not None or self.ct_registered_path) else "unregistered"

    def _current_display_electrode_space(self):
        """Return the coordinate space currently used for 3D display/export."""
        return self.electrode_display_space or self._current_electrode_space()

    def _planning_search_paths(self):
        """Return paths used to discover the patient planning worksheet."""
        paths = [
            self.mri_path,
            self.ct_path,
            self.output_dir,
        ]
        if self.controller is not None:
            paths.append(getattr(self.controller, "patient_reg_folder", None))
            paths.append(getattr(self.controller, "mri_path", None))
            paths.append(getattr(self.controller, "ct_path", None))
        return paths

    def _log_planning_message(self, message):
        """Log planning worksheet messages after the UI log exists."""
        try:
            self.log(message)
        except Exception:
            pass

    def _refresh_planning_trajectories(self, *, log_missing=False):
        """Load planned trajectory names from the patient planning worksheet."""
        path = find_planning_worksheet(self._planning_search_paths())
        if not path:
            if self.planning_worksheet_path and log_missing:
                self._log_planning_message("No SEEG planning worksheet found near the current MRI/CT paths.")
            self.planning_worksheet_path = None
            self.planning_worksheet_mtime = None
            self.planning_trajectories = []
            self.planning_shaft_order = {}
            self.planning_schema_image_count = 0
            return []

        try:
            worksheet_mtime = os.path.getmtime(path)
        except OSError:
            worksheet_mtime = None

        if (
            path == self.planning_worksheet_path
            and worksheet_mtime == self.planning_worksheet_mtime
            and self.planning_trajectories
        ):
            return self.planning_trajectories

        try:
            trajectories = read_planning_worksheet(path)
        except Exception as exc:
            self.planning_worksheet_path = path
            self.planning_worksheet_mtime = worksheet_mtime
            self.planning_trajectories = []
            self.planning_shaft_order = {}
            self.planning_schema_image_count = 0
            self._log_planning_message(f"Could not read SEEG planning worksheet '{os.path.basename(path)}': {exc}")
            return []

        self.planning_worksheet_path = path
        self.planning_worksheet_mtime = worksheet_mtime
        self.planning_trajectories = trajectories
        self.planning_shaft_order = {}
        self.planning_schema_image_count = count_embedded_images(path)
        if trajectories:
            self._log_planning_message(
                f"Loaded {len(trajectories)} planned trajectories from {os.path.basename(path)}"
            )
            if self.planning_schema_image_count:
                self._log_planning_message(
                    f"Found {self.planning_schema_image_count} embedded planning images; "
                    "MCW naming-schema prior is available for shaft assignment."
                )
        else:
            self._log_planning_message(
                f"No trajectory rows found in SEEG planning worksheet: {os.path.basename(path)}"
            )
        return trajectories

    def _refresh_electrode_map(self, *, log_missing=False):
        """Load channel/contact labels from the patient .map file."""
        path = find_electrode_map(self._planning_search_paths())
        if not path:
            if self.electrode_map_path and log_missing:
                self._log_planning_message("No electrode .map file found near the current MRI/CT paths.")
            self.electrode_map_path = None
            self.map_shafts = {}
            self.map_trajectory_assignments = {}
            self.micro_map_shafts = {}
            self.micro_map_trajectory_assignments = {}
            return {}

        if path == self.electrode_map_path and self.map_shafts:
            self._prune_map_trajectory_assignments()
            self._infer_map_trajectory_assignments()
            # Refresh micro_map_shafts in case the file was edited on disk;
            # this is cheap and keeps the manual-assignment list current.
            self.micro_map_shafts = self._read_micro_map_shafts(path) or {}
            self._prune_micro_map_trajectory_assignments()
            return self.map_shafts

        try:
            map_shafts = read_electrode_map(path)
        except Exception as exc:
            self.electrode_map_path = path
            self.map_shafts = {}
            self.map_trajectory_assignments = {}
            self.micro_map_shafts = {}
            self.micro_map_trajectory_assignments = {}
            self._log_planning_message(f"Could not read electrode .map file '{os.path.basename(path)}': {exc}")
            return {}

        self.electrode_map_path = path
        self.map_shafts = map_shafts
        # Load 1.C/1.D micro rows from the same .map file alongside macro shafts.
        self.micro_map_shafts = self._read_micro_map_shafts(path) or {}
        self._prune_micro_map_trajectory_assignments()
        self._prune_map_trajectory_assignments()
        self._infer_map_trajectory_assignments()
        if map_shafts:
            self._log_planning_message(
                f"Loaded {len(map_shafts)} mapped electrode shafts from {os.path.basename(path)}"
            )
        else:
            self._log_planning_message(
                f"No usable 1.A/1.B L/R contact rows found in electrode .map file: {os.path.basename(path)}"
            )
        return map_shafts

    @staticmethod
    def _is_generated_shaft_name(name):
        """Whether a shaft name is an auto-generated S_# label."""
        return bool(re.match(r"^[Ss][-_]?\d+$", str(name or "")))

    @staticmethod
    def _generated_shaft_sort_key(name):
        match = re.match(r"^[Ss][-_]?(\d+)$", str(name or ""))
        if match:
            return (0, int(match.group(1)), str(name))
        return (1, str(name))

    @staticmethod
    def _trajectory_base_name(trajectory):
        return (
            str(trajectory.get("trajectory_name") or "").strip()
            or str(trajectory.get("label") or "").strip()
        )

    @staticmethod
    def _trajectory_label_key(label):
        """Return a label key that preserves side markers like '*'."""
        return re.sub(r"[^a-z0-9*]+", "", str(label or "").lower())

    @staticmethod
    def _trajectory_shaft_name(base_name, label):
        """Return the canonical shaft name for a worksheet trajectory."""
        base_name = str(base_name or "").strip()
        label = str(label or "").strip()
        if label and base_name and label != base_name:
            return f"{label} - {base_name}"
        return base_name or label or "Planned Shaft"

    @staticmethod
    def _trajectory_hemisphere(trajectory):
        """Return left/right from trajectory naming or worksheet label markers."""
        name = str(trajectory.get("trajectory_name") or "").strip()
        match = re.match(r"^([LR])(?:\b|[\s_-])", name, flags=re.IGNORECASE)
        if match:
            return "left" if match.group(1).upper() == "L" else "right"

        label = str(trajectory.get("label") or "").strip()
        if label:
            return "left" if "*" in label else "right"
        return None

    @staticmethod
    def _normalize_planning_name(value):
        """Normalize worksheet/shaft names for matching."""
        return re.sub(r"[^a-z0-9]+", "", str(value or "").lower())

    def _trajectory_name_candidates(self, trajectory):
        """Return display names that can identify a planning trajectory."""
        base_name = self._trajectory_base_name(trajectory)
        label = str(trajectory.get("label") or "").strip()
        trajectory_name = str(trajectory.get("trajectory_name") or "").strip()
        canonical_name = self._trajectory_shaft_name(base_name, label)
        candidates = [canonical_name, trajectory_name, base_name, label]
        primary_key = self._trajectory_primary_key(trajectory)
        for prefix, assigned_key in getattr(self, "map_trajectory_assignments", {}).items():
            if assigned_key == primary_key:
                candidates.append(prefix)

        names = []
        seen = set()
        for candidate in candidates:
            key = str(candidate or "").strip().casefold()
            if key and key not in seen:
                names.append(candidate)
                seen.add(key)
        return names

    def _trajectory_primary_key(self, trajectory):
        """Return a stable key for a planning trajectory."""
        name_key = self._normalize_planning_name(trajectory.get("trajectory_name"))
        label_key = self._trajectory_label_key(trajectory.get("label"))
        if not name_key and not label_key:
            return None

        hemisphere = self._trajectory_hemisphere(trajectory) or ""
        parts = [hemisphere, label_key, name_key]
        return "|".join(part for part in parts if part)

    def _trajectory_by_primary_key(self, primary_key):
        """Find a loaded planning trajectory by stable key."""
        if not primary_key:
            return None
        for trajectory in self.planning_trajectories or self._refresh_planning_trajectories():
            if self._trajectory_primary_key(trajectory) == primary_key:
                return trajectory
        return None

    def _trajectory_display_name(self, trajectory):
        """Return compact text for trajectory combo boxes."""
        name = str(trajectory.get("trajectory_name") or "").strip()
        label = str(trajectory.get("label") or "").strip()
        contacts = trajectory.get("contacts")
        parts = []
        if label:
            parts.append(label)
        if name:
            parts.append(name)
        text = " - ".join(parts) if parts else "Unnamed trajectory"
        if contacts:
            text = f"{text} ({contacts})"
        return text

    def _prune_map_trajectory_assignments(self):
        """Drop map assignments that no longer match loaded shafts/trajectories."""
        valid_prefixes = set(getattr(self, "map_shafts", {}) or {})
        valid_keys = {
            self._trajectory_primary_key(trajectory)
            for trajectory in (getattr(self, "planning_trajectories", []) or [])
        }
        valid_keys.discard(None)
        self.map_trajectory_assignments = {
            prefix: primary_key
            for prefix, primary_key in getattr(self, "map_trajectory_assignments", {}).items()
            if prefix in valid_prefixes and primary_key in valid_keys
        }

    def _prune_micro_map_trajectory_assignments(self):
        """Drop micro map assignments that no longer match loaded bundles/trajectories."""
        valid_prefixes = set(getattr(self, "micro_map_shafts", {}) or {})
        valid_keys = {
            self._trajectory_primary_key(trajectory)
            for trajectory in (getattr(self, "planning_trajectories", []) or [])
        }
        valid_keys.discard(None)
        self.micro_map_trajectory_assignments = {
            prefix: primary_key
            for prefix, primary_key in getattr(self, "micro_map_trajectory_assignments", {}).items()
            if prefix in valid_prefixes and primary_key in valid_keys
        }

    @staticmethod
    def _normalized_map_label_code(label):
        """Normalize worksheet labels for comparison with map prefixes."""
        text = str(label or "").upper()
        return re.sub(r"[^A-Z0-9]", "", text)

    def _infer_map_trajectory_assignments(self):
        """Auto-map map prefixes by side plus worksheet Label.

        Examples: LN -> Label N* on the left, RN -> Label N on the right,
        RF1 -> Label F1 on the right.
        """
        map_shafts = getattr(self, "map_shafts", {}) or {}
        trajectories = self.planning_trajectories or []
        if not map_shafts or not trajectories:
            return

        self._prune_map_trajectory_assignments()
        assigned_keys = set(self.map_trajectory_assignments.values())
        for prefix in sorted(map_shafts, key=lambda name: map_shafts[name].get("order", 0)):
            if prefix in self.map_trajectory_assignments:
                continue
            if len(prefix) < 2 or prefix[0] not in {"L", "R"}:
                continue
            hemisphere = "left" if prefix[0] == "L" else "right"
            label_code = prefix[1:]
            matches = []
            for trajectory in trajectories:
                primary_key = self._trajectory_primary_key(trajectory)
                if not primary_key or primary_key in assigned_keys:
                    continue
                if self._trajectory_hemisphere(trajectory) != hemisphere:
                    continue
                if self._normalized_map_label_code(trajectory.get("label")) == label_code:
                    matches.append((primary_key, trajectory))
            if len(matches) == 1:
                primary_key, _trajectory = matches[0]
                self.map_trajectory_assignments[prefix] = primary_key
                assigned_keys.add(primary_key)

    def _map_prefix_for_trajectory(self, trajectory):
        """Return mapped .map shaft prefix for a planning trajectory, if any."""
        primary_key = self._trajectory_primary_key(trajectory)
        if not primary_key:
            return None
        prefixes = [
            prefix for prefix, assigned_key in getattr(self, "map_trajectory_assignments", {}).items()
            if assigned_key == primary_key
        ]
        if not prefixes:
            return None
        planned_contacts = self._coerce_positive_int(trajectory.get("contacts"))
        if planned_contacts is not None:
            exact = [
                prefix for prefix in prefixes
                if self._coerce_positive_int(self.map_shafts.get(prefix, {}).get("contact_count")) == planned_contacts
            ]
            if exact:
                prefixes = exact
        return sorted(prefixes, key=lambda name: self.map_shafts.get(name, {}).get("order", 0))[0]

    def _map_contact_label(self, map_prefix, contact_number):
        """Return the recorded contact label for a mapped shaft/contact number."""
        entry = (getattr(self, "map_shafts", {}) or {}).get(map_prefix)
        if entry:
            for contact in entry.get("contacts", []):
                if self._coerce_positive_int(contact.get("number")) == int(contact_number):
                    return contact.get("label")
        return f"{map_prefix}{int(contact_number):02d}"

    def _planning_trajectory_for_name(self, name):
        """Find a planning trajectory matching a displayed shaft name."""
        text = str(name or "").strip()
        exact_key = text.casefold()
        normalized_key = self._normalize_planning_name(text)
        if not exact_key and not normalized_key:
            return None
        trajectories = self.planning_trajectories or self._refresh_planning_trajectories()

        exact_matches = []
        for trajectory in trajectories:
            if exact_key in {
                str(candidate or "").strip().casefold()
                for candidate in self._trajectory_name_candidates(trajectory)
            }:
                exact_matches.append(trajectory)
        if len(exact_matches) == 1:
            return exact_matches[0]
        if len(exact_matches) > 1:
            return None

        normalized_matches = []
        for trajectory in trajectories:
            candidate_keys = {
                self._normalize_planning_name(candidate)
                for candidate in self._trajectory_name_candidates(trajectory)
            }
            if normalized_key in candidate_keys:
                normalized_matches.append(trajectory)
        if len(normalized_matches) == 1:
            return normalized_matches[0]
        return None

    def _used_planning_trajectory_keys(self, shafts_dict):
        """Return worksheet trajectories already represented by shafts or electrode metadata."""
        used = set()
        for shaft_name, indices in (shafts_dict or {}).items():
            trajectory = self._trajectory_for_existing_shaft(shaft_name, indices)
            primary_key = self._trajectory_primary_key(trajectory) if trajectory is not None else None
            if primary_key:
                used.add(primary_key)

        for electrode in getattr(self, "before_electrodes", []):
            trajectory = self._planning_trajectory_for_electrode(electrode)
            primary_key = self._trajectory_primary_key(trajectory) if trajectory is not None else None
            if primary_key:
                used.add(primary_key)
                continue
            for value in (electrode.get("map_shaft"), electrode.get("shaft")):
                trajectory = self._planning_trajectory_for_name(value)
                primary_key = self._trajectory_primary_key(trajectory) if trajectory is not None else None
                if primary_key:
                    used.add(primary_key)
        return used

    def _unused_planning_trajectory_choices(self, shafts_dict):
        """Return unused planning trajectory names to show in shaft-choice lists."""
        trajectories = self.planning_trajectories or self._refresh_planning_trajectories()
        if not trajectories:
            return []

        used_keys = self._used_planning_trajectory_keys(shafts_dict)
        used_names = set((shafts_dict or {}).keys())
        choices = []
        for trajectory in trajectories:
            primary_key = self._trajectory_primary_key(trajectory)
            if not primary_key or primary_key in used_keys:
                continue
            display_name = self._unique_planned_shaft_name(
                self._trajectory_base_name(trajectory),
                trajectory.get("label"),
                used_names,
            )
            used_names.add(display_name)
            choices.append((display_name, trajectory))
        return choices

    def _planning_trajectory_from_indices(self, indices):
        """Infer a planning trajectory from electrodes already in a shaft."""
        return self._planning_trajectory_from_electrode_metadata(indices)

    def _planning_trajectory_for_metadata(self, trajectory_name=None, trajectory_label=None):
        """Resolve a worksheet trajectory from stored trajectory fields."""
        name_text = str(trajectory_name or "").strip()
        label_text = str(trajectory_label or "").strip()
        if not name_text and not label_text:
            return None

        trajectories = self.planning_trajectories or self._refresh_planning_trajectories()
        matches = []
        for trajectory in trajectories:
            candidate_name = str(trajectory.get("trajectory_name") or "").strip()
            candidate_label = str(trajectory.get("label") or "").strip()
            name_matches = not name_text or candidate_name.casefold() == name_text.casefold()
            label_matches = not label_text or candidate_label.casefold() == label_text.casefold()
            if name_matches and label_matches:
                matches.append(trajectory)
        if len(matches) == 1:
            return matches[0]

        canonical_name = self._trajectory_shaft_name(name_text, label_text)
        if canonical_name:
            return self._planning_trajectory_for_name(canonical_name)
        return None

    def _planning_trajectory_for_electrode(self, electrode):
        """Resolve a worksheet trajectory from one electrode's planning metadata."""
        if electrode is None:
            return None
        trajectory = self._planning_trajectory_for_metadata(
            electrode.get("trajectory_name"),
            electrode.get("trajectory_label"),
        )
        if trajectory is not None:
            return trajectory
        for value in (electrode.get("map_shaft"), electrode.get("shaft")):
            trajectory = self._planning_trajectory_for_name(value)
            if trajectory is not None:
                return trajectory
        return None

    def _planning_trajectory_from_electrode_metadata(self, indices, methods=None):
        """Infer one planning trajectory from electrode metadata in a shaft."""
        primary_keys = set()
        electrodes = getattr(self, "before_electrodes", [])
        method_filter = set(methods) if methods is not None else None
        for electrode_idx in indices or []:
            try:
                electrode_idx = int(electrode_idx)
            except (TypeError, ValueError):
                continue
            if electrode_idx < 0 or electrode_idx >= len(electrodes):
                continue
            electrode = electrodes[electrode_idx]
            if method_filter is not None and electrode.get("planning_assignment_method") not in method_filter:
                continue
            trajectory = self._planning_trajectory_for_electrode(electrode)
            primary_key = self._trajectory_primary_key(trajectory) if trajectory is not None else None
            if primary_key:
                primary_keys.add(primary_key)

        if len(primary_keys) != 1:
            return None
        return self._trajectory_by_primary_key(next(iter(primary_keys)))

    def _resolve_target_planning_trajectory(self, target_shaft, existing_indices=None):
        """Resolve worksheet metadata for a manual shaft assignment."""
        return (
            self._planning_trajectory_for_name(target_shaft)
            or self._planning_trajectory_from_indices(existing_indices)
        )

    @staticmethod
    def _clear_planning_metadata(electrode):
        """Remove worksheet assignment metadata from an electrode."""
        for key in (
            "trajectory_name",
            "trajectory_label",
            "planned_contacts",
            "trajectory_order",
            "trajectory_hemisphere",
            "planning_worksheet",
            "planning_assignment_method",
            "planning_assignment_cost",
            "trajectory_schema_ap",
            "trajectory_schema_si",
            "shaft_entry_ap",
            "shaft_entry_si",
            "map_shaft",
            "map_contact",
            "map_contact_number",
            "electrode_map",
        ):
            electrode.pop(key, None)

    def _clear_planning_metadata_for_indices(self, indices):
        electrodes = getattr(self, "before_electrodes", [])
        for electrode_idx in indices or []:
            try:
                electrode_idx = int(electrode_idx)
            except (TypeError, ValueError):
                continue
            if 0 <= electrode_idx < len(electrodes):
                self._clear_planning_metadata(electrodes[electrode_idx])

    @staticmethod
    def _clear_map_metadata(electrode):
        """Remove .map-derived contact metadata from an electrode."""
        for key in ("map_shaft", "map_contact", "map_contact_number", "electrode_map"):
            electrode.pop(key, None)

    def _apply_map_contact_metadata(self, electrode, map_prefix, contact_number):
        """Apply .map contact metadata and return the contact label."""
        if not map_prefix:
            self._clear_map_metadata(electrode)
            return None
        contact_label = self._map_contact_label(map_prefix, contact_number)
        electrode["map_shaft"] = map_prefix
        electrode["map_contact"] = contact_label
        electrode["map_contact_number"] = contact_number
        electrode["electrode_map"] = self.electrode_map_path
        return contact_label

    def _apply_planning_trajectory_to_indices(self, target_shaft, trajectory, indices, method):
        """Apply worksheet trajectory metadata to all contacts in a shaft."""
        if trajectory is None:
            return
        electrodes = getattr(self, "before_electrodes", [])
        valid_indices = []
        for electrode_idx in indices or []:
            try:
                electrode_idx = int(electrode_idx)
            except (TypeError, ValueError):
                continue
            if 0 <= electrode_idx < len(electrodes):
                valid_indices.append(electrode_idx)
        valid_indices = self._order_electrode_indices_distal_first(electrodes, valid_indices)

        label = str(trajectory.get("label") or "").strip()
        map_prefix = self._map_prefix_for_trajectory(trajectory)
        for contact_number, electrode_idx in enumerate(valid_indices, start=1):
            electrode = electrodes[electrode_idx]
            electrode.shaft = target_shaft
            electrode["shaft_contact_number"] = contact_number
            electrode["trajectory_name"] = trajectory.get("trajectory_name")
            electrode["trajectory_label"] = trajectory.get("label")
            electrode["planned_contacts"] = trajectory.get("contacts")
            electrode["trajectory_order"] = trajectory.get("order")
            electrode["trajectory_hemisphere"] = self._trajectory_hemisphere(trajectory)
            electrode["planning_worksheet"] = self.planning_worksheet_path
            electrode["planning_assignment_method"] = method
            electrode["planning_assignment_cost"] = None
            map_contact = self._apply_map_contact_metadata(electrode, map_prefix, contact_number)
            if map_contact:
                electrode.name = map_contact
            elif label:
                electrode.name = f"{label}{contact_number}"

        order = trajectory.get("order")
        if order is not None:
            self.planning_shaft_order[target_shaft] = order

    def _trajectory_for_existing_shaft(self, shaft_name, indices):
        """Resolve planning trajectory metadata for an existing shaft."""
        manual_methods = {"manual_shaft_edit", "manual_assignment"}
        trajectory = self._planning_trajectory_from_electrode_metadata(indices, methods=manual_methods)
        if trajectory is not None:
            return trajectory
        trajectory = self._planning_trajectory_from_electrode_metadata(indices)
        if trajectory is not None:
            return trajectory
        return self._planning_trajectory_for_name(shaft_name)

    def _apply_map_assignments_to_existing_electrodes(self):
        """Refresh .map contact labels on already-grouped electrodes."""
        electrodes = getattr(self, "before_electrodes", [])
        if not electrodes or not getattr(self, "before_shafts", None):
            return 0
        updated = 0
        for shaft_name, indices in list(self.before_shafts.items()):
            indices = self._order_electrode_indices_distal_first(electrodes, indices)
            self.before_shafts[shaft_name] = list(indices)
            trajectory = self._trajectory_for_existing_shaft(shaft_name, indices)
            if trajectory is None:
                for contact_number, electrode_idx in enumerate(indices, start=1):
                    if 0 <= electrode_idx < len(electrodes):
                        electrodes[electrode_idx]["shaft_contact_number"] = contact_number
                        self._clear_map_metadata(electrodes[electrode_idx])
                continue
            map_prefix = self._map_prefix_for_trajectory(trajectory)
            label = str(trajectory.get("label") or "").strip()
            for contact_number, electrode_idx in enumerate(indices, start=1):
                if electrode_idx < 0 or electrode_idx >= len(electrodes):
                    continue
                electrode = electrodes[electrode_idx]
                electrode["shaft_contact_number"] = contact_number
                map_contact = self._apply_map_contact_metadata(electrode, map_prefix, contact_number)
                if map_contact:
                    electrode.name = map_contact
                    updated += 1
                elif label:
                    electrode.name = f"{label}{contact_number}"
        return updated

    def _refresh_contact_numbers_for_shaft(self, shaft_name, indices):
        """Refresh contact names for one shaft using distal-first numbering."""
        electrodes = getattr(self, "before_electrodes", [])
        if not electrodes:
            return 0
        trajectory = self._trajectory_for_existing_shaft(shaft_name, indices)
        if trajectory is None:
            for contact_number, electrode_idx in enumerate(indices, start=1):
                if 0 <= electrode_idx < len(electrodes):
                    electrodes[electrode_idx]["shaft_contact_number"] = contact_number
            return 0

        map_prefix = self._map_prefix_for_trajectory(trajectory)
        label = str(trajectory.get("label") or "").strip()
        updated = 0
        for contact_number, electrode_idx in enumerate(indices, start=1):
            if electrode_idx < 0 or electrode_idx >= len(electrodes):
                continue
            electrode = electrodes[electrode_idx]
            electrode["shaft_contact_number"] = contact_number
            map_contact = self._apply_map_contact_metadata(electrode, map_prefix, contact_number)
            if map_contact:
                electrode.name = map_contact
                updated += 1
            elif label:
                electrode.name = f"{label}{contact_number}"
                updated += 1
        return updated

    def _electrode_voxel_affine(self):
        """Return the voxel-to-world affine for hemisphere assignment."""
        for img_attr in ("before_mri_img", "before_ct_img"):
            img = getattr(self, img_attr, None)
            affine = getattr(img, "affine", None)
            if affine is not None:
                return np.asarray(affine, dtype=float)

        for path in (getattr(self, "mri_path", None), getattr(self, "ct_path", None)):
            if not path or not os.path.exists(path):
                continue
            try:
                return np.asarray(nib.load(path).affine, dtype=float)
            except Exception:
                continue
        return None

    def _shaft_hemisphere(self, electrodes, indices):
        """Classify a detected shaft by centroid world X in RAS coordinates."""
        coords = []
        for electrode_idx in indices:
            if electrode_idx >= len(electrodes):
                continue
            electrode = electrodes[electrode_idx]
            x = electrode.get("x")
            y = electrode.get("y")
            z = electrode.get("z")
            if x is None or y is None or z is None:
                continue
            coords.append([float(x), float(y), float(z)])
        if not coords:
            return None

        centroid = np.mean(np.asarray(coords, dtype=float), axis=0)
        affine = self._electrode_voxel_affine()
        if affine is None:
            return None

        world = np.append(centroid, 1.0) @ affine.T
        world_x = float(world[0])
        if abs(world_x) < 1e-6:
            return None
        return "right" if world_x > 0 else "left"

    def _shaft_world_points(self, electrodes, indices):
        """Return electrode coordinates for a shaft in RAS/world space."""
        affine = self._electrode_voxel_affine()
        if affine is None:
            return np.empty((0, 3), dtype=float)

        coords = []
        for electrode_idx in indices:
            if electrode_idx >= len(electrodes):
                continue
            electrode = electrodes[electrode_idx]
            x = electrode.get("x")
            y = electrode.get("y")
            z = electrode.get("z")
            if x is None or y is None or z is None:
                continue
            vox = np.array([float(x), float(y), float(z), 1.0], dtype=float)
            coords.append((vox @ affine.T)[:3])
        if not coords:
            return np.empty((0, 3), dtype=float)
        return np.asarray(coords, dtype=float)

    def _distal_first_order_for_positions(self, positions):
        """Return indices that order one shaft from distal/deep to proximal/entry."""
        pts = np.asarray(positions, dtype=float)
        if pts.size == 0:
            return np.array([], dtype=int)
        pts = np.reshape(pts, (-1, 3))
        if len(pts) <= 1:
            return np.arange(len(pts), dtype=int)

        centroid, direction = fit_shaft_axis_pca(pts)
        if centroid is None or direction is None or np.linalg.norm(direction) < 1e-6:
            return np.arange(len(pts), dtype=int)
        direction = direction / np.linalg.norm(direction)
        t_values = (pts - centroid) @ direction
        ascending = True

        world = self._positions_to_world(pts)
        if world is not None and len(world) == len(pts):
            order = np.argsort(t_values)
            first_lateral = float(abs(world[order[0], 0]))
            last_lateral = float(abs(world[order[-1], 0]))
            # SEEG contact 01 is conventionally the distal/deep tip. The entry
            # side is usually the most lateral end, so sort away from that end.
            if np.isfinite(first_lateral) and np.isfinite(last_lateral):
                if last_lateral + 0.5 < first_lateral:
                    ascending = False

        return np.argsort(t_values if ascending else -t_values)

    def _order_electrode_indices_distal_first(self, electrodes, indices):
        """Order electrode indices so contact 01 is the distal shaft contact."""
        valid = []
        coords = []
        for electrode_idx in indices or []:
            try:
                electrode_idx = int(electrode_idx)
            except (TypeError, ValueError):
                continue
            if electrode_idx < 0 or electrode_idx >= len(electrodes):
                continue
            electrode = electrodes[electrode_idx]
            try:
                coord = (float(electrode.x), float(electrode.y), float(electrode.z))
            except (TypeError, ValueError, AttributeError):
                continue
            if not np.all(np.isfinite(coord)):
                continue
            valid.append(electrode_idx)
            coords.append(coord)

        if len(valid) <= 1:
            return valid
        order = self._distal_first_order_for_positions(np.asarray(coords, dtype=float))
        return [valid[int(order_idx)] for order_idx in order]

    def _order_shaft_dict_distal_first(self, electrodes, shafts):
        """Return a shaft dict whose contact lists are distal-to-proximal."""
        return {
            shaft_name: self._order_electrode_indices_distal_first(electrodes, indices)
            for shaft_name, indices in (shafts or {}).items()
        }

    def _shaft_entry_world_point(self, electrodes, indices):
        """Estimate the entry-side contact as the most lateral shaft point."""
        world = self._shaft_world_points(electrodes, indices)
        if world.size == 0:
            return None
        lateral_idx = int(np.argmax(np.abs(world[:, 0])))
        return world[lateral_idx]

    def _shaft_schema_point(self, electrodes, indices):
        """Return a 2D AP/SI point used for schema-prior assignment."""
        entry = self._shaft_entry_world_point(electrodes, indices)
        if entry is None:
            return None
        return (float(entry[1]), float(entry[2]))

    def _unique_planned_shaft_name(self, base_name, label, used_names):
        """Return a unique planned shaft display name."""
        candidate = self._trajectory_shaft_name(base_name, label)
        if candidate not in used_names:
            return candidate

        suffix = 2
        while f"{candidate} ({suffix})" in used_names:
            suffix += 1
        return f"{candidate} ({suffix})"

    def _assign_planning_names_to_shafts(self, electrodes, shafts, shaft_visibility, *, log=True):
        """Rename generated shaft labels with trajectory names from the planning worksheet."""
        shafts = self._order_shaft_dict_distal_first(electrodes, shafts)
        trajectories = self.planning_trajectories or self._refresh_planning_trajectories()
        if not trajectories or not shafts:
            return shafts, shaft_visibility, 0

        generated_names = [
            name for name in shafts.keys()
            if self._is_generated_shaft_name(name)
        ]
        generated_names.sort(key=self._generated_shaft_sort_key)
        if not generated_names:
            return shafts, shaft_visibility, 0

        manual_methods = {"manual_shaft_edit", "manual_assignment"}
        manually_assigned_trajectories = {}
        for name in generated_names:
            trajectory = self._planning_trajectory_from_electrode_metadata(
                shafts.get(name, []),
                methods=manual_methods,
            )
            if trajectory is not None:
                manually_assigned_trajectories[name] = trajectory
        generated_names = [
            name for name in generated_names
            if name not in manually_assigned_trajectories
        ]
        if not generated_names:
            return shafts, shaft_visibility, 0

        generated_by_hemi = {"left": [], "right": [], None: []}
        for name in generated_names:
            generated_by_hemi.setdefault(
                self._shaft_hemisphere(electrodes, shafts.get(name, [])),
                []
            ).append(name)

        trajectories_by_hemi = {"left": [], "right": [], None: []}
        for trajectory in trajectories:
            trajectories_by_hemi.setdefault(
                self._trajectory_hemisphere(trajectory),
                []
            ).append(trajectory)

        rename_map = {}
        trajectory_by_old_name = {}
        assignment_meta_by_old_name = {}
        order_map = dict(getattr(self, "planning_shaft_order", {}))
        used_names = set(shafts.keys())
        contact_count_assignment_count = 0

        def _assign_one(old_name, trajectory, method, *, schema_point=None, detected_point=None, cost=None):
            if old_name in rename_map:
                return False
            if trajectory is None:
                return False
            new_name = self._unique_planned_shaft_name(
                self._trajectory_base_name(trajectory),
                trajectory.get("label"),
                used_names - {old_name},
            )
            rename_map[old_name] = new_name
            trajectory_by_old_name[old_name] = trajectory
            assignment_meta_by_old_name[old_name] = {
                "method": method,
                "schema_point": schema_point,
                "detected_point": detected_point,
                "cost": cost,
            }
            used_names.discard(old_name)
            used_names.add(new_name)
            order = trajectory.get("order")
            if order is not None:
                order_map[new_name] = order
            return True

        def _assign_pairs(old_names, planned, method="hemisphere_order"):
            assigned = 0
            for old_name, trajectory in zip(old_names, planned):
                if _assign_one(old_name, trajectory, method):
                    assigned += 1
            return assigned

        def _trajectory_contact_count(trajectory):
            return self._coerce_positive_int(trajectory.get("contacts"))

        def _detected_contact_count(old_name):
            return len(shafts.get(old_name, []))

        def _assign_by_contact_count(old_names, planned, *, unique_only=False):
            """Assign exact detected/planned contact-count matches within one hemisphere."""
            remaining_old = [name for name in old_names if name not in rename_map]
            remaining_planned = [trajectory for trajectory in planned if id(trajectory) not in assigned_trajectory_ids]
            if not remaining_old or not remaining_planned:
                return 0

            old_by_count = {}
            for old_name in remaining_old:
                count = _detected_contact_count(old_name)
                if count > 0:
                    old_by_count.setdefault(count, []).append(old_name)

            planned_by_count = {}
            for trajectory in remaining_planned:
                count = _trajectory_contact_count(trajectory)
                if count is not None:
                    planned_by_count.setdefault(count, []).append(trajectory)

            assigned = 0
            for count in sorted(set(old_by_count) & set(planned_by_count)):
                old_matches = [
                    name for name in old_by_count[count]
                    if name not in rename_map
                ]
                planned_matches = [
                    trajectory for trajectory in planned_by_count[count]
                    if id(trajectory) not in assigned_trajectory_ids
                ]
                if not old_matches or not planned_matches:
                    continue
                if unique_only and (len(old_matches) != 1 or len(planned_matches) != 1):
                    continue

                pair_count = min(len(old_matches), len(planned_matches))
                for old_name, trajectory in zip(old_matches[:pair_count], planned_matches[:pair_count]):
                    if _assign_one(old_name, trajectory, "hemisphere_contact_count"):
                        assigned += 1
                        assigned_trajectory_ids.add(id(trajectory))
            return assigned

        def _assign_with_schema(old_names, planned):
            has_schema_prior = bool(getattr(self, "planning_schema_image_count", 0))
            assigned = _assign_by_contact_count(old_names, planned, unique_only=has_schema_prior)
            if not has_schema_prior:
                remaining_old = [name for name in old_names if name not in rename_map]
                remaining_planned = [
                    trajectory for trajectory in planned
                    if id(trajectory) not in assigned_trajectory_ids
                ]
                return assigned + _assign_pairs(remaining_old, remaining_planned)
            detected_items = [
                {
                    "name": old_name,
                    "point": self._shaft_schema_point(electrodes, shafts.get(old_name, [])),
                    "order": idx,
                }
                for idx, old_name in enumerate(old_names)
                if old_name not in rename_map
            ]
            planned_for_schema = [
                trajectory for trajectory in planned
                if id(trajectory) not in assigned_trajectory_ids
            ]
            schema_assignments = assign_schema_prior(detected_items, planned_for_schema)
            for assignment in schema_assignments:
                old_name = assignment["old_name"]
                trajectory = assignment["trajectory"]
                if _assign_one(
                    old_name,
                    trajectory,
                    assignment.get("method", "schema_prior"),
                    schema_point=assignment.get("schema_point"),
                    detected_point=assignment.get("detected_point"),
                    cost=assignment.get("cost"),
                ):
                    assigned += 1
                    assigned_trajectory_ids.add(id(trajectory))

            remaining_old = [name for name in old_names if name not in rename_map]
            remaining_planned = [
                trajectory for trajectory in planned
                if id(trajectory) not in assigned_trajectory_ids
            ]
            assigned += _assign_pairs(remaining_old, remaining_planned)
            return assigned

        assigned_trajectory_ids = {
            id(trajectory)
            for trajectory in manually_assigned_trajectories.values()
        }
        left_right_detected = generated_by_hemi["left"] or generated_by_hemi["right"]
        if left_right_detected:
            applied_count = 0
            applied_count += _assign_with_schema(
                generated_by_hemi["right"],
                trajectories_by_hemi["right"],
            )
            applied_count += _assign_with_schema(
                generated_by_hemi["left"],
                trajectories_by_hemi["left"],
            )
            contact_count_assignment_count = sum(
                1 for meta in assignment_meta_by_old_name.values()
                if meta.get("method") == "hemisphere_contact_count"
            )
            applied_count += _assign_pairs(
                generated_by_hemi[None],
                trajectories_by_hemi[None],
            )
        else:
            applied_count = _assign_pairs(generated_names, trajectories)

        if applied_count == 0:
            return shafts, shaft_visibility, 0

        new_shafts = {}
        new_visibility = {}
        for old_name, indices in shafts.items():
            new_name = rename_map.get(old_name, old_name)
            ordered_indices = self._order_electrode_indices_distal_first(electrodes, indices)
            new_shafts[new_name] = list(ordered_indices)
            new_visibility[new_name] = shaft_visibility.get(old_name, True)

            trajectory = trajectory_by_old_name.get(old_name)
            if trajectory is None:
                continue
            assignment_meta = assignment_meta_by_old_name.get(old_name, {})
            label = str(trajectory.get("label") or "").strip()
            map_prefix = self._map_prefix_for_trajectory(trajectory)
            for contact_number, electrode_idx in enumerate(ordered_indices, start=1):
                if electrode_idx >= len(electrodes):
                    continue
                electrode = electrodes[electrode_idx]
                electrode.shaft = new_name
                electrode["trajectory_name"] = trajectory.get("trajectory_name")
                electrode["trajectory_label"] = trajectory.get("label")
                electrode["planned_contacts"] = trajectory.get("contacts")
                electrode["trajectory_order"] = trajectory.get("order")
                electrode["trajectory_hemisphere"] = self._trajectory_hemisphere(trajectory)
                electrode["planning_worksheet"] = self.planning_worksheet_path
                electrode["planning_assignment_method"] = assignment_meta.get("method")
                electrode["planning_assignment_cost"] = assignment_meta.get("cost")
                schema_point = assignment_meta.get("schema_point")
                detected_point = assignment_meta.get("detected_point")
                if schema_point is not None:
                    electrode["trajectory_schema_ap"] = schema_point[0]
                    electrode["trajectory_schema_si"] = schema_point[1]
                if detected_point is not None:
                    electrode["shaft_entry_ap"] = detected_point[0]
                    electrode["shaft_entry_si"] = detected_point[1]
                map_contact = self._apply_map_contact_metadata(electrode, map_prefix, contact_number)
                if map_contact:
                    electrode.name = map_contact
                elif label:
                    electrode.name = f"{label}{contact_number}"

        self.planning_shaft_order = order_map
        if log:
            self._log_planning_message(
                f"Applied planning worksheet names to {applied_count} detected shafts."
            )
            schema_count = sum(
                1 for meta in assignment_meta_by_old_name.values()
                if meta.get("method") == "schema_prior"
            )
            if schema_count:
                self._log_planning_message(
                    f"Used MCW naming-schema spatial prior for {schema_count} shaft assignments."
                )
            if contact_count_assignment_count:
                self._log_planning_message(
                    f"Matched {contact_count_assignment_count} shaft assignments by detected/planned contact count within hemisphere."
                )
            if len(trajectories) != len(generated_names):
                self._log_planning_message(
                    f"Planning worksheet has {len(trajectories)} trajectories; "
                    f"detected grouping has {len(generated_names)} generated shafts."
                )
            if left_right_detected:
                for hemi, label in (("right", "right"), ("left", "left")):
                    planned_count = len(trajectories_by_hemi[hemi])
                    detected_count = len(generated_by_hemi[hemi])
                    if planned_count != detected_count:
                        self._log_planning_message(
                            f"Planning worksheet has {planned_count} {label} trajectories; "
                            f"detected grouping has {detected_count} {label} shafts."
                        )
                if generated_by_hemi[None]:
                    self._log_planning_message(
                        f"Could not determine hemisphere for {len(generated_by_hemi[None])} detected shafts."
                    )
        return new_shafts, new_visibility, applied_count

    def _has_talairach_transform(self):
        """Whether a FreeSurfer/FastSurfer subject is loaded for MNI mapping."""
        if self.controller is None:
            return False
        subject_id = getattr(self.controller, "custom_subject", None)
        subjects_dir = getattr(self.controller, "custom_subjects_dir", None)
        return bool(subject_id and subjects_dir)

    def _available_electrodes_for_mni(self):
        """Return panel electrodes, falling back to controller-held detections."""
        if self.before_electrodes:
            return list(self.before_electrodes)
        if self.controller is None:
            return []
        electrodes = getattr(self.controller, "detected_electrodes", None)
        return list(electrodes or [])

    def _update_mni_button_state(self):
        """Enable/label the MRI->MNI toggle when the prerequisites are present."""
        btn = getattr(self, "transform_mni_btn", None)
        if btn is None:
            return
        in_mni_mode = self._current_display_electrode_space() == "mni_tal"
        has_electrodes = bool(self._available_electrodes_for_mni())
        has_subject = self._has_talairach_transform()
        can_transform = bool(
            has_electrodes
            and has_subject
            and self.controller is not None
        )
        btn.setText("Use Native MRI" if in_mni_mode else "MRI->MNI")
        if in_mni_mode:
            btn.setEnabled(has_electrodes)
            btn.setToolTip("Switch the 3D viewer back to native MRI-space electrode coordinates")
        else:
            btn.setEnabled(can_transform)
            if not has_electrodes:
                btn.setToolTip("Detect or load electrodes before transforming to MNI")
            elif not has_subject:
                btn.setToolTip("Load a FreeSurfer/FastSurfer subject before transforming to MNI")
            else:
                btn.setToolTip(
                    "Transform the current electrode coordinates to MNI Talairach space for fsaverage viewing"
                )

    def _transform_electrodes_to_mni(self):
        """Toggle the 3D viewer between native MRI-space and stored MNI-space electrodes."""
        electrodes = self._available_electrodes_for_mni()
        if not electrodes:
            QMessageBox.information(self, "No Electrodes", "Detect or load electrodes first.")
            return
        if not self.before_electrodes:
            self.before_electrodes = list(electrodes)

        native_space = self._current_electrode_space()
        if self._current_display_electrode_space() == "mni_tal":
            self.electrode_display_space = native_space
            self._schedule_detected_sync(self.before_electrodes)
            self._update_mni_button_state()
            self.log("Switched 3D viewer back to native MRI-space electrode coordinates.")
            return

        if self.controller is None or not hasattr(self.controller, "transform_detected_electrodes_to_mni"):
            QMessageBox.warning(self, "MNI Transform Unavailable", "The 3D viewer controller does not support MRI->MNI transforms.")
            return

        if not self._has_talairach_transform():
            QMessageBox.information(
                self,
                "FreeSurfer Subject Required",
                "Load a FreeSurfer/FastSurfer subject with talairach.xfm before transforming to MNI."
            )
            return

        try:
            transformed = self.controller.transform_detected_electrodes_to_mni(
                self.before_electrodes,
                update_display=False,
            )
        except Exception as exc:
            QMessageBox.warning(self, "MRI->MNI Failed", str(exc))
            self.log(f"MRI->MNI transform failed: {exc}")
            return

        self.electrode_display_space = "mni_tal"
        self.controller.detected_electrodes = list(self.before_electrodes)
        self.controller.detected_electrode_space = "mni_tal"
        self.controller.detected_electrode_native_space = native_space
        if hasattr(self.controller, "set_subject_view_mode"):
            try:
                self.controller.set_subject_view_mode("fsaverage")
            except Exception:
                pass
        self._schedule_detected_sync(self.before_electrodes)
        self._update_mni_button_state()
        self.log(f"Stored MNI Talairach coordinates for {transformed} electrodes and switched the 3D viewer to fsaverage.")
        
    def setup_ui(self):
        """Set up the user interface."""
        layout = QVBoxLayout()

        # Title
        title = QLabel("CT-MRI Coregistration")
        title.setStyleSheet("font-size: 18px; font-weight: bold; padding: 10px;")
        title.setAlignment(Qt.AlignCenter)
        layout.addWidget(title)
        layout.addWidget(self._create_status_row())

        # Main visualization + controls in a vertical splitter so the plot gets
        # more vertical space by default while controls remain accessible below.
        splitter = QtWidgets.QSplitter(Qt.Vertical)

        # Canvas (top / primary area)
        self.before_canvas = QWidget()
        self.before_canvas_layout = QHBoxLayout()
        self.before_canvas_layout.setContentsMargins(0, 0, 0, 0)
        self.before_canvas_layout.setSpacing(6)

        before_controls = self._create_slice_controls()
        self.before_pg_widgets = {}
        self.before_pg_containers = {}
        view_labels = {"axial": "Axial", "coronal": "Coronal", "sagittal": "Sagittal"}
        for view_key in self.before_view_order:
            glw = pg.GraphicsLayoutWidget()
            glw.setSizePolicy(QtWidgets.QSizePolicy.Expanding, QtWidgets.QSizePolicy.Expanding)
            self.before_pg_widgets[view_key] = glw

            container = QWidget()
            container.setSizePolicy(QtWidgets.QSizePolicy.Expanding, QtWidgets.QSizePolicy.Expanding)
            container_layout = QVBoxLayout()
            container_layout.setContentsMargins(0, 0, 0, 0)
            container_layout.setSpacing(4)
            header = QWidget()
            header_layout = QHBoxLayout()
            header_layout.setContentsMargins(0, 0, 0, 0)
            header_layout.setSpacing(4)
            header_label = QLabel(view_labels.get(view_key, view_key.title()))
            header_label.setStyleSheet("font-weight: bold;")
            header_layout.addWidget(header_label)
            header_layout.addStretch(1)
            header.setLayout(header_layout)
            container_layout.addWidget(header, 0)
            container_layout.addWidget(glw, 1)
            slider_widget = before_controls.get(view_key)
            if slider_widget is not None:
                container_layout.addWidget(slider_widget, 0)
            container.setLayout(container_layout)
            self.before_pg_containers[view_key] = container
            self.before_canvas_layout.addWidget(container, 1)

        self.before_canvas.setLayout(self.before_canvas_layout)
        self._setup_pg_views()
        canvas_container = QWidget()
        canvas_layout = QVBoxLayout()
        canvas_layout.setContentsMargins(0, 0, 0, 0)
        canvas_layout.addWidget(self.before_canvas)
        self.before_minimized_rail = QWidget()
        self.before_minimized_layout = QHBoxLayout()
        self.before_minimized_layout.setContentsMargins(0, 0, 0, 0)
        self.before_minimized_layout.setSpacing(6)
        self.before_minimized_layout.setAlignment(Qt.AlignLeft)
        self.before_minimized_rail.setLayout(self.before_minimized_layout)
        for view_key in self.before_view_order:
            btn = QtWidgets.QToolButton()
            btn.setText(view_labels.get(view_key, view_key.title()))
            btn.setCheckable(True)
            btn.setChecked(True)
            btn.setToolTip("Toggle this view")
            btn.toggled.connect(
                lambda checked, key=view_key: self._toggle_view_visibility(key, checked)
            )
            self.before_view_buttons[view_key] = btn
            self.before_minimized_layout.addWidget(btn)
        self.before_minimized_layout.addStretch(1)
        self.before_minimized_rail.setVisible(True)
        canvas_layout.addWidget(self.before_minimized_rail)
        canvas_container.setLayout(canvas_layout)
        splitter.addWidget(canvas_container)

        # Controls container (bottom of splitter)
        controls_container = QWidget()
        controls_layout = QVBoxLayout()
        controls_layout.setContentsMargins(6, 6, 6, 6)
        controls_layout.setSpacing(6)

        # Visualization controls (we'll place these in the HBox below)
        before_viz_controls, before_shaft_group, before_electrode_group = (
            self._create_visualization_controls()
        )

        # --- MRI / CT file selectors (moved here from the main localizer) ---
        files_widget = QWidget()
        files_layout = QGridLayout()
        files_layout.setContentsMargins(0, 0, 0, 0)
        files_layout.setHorizontalSpacing(8)
        files_layout.setVerticalSpacing(6)

        mri_label = QLabel("MRI:")
        mri_label.setMinimumWidth(60)
        self.mri_path_display = QtWidgets.QLineEdit()
        self.mri_path_display.setReadOnly(True)
        self.mri_path_display.setPlaceholderText("No MRI file selected")
        self.mri_path_display.setMinimumHeight(22)
        mri_browse_btn = QPushButton("Browse...")
        mri_browse_btn.clicked.connect(self.browse_mri_file)
        mri_browse_btn.setMinimumWidth(84)
        mri_browse_btn.setMaximumWidth(110)

        ct_label = QLabel("Postop CT:")
        ct_label.setMinimumWidth(60)
        self.ct_path_display = QtWidgets.QLineEdit()
        self.ct_path_display.setReadOnly(True)
        self.ct_path_display.setPlaceholderText("No CT file selected")
        self.ct_path_display.setMinimumHeight(22)
        ct_browse_btn = QPushButton("Browse...")
        ct_browse_btn.clicked.connect(self.browse_ct_file)
        ct_browse_btn.setMinimumWidth(84)
        ct_browse_btn.setMaximumWidth(110)

        fs_label = QLabel("FreeSurfer Subject:")
        fs_label.setMinimumWidth(60)
        self.fs_subject_path_display = QtWidgets.QLineEdit()
        self.fs_subject_path_display.setReadOnly(True)
        self.fs_subject_path_display.setPlaceholderText("No FreeSurfer subject selected")
        self.fs_subject_path_display.setMinimumHeight(22)
        fs_browse_btn = QPushButton("Browse...")
        fs_browse_btn.clicked.connect(self.browse_freesurfer_subject)
        fs_browse_btn.setMinimumWidth(84)
        fs_browse_btn.setMaximumWidth(110)

        files_layout.addWidget(mri_label, 0, 0)
        files_layout.addWidget(self.mri_path_display, 0, 1)
        files_layout.addWidget(mri_browse_btn, 0, 2)
        files_layout.addWidget(ct_label, 1, 0)
        files_layout.addWidget(self.ct_path_display, 1, 1)
        files_layout.addWidget(ct_browse_btn, 1, 2)
        files_layout.addWidget(fs_label, 2, 0)
        files_layout.addWidget(self.fs_subject_path_display, 2, 1)
        files_layout.addWidget(fs_browse_btn, 2, 2)
        subject_view_label = QLabel("Brain:")
        subject_view_label.setMinimumWidth(60)
        self.subject_view_combo = QtWidgets.QComboBox()
        self.subject_view_combo.addItem("Subject", "subject")
        self.subject_view_combo.addItem("fsaverage", "fsaverage")
        self.subject_view_combo.setToolTip("Switch between subject surfaces and fsaverage")
        self.subject_view_combo.currentIndexChanged.connect(self._on_subject_view_changed)
        files_layout.addWidget(subject_view_label, 3, 0)
        files_layout.addWidget(self.subject_view_combo, 3, 1, 1, 2)
        files_layout.setColumnStretch(0, 0)
        files_layout.setColumnStretch(1, 1)
        files_layout.setColumnStretch(2, 0)

        files_widget.setLayout(files_layout)
        files_widget.setMaximumHeight(180)
        files_widget.setSizePolicy(QtWidgets.QSizePolicy.Expanding, QtWidgets.QSizePolicy.Fixed)

        # Registration log (placed next to file selectors)
        self.log_text = QTextEdit()
        self.log_text.setReadOnly(True)
        self.log_text.setMaximumHeight(120)
        self.log_text.setStyleSheet(
            "background-color: #fbfbfb; font-family: monospace; font-size: 12px; color: #222; "
            "border: 1px solid #dcdcdc; border-radius: 4px; padding: 4px;"
        )
        log_group = QGroupBox("Registration Log")
        log_group_layout = QVBoxLayout()
        log_group_layout.setContentsMargins(4, 4, 4, 4)
        log_group_layout.addWidget(self.log_text)
        log_group.setLayout(log_group_layout)
        log_group.setSizePolicy(QtWidgets.QSizePolicy.Expanding, QtWidgets.QSizePolicy.Fixed)
        log_group.setMaximumHeight(180)
        log_group.setMinimumHeight(180)
        log_group.setMinimumWidth(240)

        files_log_widget = QWidget()
        files_log_layout = QHBoxLayout()
        files_log_layout.setContentsMargins(0, 0, 0, 0)
        files_log_layout.setSpacing(8)
        files_log_layout.addWidget(files_widget, 2)
        files_log_layout.addWidget(log_group, 3)
        files_log_widget.setLayout(files_log_layout)

        controls_layout.addWidget(files_log_widget)

        method_widget = QWidget()
        method_layout = QHBoxLayout()
        method_layout.setContentsMargins(0, 0, 0, 0)
        method_layout.setSpacing(6)
        method_label = QLabel("Registration:")
        method_combo = QtWidgets.QComboBox()
        method_combo.addItem("MNE (Rigid)", "mne_rigid")
        method_combo.addItem("ANTs (Rigid)", "ants_rigid")
        method_combo.addItem("ANTs (Affine)", "ants_affine")
        ants_default = method_combo.findData("ants_rigid")
        if ants_default >= 0:
            method_combo.setCurrentIndex(ants_default)
        method_combo.setToolTip("Select the registration backend.")
        method_layout.addWidget(method_label)
        method_layout.addWidget(method_combo)

        self.register_btn = QPushButton("Register")
        self.register_btn.setStyleSheet(
            "QPushButton { padding: 8px 12px; font-size: 13px; background-color: #4CAF50; "
            "color: white; border-radius: 6px; }"
        )
        self.register_btn.clicked.connect(self._on_register_clicked)
        btn_min_width = self.register_btn.sizeHint().width()
        btn_min_height = self.register_btn.sizeHint().height()
        self.register_btn.setMinimumWidth(btn_min_width)
        self.register_btn.setMinimumHeight(btn_min_height)

        method_combo.setSizePolicy(QtWidgets.QSizePolicy.Maximum, QtWidgets.QSizePolicy.Fixed)
        self.register_btn.setSizePolicy(QtWidgets.QSizePolicy.Maximum, QtWidgets.QSizePolicy.Fixed)

        self.progress_bar = QProgressBar()
        self.progress_bar.setRange(0, 100)
        self.progress_bar.setValue(0)
        self.progress_bar.setMaximumHeight(14)
        self.progress_bar.setTextVisible(True)
        self.progress_bar.setFormat("Ready to register")
        self.progress_bar.setSizePolicy(QtWidgets.QSizePolicy.Expanding, QtWidgets.QSizePolicy.Fixed)

        method_layout.addWidget(self.register_btn)
        method_layout.addWidget(self.progress_bar, 1)
        method_widget.setLayout(method_layout)
        controls_layout.addWidget(method_widget)
        self.registration_method_combo = method_combo

        # Transform matrix display (hidden; kept for saving/debug)
        self.transform_text = QTextEdit()
        self.transform_text.setReadOnly(True)
        self.transform_text.setLineWrapMode(QTextEdit.NoWrap)
        self.transform_text.setStyleSheet(
            "background-color: #f8f8f8; font-family: monospace; font-size: 11px; "
            "border: 1px solid #dcdcdc; border-radius: 4px; padding: 4px;"
        )
        identity_matrix = np.eye(4)
        matrix_str = "Identity Matrix (No Registration):\n"
        matrix_str += self._format_matrix(identity_matrix)
        self.transform_text.setPlainText(matrix_str)

        # Put viz controls, shafts, and electrodes side-by-side.
        hbox_widget = QWidget()
        hbox_layout = QHBoxLayout()
        hbox_layout.setContentsMargins(0, 0, 0, 0)
        hbox_layout.setSpacing(8)

        before_viz_controls.setSizePolicy(QtWidgets.QSizePolicy.Preferred, QtWidgets.QSizePolicy.Preferred)
        before_shaft_group.setSizePolicy(QtWidgets.QSizePolicy.Preferred, QtWidgets.QSizePolicy.Preferred)
        before_electrode_group.setSizePolicy(QtWidgets.QSizePolicy.Preferred, QtWidgets.QSizePolicy.Preferred)

        hbox_layout.addWidget(before_viz_controls)
        hbox_layout.addWidget(before_shaft_group)
        hbox_layout.addWidget(before_electrode_group)

        stretch = DEFAULT_CONTROLS_HBOX_STRETCH
        hbox_layout.setStretch(0, stretch["controls"])
        hbox_layout.setStretch(1, stretch["shafts"])
        hbox_layout.setStretch(2, stretch["electrodes"])

        hbox_widget.setLayout(hbox_layout)
        controls_layout.addWidget(hbox_widget)

        controls_layout.addWidget(self._create_action_bar())

        controls_container.setLayout(controls_layout)
        splitter.addWidget(controls_container)

        # Give the canvas more initial proportion of the splitter
        splitter.setSizes([800, 220])

        layout.addWidget(splitter)
        self.setLayout(layout)
        self.setStyleSheet(
            "QGroupBox { font-weight: bold; border: 1px solid #dcdcdc; border-radius: 4px; margin-top: 6px; }"
            "QGroupBox::title { subcontrol-origin: margin; left: 8px; padding: 0 3px 0 3px; }"
            "QPushButton { padding: 6px 10px; font-size: 12px; }"
            "QLineEdit, QTextEdit, QListWidget { background-color: #fbfbfb; border: 1px solid #dcdcdc; border-radius: 4px; padding: 4px; }"
            "QCheckBox { spacing: 6px; }"
        )
        self.registration_complete = False
        self._update_accept_button()
        self._update_status_chips()

    def _make_status_chip(self, text, tooltip):
        chip = QLabel(text)
        chip.setAlignment(Qt.AlignCenter)
        chip.setToolTip(tooltip)
        chip.setMinimumHeight(20)
        chip.setMinimumWidth(110)
        chip.setStyleSheet(
            "QLabel { background-color: #f2f2f2; color: #555555; "
            "border: 1px solid #d0d0d0; border-radius: 10px; padding: 2px 8px; "
            "font-size: 11px; }"
        )
        return chip

    def _set_status_chip(self, chip, text, state):
        colors = {
            "ok": ("#e6f4ea", "#1e7e34", "#a8d5b9"),
            "warn": ("#fff4e5", "#a05a00", "#f0c36d"),
            "info": ("#e8f1ff", "#3056d3", "#b4c8ff"),
            "busy": ("#eaf4ff", "#1f6feb", "#8ec5ff"),
        }
        bg, fg, border = colors.get(state, ("#f2f2f2", "#555555", "#d0d0d0"))
        chip.setText(text)
        chip.setStyleSheet(
            "QLabel { "
            f"background-color: {bg}; color: {fg}; border: 1px solid {border}; "
            "border-radius: 10px; padding: 2px 8px; font-size: 11px; }"
        )

    def _create_status_row(self):
        row = QWidget()
        row_layout = QHBoxLayout()
        row_layout.setContentsMargins(6, 0, 6, 0)
        row_layout.setSpacing(6)

        label = QLabel("Status:")
        label.setStyleSheet("font-weight: bold; color: #555555;")
        row_layout.addWidget(label)

        self.status_mri_chip = self._make_status_chip("MRI: --", "MRI file status")
        self.status_ct_chip = self._make_status_chip("CT: --", "CT file status")
        self.status_fs_chip = self._make_status_chip("FS: --", "FreeSurfer subject status")
        self.status_reg_chip = self._make_status_chip("Reg: --", "Registration status")

        row_layout.addWidget(self.status_mri_chip)
        row_layout.addWidget(self.status_ct_chip)
        row_layout.addWidget(self.status_fs_chip)
        row_layout.addWidget(self.status_reg_chip)
        row_layout.addStretch(1)
        row.setLayout(row_layout)

        self._update_status_chips()
        return row

    def _create_action_bar(self):
        actions_widget = QWidget()
        actions_layout = QHBoxLayout()
        actions_layout.setContentsMargins(0, 0, 0, 0)
        actions_layout.setSpacing(6)

        self.save_transform_btn = QPushButton("Save Transform")
        self.save_transform_btn.setToolTip("Save the current transform matrix")
        self.save_transform_btn.clicked.connect(self.save_transformation)
        self.save_transform_btn.setEnabled(False)
        actions_layout.addWidget(self.save_transform_btn)

        actions_layout.addStretch(1)

        apply_label = "Apply to Viewer" if self.as_panel else "Apply"
        close_label = "Close Panel" if self.as_panel else "Cancel"

        self.apply_btn = QPushButton(apply_label)
        self.apply_btn.setToolTip("Apply registration results and close the panel")
        self.apply_btn.clicked.connect(self._apply_results)
        self.apply_btn.setEnabled(False)
        actions_layout.addWidget(self.apply_btn)

        self.cancel_btn = QPushButton(close_label)
        self.cancel_btn.clicked.connect(self._on_cancel_clicked)
        actions_layout.addWidget(self.cancel_btn)

        actions_widget.setLayout(actions_layout)
        self.accept_btn = self.apply_btn
        return actions_widget

    def _update_status_chips(self):
        if not all([self.status_mri_chip, self.status_ct_chip, self.status_fs_chip, self.status_reg_chip]):
            return

        has_mri = bool(getattr(self, "mri_path", None))
        has_ct = bool(getattr(self, "ct_path", None))
        has_fs = bool(getattr(self, "fs_subject_path", None))
        if not has_fs and self.controller is not None:
            has_fs = bool(
                getattr(self.controller, "custom_subject", None)
                and getattr(self.controller, "custom_subjects_dir", None)
            )

        self._set_status_chip(
            self.status_mri_chip,
            "MRI: Loaded" if has_mri else "MRI: Missing",
            "ok" if has_mri else "warn",
        )
        self._set_status_chip(
            self.status_ct_chip,
            "CT: Loaded" if has_ct else "CT: Missing",
            "ok" if has_ct else "warn",
        )
        self._set_status_chip(
            self.status_fs_chip,
            "FS: Loaded" if has_fs else "FS: Optional",
            "ok" if has_fs else "info",
        )

        has_registration = bool(
            getattr(self, "registration_complete", False)
            or getattr(self, "trans", None) is not None
            or getattr(self, "ct_registered_path", None)
        )
        if getattr(self, "registration_running", False):
            reg_text = "Reg: Running"
            reg_state = "busy"
        elif has_registration:
            reg_text = "Reg: Complete"
            reg_state = "ok"
        elif has_mri and has_ct:
            reg_text = "Reg: Ready"
            reg_state = "info"
        else:
            reg_text = "Reg: Waiting"
            reg_state = "warn"
        self._set_status_chip(self.status_reg_chip, reg_text, reg_state)

    def _update_action_buttons(self):
        has_mri = bool(getattr(self, "mri_path", None))
        has_ct = bool(getattr(self, "ct_path", None))
        has_registration = bool(
            getattr(self, "registration_complete", False)
            or getattr(self, "trans", None) is not None
            or getattr(self, "ct_registered_path", None)
        )
        can_apply = has_mri and has_ct and not getattr(self, "registration_running", False)

        if self.apply_btn is not None:
            apply_label = "Apply to Viewer" if self.as_panel else "Apply"
            skip_label = "Use CT As-Is" if self.as_panel else "Use CT As-Is"
            if can_apply:
                if has_registration:
                    self.apply_btn.setText(apply_label)
                    self.apply_btn.setToolTip("Apply registered CT/transform to the viewer")
                else:
                    self.apply_btn.setText(skip_label)
                    self.apply_btn.setToolTip("Skip registration and use the CT as already aligned")
                self.apply_btn.setEnabled(True)
            else:
                self.apply_btn.setText(apply_label)
                self.apply_btn.setToolTip("Load MRI and CT to enable apply")
                self.apply_btn.setEnabled(False)

        if self.save_transform_btn is not None:
            self.save_transform_btn.setEnabled(bool(getattr(self, "trans", None)) and not self.registration_running)
        self._update_mni_button_state()

    def _setup_pg_views(self):
        """Initialize the PyQtGraph slice views."""
        self.before_pg_views = {}
        view_order = [("axial", "Axial"), ("coronal", "Coronal"), ("sagittal", "Sagittal")]

        for col, (view_key, view_label) in enumerate(view_order):
            glw = self.before_pg_widgets.get(view_key)
            if glw is None:
                glw = pg.GraphicsLayoutWidget()
                glw.setSizePolicy(QtWidgets.QSizePolicy.Expanding, QtWidgets.QSizePolicy.Expanding)
                self.before_pg_widgets[view_key] = glw
                container = self.before_pg_containers.get(view_key)
                if container is None:
                    container = QWidget()
                    container_layout = QVBoxLayout()
                    container_layout.setContentsMargins(0, 0, 0, 0)
                    container_layout.setSpacing(4)
                    container.setLayout(container_layout)
                    self.before_pg_containers[view_key] = container
                    self.before_canvas_layout.addWidget(container, 1)
                container_layout = container.layout()
                if container_layout is None:
                    container_layout = QVBoxLayout()
                    container_layout.setContentsMargins(0, 0, 0, 0)
                    container_layout.setSpacing(4)
                    container.setLayout(container_layout)
                if glw.parent() is None:
                    container_layout.insertWidget(0, glw, 1)
            glw.clear()

            view_box = SliceViewBox(view_key)
            view_box.setAspectLocked(True)
            view_box.invertY(False)
            view_box.marking_check = lambda: bool(
                getattr(self, "before_electrode_checkbox", None)
                and self.before_electrode_checkbox.isChecked()
            )
            glw.addItem(view_box, row=0, col=0)

            colorbar = pg.ColorBarItem(
                values=(0.0, 1.0),
                colorMap=self._get_ct_colormap(),
                interactive=False,
                orientation='horizontal',
                width=8,
            )
            colorbar.axis.setStyle(tickFont=QtGui.QFont("Sans", 7))
            colorbar.setVisible(False)
            glw.addItem(colorbar, row=1, col=0)
            try:
                glw.ci.layout.setRowStretchFactor(0, 1)
                glw.ci.layout.setRowStretchFactor(1, 0)
                glw.ci.layout.setRowMinimumHeight(1, 20)
            except Exception:
                pass

            mri_item = pg.ImageItem()
            ct_item = pg.ImageItem()
            parcellation_item = pg.ImageItem()
            mri_item.setOpts(axisOrder='row-major')
            ct_item.setOpts(axisOrder='row-major')
            parcellation_item.setOpts(axisOrder='row-major')
            ct_item.setOpacity(0.6)
            parcellation_item.setZValue(10)
            view_box.addItem(mri_item)
            view_box.addItem(ct_item)
            view_box.addItem(parcellation_item)

            scatter = pg.ScatterPlotItem()
            scatter.setPxMode(False)
            scatter.setZValue(30)
            view_box.addItem(scatter)

            vline = pg.InfiniteLine(angle=90, pen=pg.mkPen('c', width=1, style=Qt.DashLine))
            hline = pg.InfiniteLine(angle=0, pen=pg.mkPen('c', width=1, style=Qt.DashLine))
            vline.setZValue(25)
            hline.setZValue(25)
            view_box.addItem(vline)
            view_box.addItem(hline)

            left_label = "L"
            right_label = "R"
            if view_key == "sagittal":
                left_label = "P"
                right_label = "A"
            labels = {
                "left": pg.TextItem(left_label, color='w', anchor=(0, 0.5)),
                "right": pg.TextItem(right_label, color='w', anchor=(1, 0.5)),
                "top": pg.TextItem("", color='w', anchor=(0.5, 1)),
                "bottom": pg.TextItem("", color='w', anchor=(0.5, 0)),
            }
            for item in labels.values():
                view_box.addItem(item)

            view_box.sigMouseHover.connect(self._on_before_pg_hover)
            view_box.sigMouseClick.connect(self._on_before_pg_click)
            view_box.sigDragSelect.connect(self._on_before_pg_drag_select)
            view_box.pick_electrode = self._pick_pg_electrode
            view_box.drag_electrode = self._drag_pg_electrode

            self.before_pg_views[view_key] = {
                "view": view_box,
                "mri": mri_item,
                "ct": ct_item,
                "parcellation": parcellation_item,
                "scatter": scatter,
                "vline": vline,
                "hline": hline,
                "labels": labels,
                "colorbar": colorbar,
                "title": view_label
            }
            # No zoom overlay in PyQtGraph views.
    
    def browse_mri_file(self):
        """Open file browser to select MRI file (used in-dialog)."""
        file_path, _ = QtWidgets.QFileDialog.getOpenFileName(
            None,
            "Select MRI File",
            "",
            "NIfTI Files (*.nii *.nii.gz);;MGZ Files (*.mgz);;All Files (*.*)"
        )
        if file_path:
            self.mri_path = file_path
            try:
                self.mri_path_display.setText(file_path)
            except Exception:
                pass
            self._sync_output_dir()
            self._refresh_planning_trajectories()
            self._refresh_electrode_map()
            self._update_button_states()
            self._set_progress_message(f"MRI loaded - {os.path.basename(file_path)}")
            if self.controller is not None:
                self.controller.mri_path = file_path
                if hasattr(self.controller, "update_button_states"):
                    self.controller.update_button_states()
                if hasattr(self.controller, "coreg_status_label"):
                    self.controller.coreg_status_label.setText(
                        f"Status: MRI loaded - {os.path.basename(file_path)}"
                    )
            # Load and display MRI immediately
            try:
                self._load_mri_and_display(file_path)
            except Exception as e:
                self.log(f"Error loading MRI: {e}")

    def browse_ct_file(self):
        """Open file browser to select postoperative CT file (used in-dialog)."""
        file_path, _ = QtWidgets.QFileDialog.getOpenFileName(
            None,
            "Select Postoperative CT File",
            "",
            "NIfTI Files (*.nii *.nii.gz);;DICOM Files (*.dcm);;All Files (*.*)"
        )
        if file_path:
            self.ct_path = file_path
            try:
                self.ct_path_display.setText(file_path)
            except Exception:
                pass
            self._refresh_planning_trajectories()
            self._refresh_electrode_map()
            self._update_button_states()
            self._set_progress_message(f"CT loaded - {os.path.basename(file_path)}")
            if self.controller is not None:
                self.controller.ct_path = file_path
                if hasattr(self.controller, "update_button_states"):
                    self.controller.update_button_states()
                if hasattr(self.controller, "coreg_status_label"):
                    self.controller.coreg_status_label.setText(
                        f"Status: CT loaded - {os.path.basename(file_path)}"
                    )
            # Load and display CT immediately (overlay if MRI present)
            try:
                self._load_ct_and_display(file_path)
            except Exception as e:
                self.log(f"Error loading CT: {e}")

    def browse_freesurfer_subject(self):
        """Open file browser to select FreeSurfer subject folder."""
        subject_dir = QtWidgets.QFileDialog.getExistingDirectory(
            None,
            "Select FreeSurfer/FastSurfer Subject Folder",
            ""
        )
        if not subject_dir:
            return

        subject_dir = os.path.abspath(subject_dir)
        self.fs_subject_path_display.setText(subject_dir)
        self.fs_subject_path = subject_dir
        self._update_status_chips()

        if self.controller is None:
            return

        surf_dir = os.path.join(subject_dir, "surf")
        mri_dir = os.path.join(subject_dir, "mri")
        if not (os.path.isdir(surf_dir) and os.path.isdir(mri_dir)):
            QtWidgets.QMessageBox.warning(
                None,
                "Invalid Subject Folder",
                "Selected folder does not look like a FreeSurfer/FastSurfer subject.\n\n"
                "Expected subfolders:\n  - surf\n  - mri"
            )
            return

        subject_id = os.path.basename(subject_dir)
        subjects_dir = os.path.dirname(subject_dir)
        self.controller._load_freesurfer_subject(subjects_dir, subject_id)
        if hasattr(self.controller, "_update_classify_button_state"):
            self.controller._update_classify_button_state()
        self._refresh_subject_view_combo()
        self.fs_subject_path = subject_dir
        self._load_before_parcellation_volume()
        self._update_mni_button_state()

    def _sync_output_dir(self):
        """Update output directory based on current MRI selection."""
        if self.controller is not None and getattr(self.controller, "patient_reg_folder", None):
            self.output_dir = self.controller.patient_reg_folder
            return
        if self.mri_path:
            self.output_dir = str(Path(self.mri_path).parent / "coregistration")

    def _load_mri_and_display(self, file_path):
        """Load MRI file and update the before-view display immediately."""
        try:
            import traceback
            img = nib.load(file_path)
            data = img.get_fdata(dtype=np.float32)

            # Store images and data
            self.before_mri_img = img
            self.before_mri_data = np.asarray(data)

            # Initialize slice indices to volume center
            sx, sy, sz = self.before_mri_data.shape
            self.before_slice_x = sx // 2
            self.before_slice_y = sy // 2
            self.before_slice_z = sz // 2

            # Prepare CT placeholder if missing
            if getattr(self, 'before_ct_data', None) is None:
                self.before_ct_data = np.zeros_like(self.before_mri_data)

            # If a CT path exists, resample CT to the newly loaded MRI grid.
            if self.ct_path and os.path.exists(self.ct_path):
                self._load_ct_and_display(self.ct_path)
                return

            # Update sliders and labels if present
            try:
                self.before_sag_slider.setMaximum(self.before_mri_data.shape[0] - 1)
                self.before_sag_slider.setValue(self.before_slice_x)
                self.before_sag_slider.setEnabled(True)
                self.before_cor_slider.setMaximum(self.before_mri_data.shape[1] - 1)
                self.before_cor_slider.setValue(self.before_slice_y)
                self.before_cor_slider.setEnabled(True)
                self.before_ax_slider.setMaximum(self.before_mri_data.shape[2] - 1)
                self.before_ax_slider.setValue(self.before_slice_z)
                self.before_ax_slider.setEnabled(True)
                self.before_sag_label.setText(f"Sagittal Slice: {self.before_slice_x}/{self.before_mri_data.shape[0]-1}")
                self.before_cor_label.setText(f"Coronal Slice: {self.before_slice_y}/{self.before_mri_data.shape[1]-1}")
                self.before_ax_label.setText(f"Axial Slice: {self.before_slice_z}/{self.before_mri_data.shape[2]-1}")
            except Exception:
                pass

            # If CT exists on disk and was previously loaded, try to keep it aligned; otherwise overlay zeros
            try:
                if self.fs_subject_path:
                    self._load_before_parcellation_volume()
                self._plot_slices(self.before_mri_data, self.before_ct_data, self.before_canvas,
                                    "Unaligned CT Overlaid on MRI",
                                    self.before_slice_x, self.before_slice_y, self.before_slice_z)
                self.log(f"Loaded MRI: {os.path.basename(file_path)}")
            except Exception:
                self.log("Error plotting MRI after load:\n" + traceback.format_exc())

        except Exception as e:
            self.log(f"Error loading MRI: {e}")

    def _load_ct_and_display(self, file_path):
        """Load CT file and display it, overlaying onto MRI if MRI is present."""
        try:
            import traceback
            ct_img = nib.load(file_path)
            ct_data = ct_img.get_fdata(dtype=np.float32)
            self.before_ct_img = ct_img

            # If MRI present, resample CT to MRI space/shape
            if getattr(self, 'before_mri_img', None) is not None and getattr(self, 'before_mri_data', None) is not None:
                ct_res = None
                transform_matrix = None
                if not getattr(self, "_suppress_transform_load", False):
                    transform_matrix = self._load_transform_matrix(file_path)
                if transform_matrix is not None:
                    try:
                        ct_aligned = mne.transforms.apply_volume_registration(
                            ct_img, self.before_mri_img, transform_matrix, cval='1%'
                        )
                        ct_res = np.asarray(ct_aligned.get_fdata(dtype=np.float32))
                        self.trans = mne.transforms.Transform(
                            fro="head", to="mri", trans=transform_matrix
                        )
                        self.registration_complete = True
                        self._update_transform_display(transform_matrix, "Loaded Transform Matrix")
                        self._update_accept_button()
                        self._update_button_states()
                        if hasattr(self, "register_btn"):
                            self.register_btn.setText("Reset")
                            self.register_btn.setStyleSheet(
                                "QPushButton { padding: 8px 12px; font-size: 13px; "
                                "background-color: #ff9800; color: white; border-radius: 6px; }"
                            )
                        self.log(f"Applied saved transform: {self._transform_path(file_path)}")
                    except Exception as e:
                        self.log(f"Failed to apply saved transform: {e}")

                try:
                    # Prefer nibabel processing if available
                    try:
                        if ct_res is None:
                            from nibabel.processing import resample_from_to
                            target = (self.before_mri_img.shape, self.before_mri_img.affine)
                            resampled = resample_from_to(ct_img, target)
                            ct_res = np.asarray(resampled.get_fdata(dtype=np.float32))
                    except Exception:
                        # Fall back to simple zoom/resampling to match shapes
                        if ct_res is None:
                            zoom = np.array(self.before_mri_data.shape) / np.array(ct_data.shape)
                            ct_res = ndi.zoom(ct_data, zoom, order=1)
                except Exception:
                    self.log("CT resampling failed; using crude fallback")
                    ct_res = np.zeros_like(self.before_mri_data)

                self.before_ct_data = np.asarray(ct_res)
                # Keep slices centered on MRI
                sx, sy, sz = self.before_mri_data.shape
                self.before_slice_x = sx // 2
                self.before_slice_y = sy // 2
                self.before_slice_z = sz // 2

            else:
                # No MRI: display CT alone (use CT as the primary volume)
                self.before_ct_data = np.asarray(ct_data)
                # Use CT volume as MRI placeholder for visualization
                self.before_mri_data = np.zeros_like(self.before_ct_data)
                sx, sy, sz = self.before_ct_data.shape
                self.before_slice_x = sx // 2
                self.before_slice_y = sy // 2
                self.before_slice_z = sz // 2

            # Update sliders and labels if present
            try:
                self.before_sag_slider.setMaximum(self.before_mri_data.shape[0] - 1)
                self.before_sag_slider.setValue(self.before_slice_x)
                self.before_sag_slider.setEnabled(True)
                self.before_cor_slider.setMaximum(self.before_mri_data.shape[1] - 1)
                self.before_cor_slider.setValue(self.before_slice_y)
                self.before_cor_slider.setEnabled(True)
                self.before_ax_slider.setMaximum(self.before_mri_data.shape[2] - 1)
                self.before_ax_slider.setValue(self.before_slice_z)
                self.before_ax_slider.setEnabled(True)
                self.before_sag_label.setText(f"Sagittal Slice: {self.before_slice_x}/{self.before_mri_data.shape[0]-1}")
                self.before_cor_label.setText(f"Coronal Slice: {self.before_slice_y}/{self.before_mri_data.shape[1]-1}")
                self.before_ax_label.setText(f"Axial Slice: {self.before_slice_z}/{self.before_mri_data.shape[2]-1}")
            except Exception:
                pass

            # Update threshold spinbox ranges if possible
            try:
                ct_min = np.nanmin(self.before_ct_data)
                ct_max = np.nanmax(self.before_ct_data)
                ct_mid = min(max(ct_max * 0.4, ct_min), ct_max)
                self.before_threshold_lower.setRange(ct_min, ct_max)
                self.before_threshold_upper.setRange(ct_min, ct_max)
                self.before_threshold_lower.setValue(ct_mid)
                self.before_threshold_upper.setValue(ct_max)
            except Exception:
                pass

            try:
                self._load_electrode_locations(file_path)
            except Exception:
                pass

            # Plot overlay
            try:
                if self.fs_subject_path and getattr(self, "before_mri_img", None) is not None:
                    self._load_before_parcellation_volume()
                title = "Registered CT Overlaid on MRI" if (
                    getattr(self, "trans", None) is not None
                    and getattr(self, "registration_complete", False)
                ) else "Unaligned CT Overlaid on MRI"
                self._plot_slices(self.before_mri_data, self.before_ct_data, self.before_canvas,
                                    title,
                                    self.before_slice_x, self.before_slice_y, self.before_slice_z)
                self.log(f"Loaded CT: {os.path.basename(file_path)}")
            except Exception:
                self.log("Error plotting CT after load:\n" + traceback.format_exc())

        except Exception as e:
            self.log(f"Error loading CT: {e}")

    def _update_button_states(self):
        """Enable/disable buttons based on file selection and registration state."""
        has_mri = getattr(self, 'mri_path', None) is not None
        has_ct = getattr(self, 'ct_path', None) is not None
        try:
            self.register_btn.setEnabled(has_mri and has_ct)
        except Exception:
            pass
        self._update_action_buttons()
        self._update_status_chips()

    def save_transformation(self):
        """Save current transformation matrix to file (text or .npy)."""
        if self.trans is None:
            QMessageBox.warning(self, "No Transform", "No transformation available to save.")
            return
        if isinstance(self.trans, dict) and self.trans.get("ants_fwdtransforms"):
            QMessageBox.information(
                self,
                "ANTs Transform",
                "ANTs transform files were generated by the registration backend.\n"
                "Use the ANTs transform files directly from the output directory.",
            )
            return
        default_name = self._transform_path(self.ct_path) if self.ct_path else "transform.npy"
        fname, _ = QtWidgets.QFileDialog.getSaveFileName(
            None,
            "Save Transform",
            default_name,
            "Transform files (*.npy *.txt);;All Files (*.*)"
        )
        if not fname:
            return
        arr = None
        try:
            arr = np.array(self.trans['trans']) if isinstance(self.trans, dict) and 'trans' in self.trans else np.array(self.trans)
        except Exception:
            try:
                arr = np.array(self.trans)
            except Exception:
                arr = None
        if arr is None:
            QMessageBox.critical(self, "Save Error", "Could not extract transform matrix to save.")
            return
        try:
            if fname.lower().endswith('.npy'):
                np.save(fname, arr)
            else:
                np.savetxt(fname, arr, fmt='%0.6f')
            QMessageBox.information(self, "Saved", f"Transform saved to: {fname}")
        except Exception as e:
            QMessageBox.critical(self, "Save Error", f"Error saving transform: {e}")

    def _save_transform_for_ct(self):
        """Auto-save the current transform alongside the CT if available."""
        if not self.ct_path or self.trans is None:
            return
        if isinstance(self.trans, dict) and self.trans.get("ants_fwdtransforms"):
            return
        dest = self._transform_path(self.ct_path)
        if not dest:
            return
        arr = None
        try:
            arr = np.array(self.trans['trans']) if isinstance(self.trans, dict) and 'trans' in self.trans else np.array(self.trans)
        except Exception:
            try:
                arr = np.array(self.trans)
            except Exception:
                arr = None
        if arr is None:
            self.log("Could not extract transform matrix for auto-save.")
            return
        try:
            np.save(dest, arr)
            self.log(f"Saved transform: {dest}")
        except Exception as e:
            self.log(f"Failed to auto-save transform: {e}")

    def _export_value(self, value):
        """Return JSON-serializable scalar values."""
        if isinstance(value, np.generic):
            return value.item()
        return value

    def _is_synthetic_micro_export_row(self, row):
        """Return whether an export row is a derived micro contact, not a real marked electrode."""
        if not row:
            return False
        synthetic_method = str(row.get("synthetic_3d_method") or "").strip().lower()
        if synthetic_method in {"macro_first_contact", "unresolved_macro_source"}:
            return True

        micro_values = (
            row.get("map_shaft"),
            row.get("map_contact"),
            row.get("shaft"),
            row.get("name"),
        )
        has_micro_label = any(self._is_micro_like_label(value) for value in micro_values)
        if not has_micro_label:
            return False

        assignment_method = str(row.get("planning_assignment_method") or "").strip().lower()
        if assignment_method in {"synthetic_micro_target", "manual_micro_map_trajectory"}:
            return True

        detected_text = str(row.get("detected") or "").strip().lower()
        is_detected = row.get("detected") is True or detected_text in {"true", "1", "yes", "y", "t"}
        return not is_detected and any(
            self._is_micro_like_label(value)
            for value in (row.get("map_shaft"), row.get("map_contact"))
        )

    def _restore_stale_micro_source_row(self, row):
        """Repair old exports where a macro source row was renamed to a micro shaft."""
        if not row or not self._is_micro_like_label(row.get("shaft")):
            return row
        if self._is_micro_like_label(row.get("name")) or self._is_micro_like_label(row.get("map_contact")):
            return row
        map_shaft = row.get("map_shaft")
        if not map_shaft or self._is_micro_like_label(map_shaft):
            return row
        row["shaft"] = map_shaft
        return row

    def _stale_micro_source_contact_key(self, row):
        """Return a macro contact key for old exports that copied source rows into micro shafts."""
        if not row or not self._is_micro_like_label(row.get("shaft")):
            return None
        if self._is_micro_like_label(row.get("name")) or self._is_micro_like_label(row.get("map_contact")):
            return None
        map_shaft = row.get("map_shaft")
        if not map_shaft or self._is_micro_like_label(map_shaft):
            return None
        for value in (row.get("map_contact"), row.get("name")):
            key = self._normalize_electrode_match_key(value)
            if key:
                return key
        return None

    def _clean_stale_micro_source_rows(self, rows):
        """Drop duplicate old micro-source copies, or restore them when they are the only source."""
        if not rows:
            return rows, 0, 0

        original_map_shaft_keys = set()
        for row in rows:
            if self._is_micro_like_label(row.get("shaft")):
                continue
            for value in (row.get("map_shaft"), row.get("shaft")):
                if self._is_micro_like_label(value):
                    continue
                key = self._normalize_electrode_match_key(value)
                if key:
                    original_map_shaft_keys.add(key)

        cleaned_rows = []
        restored_count = 0
        dropped_duplicate_count = 0
        restored_map_shaft_keys = set()
        for row in rows:
            contact_key = self._stale_micro_source_contact_key(row)
            if not contact_key:
                cleaned_rows.append(row)
                continue
            map_shaft_key = self._normalize_electrode_match_key(row.get("map_shaft"))
            if map_shaft_key in original_map_shaft_keys or map_shaft_key in restored_map_shaft_keys:
                dropped_duplicate_count += 1
                continue
            detected_text = str(row.get("detected") or "").strip().lower()
            is_detected = row.get("detected") is True or detected_text in {"true", "1", "yes", "y", "t"}
            if not is_detected and map_shaft_key:
                # Old buggy exports cloned every source contact into the micro shaft.
                # Keep only one fallback source row when no real source shaft exists.
                restored_map_shaft_keys.add(map_shaft_key)
            before_shaft = row.get("shaft")
            row = self._restore_stale_micro_source_row(row)
            if row.get("shaft") != before_shaft:
                restored_count += 1
            cleaned_rows.append(row)
        return cleaned_rows, restored_count, dropped_duplicate_count

    def _restore_micro_map_assignments_from_export_rows(self, rows):
        """Recover persisted manual micro-trajectory choices before synthetic rows are skipped."""
        if not rows:
            return 0
        valid_micro_shafts = set(getattr(self, "micro_map_shafts", {}) or {})
        if not valid_micro_shafts:
            return 0

        assignments = dict(getattr(self, "micro_map_trajectory_assignments", {}) or {})
        changed = 0
        for row in rows:
            assignment_method = str(row.get("planning_assignment_method") or "").strip().lower()
            if assignment_method != "manual_micro_map_trajectory":
                continue
            shaft_name = row.get("map_shaft") or row.get("shaft")
            if shaft_name not in valid_micro_shafts:
                continue
            trajectory_stub = {
                "trajectory_name": row.get("trajectory_name"),
                "label": row.get("trajectory_label"),
            }
            primary_key = self._trajectory_primary_key(trajectory_stub)
            if not primary_key or self._trajectory_by_primary_key(primary_key) is None:
                continue
            if assignments.get(shaft_name) != primary_key:
                assignments[shaft_name] = primary_key
                changed += 1

        if changed:
            self.micro_map_trajectory_assignments = assignments
            self._prune_micro_map_trajectory_assignments()
        return changed

    def _electrode_info_path(self, ct_path):
        if not ct_path:
            return None
        ct_path = Path(ct_path)
        base_name = ct_path.name
        if base_name.endswith(".nii.gz"):
            base_name = base_name[:-7]
        else:
            base_name = ct_path.stem
        return str(ct_path.parent / f"{base_name}_elec_info.csv")

    def _electrode_info_xlsx_path(self, ct_path):
        if not ct_path:
            return None
        ct_path = Path(ct_path)
        base_name = ct_path.name
        if base_name.endswith(".nii.gz"):
            base_name = base_name[:-7]
        else:
            base_name = ct_path.stem
        return str(ct_path.parent / f"{base_name}_elec_info.xlsx")

    def _electrodes_mat_search_dirs(self):
        """Return likely folders that may contain Electrodes.mat."""
        paths = [
            getattr(self.controller, "patient_reg_folder", None) if self.controller is not None else None,
            self.output_dir,
            self.mri_path,
            self.ct_path,
        ]
        seen = set()
        dirs = []
        for path in paths:
            if not path:
                continue
            try:
                candidate = Path(path).expanduser()
            except Exception:
                continue
            if candidate.suffix:
                candidate = candidate.parent
            for directory in (candidate, candidate.parent):
                try:
                    resolved = str(directory.resolve())
                except Exception:
                    resolved = str(directory)
                if resolved and resolved not in seen:
                    seen.add(resolved)
                    dirs.append(directory)
        return dirs

    @staticmethod
    def _find_named_file_case_insensitive(directory, filename):
        try:
            exact = Path(directory) / filename
            if exact.exists():
                return exact
            target = filename.lower()
            for child in Path(directory).iterdir():
                if child.is_file() and child.name.lower() == target:
                    return child
        except Exception:
            return None
        return None

    def _find_electrodes_mat_path(self):
        """Find a neighboring Electrodes.mat file, if one is available."""
        for directory in self._electrodes_mat_search_dirs():
            match = self._find_named_file_case_insensitive(directory, "Electrodes.mat")
            if match is not None:
                return str(match)
        return None

    @staticmethod
    def _mat_output_key(key):
        text = re.sub(r"[^A-Za-z0-9]+", "_", str(key or "").strip()).strip("_").lower()
        return text or "value"

    @staticmethod
    def _normalize_electrode_match_key(value):
        text = str(value or "").strip().lower()
        if not text:
            return None
        text = re.sub(r"[^a-z0-9]+", "", text)
        match = re.match(r"^([a-z]+)0*([0-9]+)$", text)
        if match:
            return f"{match.group(1)}{int(match.group(2))}"
        return text or None

    @staticmethod
    def _electrode_subshaft_roman_prefix(map_shaft):
        text = re.sub(r"[^A-Za-z0-9]+", "", str(map_shaft or "").strip()).upper()
        match = re.match(r"^([A-Z]+)([12])$", text)
        if not match:
            return None
        roman_suffix = {"1": "I", "2": "II"}.get(match.group(2))
        if not roman_suffix:
            return None
        return f"{match.group(1)}{roman_suffix}"

    def _electrode_row_match_keys(self, row):
        candidates = [
            row.get("name"),
            row.get("map_contact"),
        ]
        label = row.get("trajectory_label")
        map_contact_number = row.get("map_contact_number")
        if label and map_contact_number not in (None, ""):
            candidates.append(f"{label}{map_contact_number}")
        shaft_contact_number = row.get("shaft_contact_number")
        if label and shaft_contact_number not in (None, ""):
            candidates.append(f"{label}{shaft_contact_number}")
        keys = []
        for value in candidates:
            key = self._normalize_electrode_match_key(value)
            if key and key not in keys:
                keys.append(key)
        roman_prefix = self._electrode_subshaft_roman_prefix(row.get("map_shaft"))
        roman_contact_number = self._coerce_positive_int(row.get("map_contact_number"))
        if roman_contact_number is None:
            roman_contact_number = self._coerce_positive_int(row.get("shaft_contact_number"))
        if roman_prefix and roman_contact_number is not None:
            key = self._normalize_electrode_match_key(f"{roman_prefix}{roman_contact_number}")
            if key and key not in keys:
                keys.append(key)
        return keys

    @staticmethod
    def _normalized_trajectory_label_code(label):
        text = str(label or "").strip().upper().replace("*", "")
        text = re.sub(r"[^A-Z0-9]+", "", text)
        return text or None

    @staticmethod
    def _micro_bundle_key(value):
        text = str(value or "").strip()
        if not text:
            return None
        text = re.sub(r"\braw\b.*$", "", text, flags=re.IGNORECASE)
        text = re.sub(r"[^A-Za-z0-9]+", "", text).lower()
        if not text.startswith("m"):
            return None
        text = re.sub(r"ref$", "", text)
        text = re.sub(r"\d+$", "", text)
        if len(text) < 4:
            return None
        return text or None

    @staticmethod
    def _is_micro_like_label(value):
        text = re.sub(r"[^A-Za-z0-9]+", "", str(value or "").strip())
        return bool(re.match(r"^m[lr]", text, flags=re.IGNORECASE))

    @staticmethod
    def _side_code_from_value(value):
        text = str(value or "").strip()
        if not text:
            return None
        compact = re.sub(r"[^A-Za-z0-9]+", "", text)
        if not compact:
            return None
        lower = compact.lower()
        if lower.startswith("ml") or lower.startswith("l"):
            return "l"
        if lower.startswith("mr") or lower.startswith("r"):
            return "r"
        if re.search(r"\bleft\b", text, flags=re.IGNORECASE):
            return "l"
        if re.search(r"\bright\b", text, flags=re.IGNORECASE):
            return "r"
        return None

    @classmethod
    def _micro_target_code_candidates(cls, value):
        text = str(value or "").strip()
        if not text:
            return []

        upper = re.sub(r"[^A-Z0-9]+", " ", text.upper()).strip()
        compact = upper.replace(" ", "")
        codes = []

        def add(code):
            code = re.sub(r"[^A-Z0-9]+", "", str(code or "").upper())
            if code and code not in codes:
                codes.append(code)

        if "AMYGD" in compact:
            add("AMY")
        if "ANGULAR" in compact:
            add("ANG")
        if "FUSIFORM" in compact:
            add("FUS")
        if "PARAHIPPOCAMP" in compact and "GYRUS" in compact:
            add("PHG")
        if "HIPPOCAMP" in compact or "HIPPO" in compact:
            if "HEAD" in compact:
                add("HIPH")
            if "BODY" in compact:
                add("HIPB")
            if "TAIL" in compact:
                add("HIPT")
            add("HIP")
        if all(token in compact for token in ("BASAL", "TEMPORAL", "GYRUS")):
            if "POSTERIOR" in compact:
                add("PBTG")
            elif "ANTERIOR" in compact:
                add("ABTG")
            add("BTG")

        stopwords = {
            "LEFT",
            "RIGHT",
            "MICRO",
            "MACRO",
            "TARGET",
            "TRAJECTORY",
            "SHAFT",
            "ELECTRODE",
            "ELECTRODES",
            "CONTACT",
            "CONTACTS",
        }
        tokens = [
            token for token in re.findall(r"[A-Z0-9]+", upper)
            if token and token not in stopwords
        ]
        if not codes and len(tokens) == 1 and len(tokens[0]) >= 3:
            add(tokens[0][:3])
        elif not codes and len(tokens) >= 2:
            add("".join(token[0] for token in tokens if token))
        return codes

    def _electrode_row_micro_bundle_keys(self, row):
        keys = []
        is_micro_row = any(
            self._is_micro_like_label(value)
            for value in (row.get("name"), row.get("map_contact"), row.get("shaft"))
        )
        if not is_micro_row:
            return keys

        for value in (row.get("name"), row.get("map_contact")):
            key = self._micro_bundle_key(value)
            if key and key not in keys:
                keys.append(key)

        side = None
        for value in (
            row.get("trajectory_hemisphere"),
            row.get("name"),
            row.get("map_contact"),
            row.get("shaft"),
            row.get("trajectory_name"),
        ):
            side = self._side_code_from_value(value)
            if side:
                break
        if side:
            for code in self._micro_target_code_candidates(row.get("trajectory_name")):
                key = self._micro_bundle_key(f"m{side}{code}")
                if key and key not in keys:
                    keys.append(key)
        return keys

    def _electrode_row_shared_mat_keys(self, row):
        keys = []
        is_micro_row = any(
            self._is_micro_like_label(value)
            for value in (row.get("name"), row.get("map_contact"), row.get("shaft"))
        )
        if not is_micro_row:
            return keys

        side = None
        for value in (
            row.get("trajectory_hemisphere"),
            row.get("name"),
            row.get("map_contact"),
            row.get("shaft"),
            row.get("trajectory_name"),
        ):
            side = self._side_code_from_value(value)
            if side:
                break

        label = self._normalized_trajectory_label_code(row.get("trajectory_label"))
        if side and label:
            key = self._normalize_electrode_match_key(f"m{side}{label}")
            if key and key not in keys:
                keys.append(key)
            return keys

        for value in (row.get("name"), row.get("map_contact")):
            bundle_key = self._micro_bundle_key(value)
            if not bundle_key:
                continue
            match = re.match(r"^(m[lr][a-z]\d?)", bundle_key)
            if not match:
                continue
            key = self._normalize_electrode_match_key(match.group(1))
            if key and key not in keys:
                keys.append(key)
        return keys

    def _normalize_micro_shared_key(self, value):
        text = str(value or "").strip()
        if not text:
            return None
        text = re.sub(r"\braw\b.*$", "", text, flags=re.IGNORECASE)
        text = re.sub(r"[^A-Za-z0-9]+", "", text)
        text = re.sub(r"ref$", "", text, flags=re.IGNORECASE)
        if not text:
            return None
        match = re.match(r"^(m[lr][a-z]\d?)", text, flags=re.IGNORECASE)
        if match:
            return self._normalize_electrode_match_key(match.group(1))
        match = re.match(r"^([lr][a-z]\d?)(?:\d+)?$", text, flags=re.IGNORECASE)
        if match:
            return self._normalize_electrode_match_key(f"m{match.group(1)}")
        return None

    def _source_row_micro_bundle_keys(self, row):
        keys = []
        side = None
        for value in (
            row.get("trajectory_hemisphere"),
            row.get("name"),
            row.get("map_contact"),
            row.get("shaft"),
            row.get("trajectory_name"),
        ):
            side = self._side_code_from_value(value)
            if side:
                break
        if not side:
            return keys

        for code in self._micro_target_code_candidates(row.get("trajectory_name")):
            key = self._micro_bundle_key(f"m{side}{code}")
            if key and key not in keys:
                keys.append(key)
        return keys

    def _source_row_shared_micro_key(self, row):
        side = None
        for value in (
            row.get("trajectory_hemisphere"),
            row.get("name"),
            row.get("map_contact"),
            row.get("shaft"),
            row.get("trajectory_name"),
        ):
            side = self._side_code_from_value(value)
            if side:
                break

        label = self._normalized_trajectory_label_code(row.get("trajectory_label"))
        if side and label:
            key = self._normalize_electrode_match_key(f"m{side}{label}")
            if key:
                return key

        for value in (row.get("name"), row.get("map_contact"), row.get("shaft")):
            key = self._normalize_micro_shared_key(value)
            if key:
                return key
        return None

    @staticmethod
    def _micro_target_code(value):
        text = str(value or "").strip()
        if not text:
            return None, None
        text = re.sub(r"\braw\b.*$", "", text, flags=re.IGNORECASE)
        text = re.sub(r"[^A-Za-z0-9]+", "", text)
        text = re.sub(r"ref$", "", text, flags=re.IGNORECASE)
        text = re.sub(r"\d+$", "", text)
        match = re.match(r"^m([lr])([a-z0-9]+)$", text, flags=re.IGNORECASE)
        if not match:
            return None, None
        return match.group(1).lower(), match.group(2).upper()

    @staticmethod
    def _micro_target_name_patterns(target_code):
        code = str(target_code or "").upper()
        pattern_map = {
            "AMY": [("amygd",), ("amy",)],
            "ANG": [("angular",), ("ang",)],
            "FUS": [("fusiform",), ("fus",)],
            "PHG": [("parahipp", "gyr"), ("phg",)],
            "HIPH": [("hippocamp", "head"), ("hippo", "head"), ("hiph",)],
            "HIPB": [("hippocamp", "body"), ("hippo", "body"), ("hipb",)],
            "HIPT": [("hippocamp", "tail"), ("hippo", "tail"), ("hipt",)],
            "HIP": [("hippocamp",), ("hippo",), ("hip",)],
            "PBTG": [("posterior", "basal", "temporal", "gyr"), ("pbtg",)],
            "ABTG": [("anterior", "basal", "temporal", "gyr"), ("abtg",)],
            "BTG": [("basal", "temporal", "gyr"), ("btg",)],
        }
        return pattern_map.get(code, [(code.lower(),)] if code else [])

    def _trajectory_matches_micro_target(self, trajectory, side_code, target_code):
        if trajectory is None or not target_code:
            return False
        if side_code:
            hemisphere = self._trajectory_hemisphere(trajectory)
            expected = "left" if side_code == "l" else "right"
            if hemisphere and hemisphere != expected:
                return False

        candidate_texts = []
        for value in (
            trajectory.get("trajectory_name"),
            trajectory.get("label"),
            self._trajectory_base_name(trajectory),
        ):
            key = self._normalize_planning_name(value)
            if key and key not in candidate_texts:
                candidate_texts.append(key)

        target_key = self._normalize_planning_name(target_code)
        for text in candidate_texts:
            if target_key and target_key in text:
                return True
            for pattern in self._micro_target_name_patterns(target_code):
                if all(token in text for token in pattern):
                    return True
        return False

    def _planning_trajectory_for_micro_target(self, shaft_name, source_row=None):
        trajectories = self.planning_trajectories or self._refresh_planning_trajectories()
        if not trajectories:
            return None

        side_code, target_code = self._micro_target_code(shaft_name)
        if source_row is not None:
            if not side_code:
                side_code = self._side_code_from_value(
                    source_row.get("trajectory_hemisphere")
                    or source_row.get("name")
                    or source_row.get("shaft")
                )
            if not target_code:
                for value in (source_row.get("trajectory_name"), source_row.get("name"), source_row.get("shaft")):
                    side_guess, code_guess = self._micro_target_code(value)
                    if not side_code and side_guess:
                        side_code = side_guess
                    if code_guess:
                        target_code = code_guess
                        break
        if not target_code:
            return None

        matches = [
            trajectory for trajectory in trajectories
            if self._trajectory_matches_micro_target(trajectory, side_code, target_code)
        ]
        if len(matches) == 1:
            return matches[0]
        if len(matches) > 1:
            exact = [
                trajectory for trajectory in matches
                if self._normalize_planning_name(trajectory.get("trajectory_name")) == self._normalize_planning_name(target_code)
            ]
            if len(exact) == 1:
                return exact[0]
            matches.sort(key=lambda trajectory: (
                self._coerce_positive_int(trajectory.get("order")) or 10**9,
                str(trajectory.get("trajectory_name") or ""),
            ))
            return matches[0]
        return None

    @staticmethod
    def _parse_micro_map_contact_label(label):
        text = str(label or "").strip()
        if not text:
            return None, None
        # Strip "ref" tokens wherever they appear (and stray separators) so that
        # labels like "mLAMYref01", "mLAMY01ref", and "mLAMY_01_ref" all yield
        # the same bundle ("mLAMY") with the same contact number. Without this,
        # the regex below would create a separate "mLAMYref" bundle for any
        # label where "ref" appears before the digit pair, doubling the number
        # of synthesized micro rows.
        normalized = re.sub(r"ref", "", text, flags=re.IGNORECASE)
        normalized = re.sub(r"[_\s\-]+", "", normalized)
        match = re.match(r"^(m[LR][A-Za-z0-9]*?)(\d{2})(?:\D.*)?$", normalized)
        if not match:
            match = re.match(r"^(m[LR][A-Za-z0-9]*?)(\d+)", normalized)
        if not match:
            return None, None
        bundle = match.group(1)
        try:
            contact_number = int(match.group(2))
        except ValueError:
            return None, None
        return bundle, contact_number

    @staticmethod
    def _canonical_micro_map_contact_label(bundle, contact_number, raw_label=None):
        try:
            contact_number = int(contact_number)
        except (TypeError, ValueError):
            return None
        suffix = "ref" if re.search(r"ref", str(raw_label or ""), flags=re.IGNORECASE) else ""
        return f"{bundle}{contact_number:02d}{suffix}"

    @staticmethod
    def _micro_map_address_rank(address):
        text = str(address or "").strip().upper()
        if text.startswith("1.C"):
            return 0
        if text.startswith("1.D"):
            return 1
        return 2

    def _read_micro_map_shafts(self, path):
        shafts = {}
        row_order = 0
        try:
            with open(path, "r", encoding="utf-8", errors="replace") as handle:
                for line_number, raw_line in enumerate(handle, start=1):
                    line = raw_line.strip()
                    if not line or line.startswith("#"):
                        continue
                    fields = [field.strip() for field in line.split(";")]
                    if len(fields) < 2:
                        continue
                    address = fields[0]
                    if not re.match(r"^1\.[CD](?:\.|$)", address, flags=re.IGNORECASE):
                        continue
                    contact_label = fields[1]
                    bundle, contact_number = self._parse_micro_map_contact_label(contact_label)
                    if bundle is None or contact_number is None:
                        continue

                    entry = shafts.setdefault(
                        bundle,
                        {
                            "shaft": bundle,
                            "contacts": [],
                            "order": row_order,
                        },
                    )
                    if not entry["contacts"]:
                        row_order += 1
                    canonical_label = self._canonical_micro_map_contact_label(
                        bundle,
                        contact_number,
                        contact_label,
                    )
                    entry["contacts"].append({
                        "address": address,
                        "label": canonical_label or contact_label,
                        "raw_label": contact_label,
                        "number": contact_number,
                        "line": line_number,
                        "address_rank": self._micro_map_address_rank(address),
                    })
        except OSError:
            return {}

        for entry in shafts.values():
            contacts = entry["contacts"]
            contacts.sort(key=lambda item: (item["address_rank"], item["line"], item["number"]))
            entry["contact_count"] = len(contacts)
            entry["first_contact"] = contacts[0]["label"] if contacts else None
        return shafts

    def _synthetic_micro_rows_from_map(self, rows, map_path):
        if not map_path:
            return [], []

        micro_shafts = self._read_micro_map_shafts(map_path)
        if not micro_shafts:
            return [], []

        def row_sort_key(row):
            shaft_contact = self._coerce_positive_int(row.get("shaft_contact_number"))
            map_contact = self._coerce_positive_int(row.get("map_contact_number"))
            index = self._coerce_positive_int(row.get("index"))
            return (
                shaft_contact if shaft_contact is not None else 10**9,
                map_contact if map_contact is not None else 10**9,
                index if index is not None else 10**9,
                str(row.get("name") or ""),
            )

        existing_names = {
            self._normalize_electrode_match_key(row.get("name"))
            for row in rows
            if row.get("name")
        }
        source_by_bundle_key = {}
        source_by_shared_key = {}

        def consider(mapping, key, row):
            if not key:
                return
            current = mapping.get(key)
            if current is None or row_sort_key(row) < row_sort_key(current):
                mapping[key] = row

        for row in rows:
            if any(
                self._is_micro_like_label(value)
                for value in (row.get("name"), row.get("map_contact"), row.get("shaft"))
            ):
                continue
            if any(row.get(axis) in (None, "") for axis in ("x", "y", "z")):
                continue
            for key in self._source_row_micro_bundle_keys(row):
                consider(source_by_bundle_key, key, row)
            consider(source_by_shared_key, self._source_row_shared_micro_key(row), row)

        def source_row_for_trajectory(trajectory):
            if trajectory is None:
                return None
            side = self._side_code_from_value(
                self._trajectory_hemisphere(trajectory)
                or trajectory.get("trajectory_name")
                or trajectory.get("label")
            )
            if side:
                for code in self._micro_target_code_candidates(trajectory.get("trajectory_name")):
                    key = self._micro_bundle_key(f"m{side}{code}")
                    source_row = source_by_bundle_key.get(key)
                    if source_row is not None:
                        return source_row
                label = self._normalized_trajectory_label_code(trajectory.get("label"))
                if label:
                    source_row = source_by_shared_key.get(
                        self._normalize_electrode_match_key(f"m{side}{label}")
                    )
                    if source_row is not None:
                        return source_row
            return None

        next_index = max(
            [self._coerce_positive_int(row.get("index")) or -1 for row in rows] or [-1]
        ) + 1
        synthetic_rows = []
        unresolved_shafts = []

        for shaft_name in sorted(micro_shafts, key=lambda name: micro_shafts[name].get("order", 0)):
            shaft_entry = micro_shafts[shaft_name]
            bundle_key = self._micro_bundle_key(shaft_name)
            manual_key = getattr(self, "micro_map_trajectory_assignments", {}).get(shaft_name)
            target_trajectory = self._trajectory_by_primary_key(manual_key) if manual_key else None
            assignment_method = "manual_micro_map_trajectory" if target_trajectory is not None else None
            if target_trajectory is None:
                target_trajectory = self._planning_trajectory_for_micro_target(shaft_name)
                assignment_method = "synthetic_micro_target" if target_trajectory is not None else None
            source_row = (
                source_row_for_trajectory(target_trajectory)
                or source_by_bundle_key.get(bundle_key)
            )
            if source_row is None:
                unresolved_shafts.append(shaft_name)

            for contact in shaft_entry.get("contacts", []):
                contact_label = contact.get("label")
                name_key = self._normalize_electrode_match_key(contact_label)
                if not contact_label or (name_key and name_key in existing_names):
                    continue

                new_row = dict(source_row) if source_row is not None else {}
                new_row.update({
                    "index": next_index,
                    "source_index": source_row.get("source_index") if source_row is not None else None,
                    "name": contact_label,
                    "shaft": shaft_name,
                    "shaft_contact_number": contact.get("number"),
                    "detected": False,
                    "map_shaft": shaft_name,
                    "map_contact": contact_label,
                    "map_contact_number": contact.get("number"),
                    "electrode_map": map_path,
                    "synthetic_3d_method": "macro_first_contact",
                    "synthetic_3d_source_contact": source_row.get("name") if source_row is not None else None,
                    "synthetic_3d_source_shaft": source_row.get("shaft") if source_row is not None else None,
                })
                if source_row is None:
                    new_row["synthetic_3d_method"] = "unresolved_macro_source"
                if target_trajectory is not None:
                    new_row.update({
                        "trajectory_name": target_trajectory.get("trajectory_name"),
                        "trajectory_label": target_trajectory.get("label"),
                        "planned_contacts": target_trajectory.get("contacts"),
                        "trajectory_order": target_trajectory.get("order"),
                        "trajectory_hemisphere": self._trajectory_hemisphere(target_trajectory),
                        "planning_worksheet": (
                            self.planning_worksheet_path
                            or (source_row.get("planning_worksheet") if source_row is not None else None)
                        ),
                        "planning_assignment_method": assignment_method,
                        "planning_assignment_cost": None,
                    })
                synthetic_rows.append({
                    key: self._export_value(value)
                    for key, value in new_row.items()
                })
                next_index += 1

        return synthetic_rows, unresolved_shafts

    @staticmethod
    def _mat_record_text_values(record):
        values = []
        for key, value in (record or {}).items():
            lower = str(key).lower()
            if lower.endswith("_source") or lower == "mat_label":
                continue
            if not isinstance(value, str):
                continue
            if any(token in lower for token in ("target", "trajectory", "implant", "anatom", "region")):
                text = value.strip()
                if text and text not in values:
                    values.append(text)
        return values

    def _mat_record_micro_bundle_keys(self, record):
        keys = []
        label_key = self._normalize_electrode_match_key(record.get("mat_label"))
        if not label_key:
            return keys

        side_match = re.match(r"^m([lr])", label_key)
        if not side_match:
            return keys
        side = side_match.group(1)

        for value in self._mat_record_text_values(record):
            for code in self._micro_target_code_candidates(value):
                key = self._micro_bundle_key(f"m{side}{code}")
                if key and key not in keys:
                    keys.append(key)
        return keys

    def _mat_record_shared_mat_keys(self, record):
        keys = []
        label_key = self._normalize_electrode_match_key(record.get("mat_label"))
        if label_key and self._is_micro_like_label(record.get("mat_label")):
            keys.append(label_key)
        return keys

    def _record_matches_preferred_shared_key(self, record, preferred_key, row_bundle_keys=None):
        if not preferred_key:
            shared_key_ok = True
        else:
            shared_key_ok = preferred_key in self._mat_record_shared_mat_keys(record)
        if not shared_key_ok:
            return False

        if row_bundle_keys:
            record_bundle_keys = self._mat_record_micro_bundle_keys(record)
            if record_bundle_keys:
                row_bundle_key_set = set(row_bundle_keys)
                return any(key in row_bundle_key_set for key in record_bundle_keys)
        return True

    @staticmethod
    def _text_from_mat_scalar(value):
        if value is None:
            return None
        if isinstance(value, bytes):
            return value.decode("utf-8", errors="ignore").strip() or None
        if isinstance(value, str):
            return value.strip() or None
        if isinstance(value, np.generic):
            value = value.item()
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            return None

        try:
            arr = np.asarray(value)
        except Exception:
            return str(value).strip() or None
        if arr.size == 0:
            return None
        if arr.dtype.kind == "S":
            if arr.ndim == 0:
                return bytes(arr.item()).decode("utf-8", errors="ignore").strip() or None
            parts = []
            for item in arr.reshape(-1):
                text = bytes(item).decode("utf-8", errors="ignore")
                if text:
                    parts.append(text)
            return "".join(parts).strip() or None
        if arr.dtype.kind == "U":
            if arr.ndim == 0:
                return str(arr.item()).strip() or None
            parts = [str(item) for item in arr.reshape(-1)]
            if all(len(part) <= 1 for part in parts):
                return "".join(parts).strip() or None
            return " ".join(part.strip() for part in parts if part.strip()) or None
        if arr.dtype.kind in "ui" and arr.ndim <= 2 and arr.size <= 512:
            try:
                flat = [int(v) for v in arr.reshape(-1) if int(v) > 0]
                if flat and max(flat) <= 65535:
                    return "".join(chr(v) for v in flat).strip() or None
            except Exception:
                return None
        return None

    @classmethod
    def _mat_value_to_text_list(cls, key, value):
        """Convert common MATLAB string/cell encodings to a Python text list."""
        key_lower = str(key or "").lower()
        key_looks_text = any(token in key_lower for token in ("name", "label", "chan", "contact", "electrode"))

        if isinstance(value, (list, tuple)):
            texts = []
            for item in value:
                if isinstance(item, (list, tuple)):
                    nested = cls._mat_value_to_text_list(key, item)
                    if nested:
                        texts.extend(nested)
                    continue
                text = cls._text_from_mat_scalar(item)
                if text is not None:
                    texts.append(text)
            return texts or None

        try:
            arr = np.asarray(value)
        except Exception:
            text = cls._text_from_mat_scalar(value)
            return [text] if text else None

        if arr.size == 0:
            return None
        if arr.dtype == object:
            texts = []
            for item in arr.reshape(-1):
                nested = cls._mat_value_to_text_list(key, item)
                if nested:
                    if len(nested) == 1:
                        texts.append(nested[0])
                    else:
                        texts.extend(nested)
            return texts or None
        if arr.dtype.kind in "SU":
            if arr.ndim == 0:
                text = cls._text_from_mat_scalar(arr.item())
                return [text] if text else None
            if arr.ndim == 1:
                parts = [cls._text_from_mat_scalar(item) or "" for item in arr]
                if parts and all(len(part) <= 1 for part in parts):
                    text = "".join(parts).strip()
                    return [text] if text else None
                return [part for part in parts if part]
            if arr.ndim == 2:
                row_texts = []
                for row in arr:
                    text = cls._text_from_mat_scalar(row)
                    if text:
                        row_texts.append(text)
                if row_texts:
                    return row_texts
        if key_looks_text and arr.dtype.kind in "ui" and arr.ndim <= 2:
            if arr.ndim == 1:
                text = cls._text_from_mat_scalar(arr)
                return [text] if text else None
            row_texts = []
            for row in arr:
                text = cls._text_from_mat_scalar(row)
                if text:
                    row_texts.append(text)
            return row_texts or None
        return None

    @staticmethod
    def _mat_coordinate_array(key, value):
        key_lower = str(key or "").lower()
        if not any(token in key_lower for token in ("xyz", "coord", "pos", "ras", "mni", "voxel", "vox")):
            return None
        try:
            arr = np.asarray(value)
        except Exception:
            return None
        if arr.dtype == object or arr.dtype.kind not in "fiu":
            return None
        arr = np.asarray(arr, dtype=float)
        arr = np.squeeze(arr)
        if arr.ndim != 2:
            return None
        if arr.shape[1] == 3:
            coords = arr
        elif arr.shape[0] == 3:
            coords = arr.T
        else:
            return None
        if coords.size == 0:
            return None
        return coords

    @staticmethod
    def _mat_column_values(value, n_rows):
        try:
            arr = np.asarray(value)
        except Exception:
            return None
        if arr.dtype == object or arr.size == 0:
            return None
        arr = np.squeeze(arr)
        if arr.ndim == 0:
            return None
        if arr.ndim == 1 and arr.shape[0] == n_rows:
            return arr.tolist()
        if arr.ndim == 2 and 1 in arr.shape and max(arr.shape) == n_rows:
            return arr.reshape(-1).tolist()
        return None

    @staticmethod
    def _mat_coordinate_kind(key):
        lower = str(key or "").lower()
        if "mni" in lower or "tal" in lower:
            return "mni"
        return "native"

    @staticmethod
    def _mat_coordinate_preference(key):
        lower = str(key or "").lower()
        preference = 50
        preferred = [
            ("elecxyzprojraw", 0),
            ("elecxyzraw", 1),
            ("elecxyzproj", 2),
            ("elecxyz", 3),
            ("mni", 4),
            ("ras", 5),
            ("coord", 6),
            ("pos", 7),
            ("vox", 8),
        ]
        compact = re.sub(r"[^a-z0-9]+", "", lower)
        for token, score in preferred:
            if token in compact:
                preference = min(preference, score)
        return preference

    @staticmethod
    def _mat_label_preference(key):
        lower = str(key or "").lower()
        compact = re.sub(r"[^a-z0-9]+", "", lower)
        if "electrodename" in compact or "elecname" in compact:
            return 0
        if "channame" in compact or "channelname" in compact:
            return 1
        if "name" in compact:
            return 2
        if "contact" in compact:
            return 3
        if "label" in compact:
            return 4
        return 20

    def _collect_scipy_mat_values(self, prefix, value, values):
        if prefix.startswith("__"):
            return
        if isinstance(value, dict):
            for key, item in value.items():
                child = f"{prefix}.{key}" if prefix else str(key)
                self._collect_scipy_mat_values(child, item, values)
            return
        if hasattr(value, "_fieldnames"):
            for key in getattr(value, "_fieldnames", []) or []:
                child = f"{prefix}.{key}" if prefix else str(key)
                self._collect_scipy_mat_values(child, getattr(value, key), values)
            return
        if isinstance(value, np.ndarray) and value.dtype.names:
            for key in value.dtype.names:
                child = f"{prefix}.{key}" if prefix else str(key)
                self._collect_scipy_mat_values(child, value[key], values)
            return
        values[prefix] = value

    def _load_scipy_mat_values(self, mat_path):
        from scipy.io import loadmat

        raw = loadmat(mat_path, squeeze_me=True, struct_as_record=False)
        values = {}
        for key, value in raw.items():
            self._collect_scipy_mat_values(str(key), value, values)
        return values

    def _load_hdf5_mat_values(self, mat_path):
        import h5py

        values = {}

        def decode_value(handle, value):
            if isinstance(value, h5py.Dataset):
                return decode_value(handle, value[()])
            if isinstance(value, h5py.Reference):
                if not value:
                    return None
                return decode_value(handle, handle[value])
            if isinstance(value, bytes):
                return value.decode("utf-8", errors="ignore")
            try:
                arr = np.asarray(value)
            except Exception:
                return value
            if h5py.check_dtype(ref=arr.dtype) is not None:
                decoded = []
                for ref in arr.reshape(-1):
                    if ref:
                        decoded.append(decode_value(handle, ref))
                return decoded
            return value

        with h5py.File(mat_path, "r") as handle:
            def visit(name, obj):
                if name.startswith("#refs#"):
                    return
                if isinstance(obj, h5py.Dataset):
                    try:
                        values[name] = decode_value(handle, obj)
                    except Exception:
                        pass
            handle.visititems(visit)
        return values

    def _load_mat_values(self, mat_path):
        try:
            return self._load_scipy_mat_values(mat_path)
        except NotImplementedError:
            return self._load_hdf5_mat_values(mat_path)
        except ValueError as exc:
            if "Unknown mat file type" in str(exc) or "Please use HDF reader" in str(exc):
                return self._load_hdf5_mat_values(mat_path)
            raise

    def _mat_records_from_values(self, values):
        coordinate_candidates = []
        text_candidates = []
        for key, value in (values or {}).items():
            coords = self._mat_coordinate_array(key, value)
            if coords is not None:
                coordinate_candidates.append((
                    self._mat_coordinate_kind(key),
                    self._mat_coordinate_preference(key),
                    key,
                    coords,
                ))
                continue
            texts = self._mat_value_to_text_list(key, value)
            if texts and len(texts) > 1:
                text_candidates.append((self._mat_label_preference(key), key, texts))

        if not coordinate_candidates:
            return [], {}

        native_candidates = sorted(
            [item for item in coordinate_candidates if item[0] == "native"],
            key=lambda item: (item[1], item[2]),
        )
        mni_candidates = sorted(
            [item for item in coordinate_candidates if item[0] == "mni"],
            key=lambda item: (item[1], item[2]),
        )
        native_key, native_coords = (None, None)
        mni_key, mni_coords = (None, None)
        if native_candidates:
            _, _, native_key, native_coords = native_candidates[0]
        if mni_candidates:
            _, _, mni_key, mni_coords = mni_candidates[0]

        n_rows = 0
        for coords in (native_coords, mni_coords):
            if coords is not None:
                n_rows = max(n_rows, int(coords.shape[0]))
        if n_rows <= 0:
            return [], {}

        label_key = None
        labels = None
        matching_texts = [
            (score, key, texts)
            for score, key, texts in text_candidates
            if len(texts) == n_rows
        ]
        if matching_texts:
            _, label_key, labels = sorted(matching_texts, key=lambda item: (item[0], item[1]))[0]

        extra_columns = {}
        coordinate_keys = {native_key, mni_key}
        for key, value in (values or {}).items():
            if key in coordinate_keys:
                continue
            output_key = f"mat_{self._mat_output_key(key)}"
            if output_key in {"mat_index", "mat_label"}:
                continue
            texts = self._mat_value_to_text_list(key, value)
            if texts and len(texts) == n_rows:
                extra_columns[output_key] = texts
                continue
            column = self._mat_column_values(value, n_rows)
            if column is not None:
                extra_columns[output_key] = column

        records = []
        for idx in range(n_rows):
            record = {
                "mat_index": idx + 1,
                "mat_label": labels[idx] if labels and idx < len(labels) else None,
                "mat_label_source": label_key,
                "mat_xyz_source": native_key,
                "mat_mni_source": mni_key,
            }
            if native_coords is not None and idx < native_coords.shape[0]:
                record["mat_x"] = float(native_coords[idx, 0])
                record["mat_y"] = float(native_coords[idx, 1])
                record["mat_z"] = float(native_coords[idx, 2])
            if mni_coords is not None and idx < mni_coords.shape[0]:
                record["mat_mni_x"] = float(mni_coords[idx, 0])
                record["mat_mni_y"] = float(mni_coords[idx, 1])
                record["mat_mni_z"] = float(mni_coords[idx, 2])
            for key, values_for_key in extra_columns.items():
                if idx < len(values_for_key):
                    record[key] = self._export_value(values_for_key[idx])
            records.append(record)

        summary = {
            "native_source": native_key,
            "mni_source": mni_key,
            "label_source": label_key,
            "record_count": len(records),
        }
        return records, summary

    def _load_electrodes_mat_records(self, mat_path):
        values = self._load_mat_values(mat_path)
        return self._mat_records_from_values(values)

    @staticmethod
    def _coordinate_distance(row, row_keys, mat_record, mat_keys):
        try:
            row_coords = [float(row.get(key)) for key in row_keys]
            mat_coords = [float(mat_record.get(key)) for key in mat_keys]
        except (TypeError, ValueError):
            return None
        if any(not np.isfinite(value) for value in row_coords + mat_coords):
            return None
        return float(np.linalg.norm(np.asarray(row_coords) - np.asarray(mat_coords)))

    def _merge_electrodes_mat_info(self, rows, mat_records, mat_path):
        """Merge Electrodes.mat rows into exported electrode rows and add comparison status."""
        if not mat_records:
            for row in rows:
                row["mat_source"] = mat_path or ""
                row["mat_match_status"] = "no_mat_file" if not mat_path else "no_mat_records"
            return rows, {"matched": 0, "close": 0, "off": 0, "review": 0, "unmatched": len(rows)}

        records_by_key = {}
        duplicate_keys = set()
        records_by_micro_bundle = {}
        duplicate_micro_bundle_keys = set()
        records_by_shared_mat_key = {}
        duplicate_shared_mat_keys = set()
        for record in mat_records:
            candidate_values = [record.get("mat_label")]
            for key, value in record.items():
                lower = str(key).lower()
                if lower.endswith("_source"):
                    continue
                if any(token in lower for token in ("name", "label", "chan", "contact")):
                    candidate_values.append(value)
            for value in candidate_values:
                key = self._normalize_electrode_match_key(value)
                if not key:
                    continue
                if key in records_by_key and records_by_key[key] is not record:
                    duplicate_keys.add(key)
                else:
                    records_by_key[key] = record
            for key in self._mat_record_micro_bundle_keys(record):
                if key in records_by_micro_bundle and records_by_micro_bundle[key] is not record:
                    duplicate_micro_bundle_keys.add(key)
                else:
                    records_by_micro_bundle[key] = record
            for key in self._mat_record_shared_mat_keys(record):
                if key in records_by_shared_mat_key and records_by_shared_mat_key[key] is not record:
                    duplicate_shared_mat_keys.add(key)
                else:
                    records_by_shared_mat_key[key] = record
        for key in duplicate_keys:
            records_by_key.pop(key, None)
        for key in duplicate_micro_bundle_keys:
            records_by_micro_bundle.pop(key, None)
        for key in duplicate_shared_mat_keys:
            records_by_shared_mat_key.pop(key, None)

        used_records = set()
        counts = {"matched": 0, "close": 0, "off": 0, "review": 0, "unmatched": 0}
        for row_index, row in enumerate(rows):
            match = None
            shared_match = False
            preferred_bundle_keys = self._electrode_row_micro_bundle_keys(row)
            preferred_shared_keys = self._electrode_row_shared_mat_keys(row)
            preferred_shared_key = preferred_shared_keys[0] if preferred_shared_keys else None
            for key in self._electrode_row_match_keys(row):
                candidate = records_by_key.get(key)
                if candidate is not None and id(candidate) not in used_records:
                    match = candidate
                    break
            if match is None:
                for key in preferred_bundle_keys:
                    candidate = records_by_micro_bundle.get(key)
                    if candidate is not None and self._record_matches_preferred_shared_key(
                        candidate,
                        preferred_shared_key,
                        preferred_bundle_keys,
                    ):
                        match = candidate
                        shared_match = True
                        break
            if match is None:
                for key in preferred_shared_keys:
                    candidate = records_by_shared_mat_key.get(key)
                    if candidate is not None and self._record_matches_preferred_shared_key(
                        candidate,
                        preferred_shared_key,
                        preferred_bundle_keys,
                    ):
                        match = candidate
                        shared_match = True
                        break
            # Only use row-order matching when Electrodes.mat has no usable labels;
            # labeled files are often not stored in export order.
            if (
                match is None
                and not records_by_key
                and row_index < len(mat_records)
                and id(mat_records[row_index]) not in used_records
            ):
                match = mat_records[row_index]

            row["mat_source"] = mat_path or ""
            if match is None:
                row["mat_match_status"] = "no_match"
                counts["unmatched"] += 1
                continue

            if not shared_match:
                used_records.add(id(match))
            counts["matched"] += 1
            for key, value in match.items():
                if key not in row:
                    row[key] = self._export_value(value)

            xyz_distance = self._coordinate_distance(
                row, ("x", "y", "z"), match, ("mat_x", "mat_y", "mat_z")
            )
            mri_world_distance = self._coordinate_distance(
                row,
                ("mri_world_x", "mri_world_y", "mri_world_z"),
                match,
                ("mat_x", "mat_y", "mat_z"),
            )
            mni_distance = self._coordinate_distance(
                row,
                ("mni_x", "mni_y", "mni_z"),
                match,
                ("mat_mni_x", "mat_mni_y", "mat_mni_z"),
            )
            native_distances = [
                value for value in (xyz_distance, mri_world_distance)
                if value is not None
            ]
            native_best = min(native_distances) if native_distances else None
            comparison_distances = [
                value for value in (native_best, mni_distance)
                if value is not None
            ]

            row["mat_xyz_distance"] = xyz_distance
            row["mat_mri_world_distance"] = mri_world_distance
            row["mat_native_best_distance"] = native_best
            row["mat_mni_distance"] = mni_distance

            if not comparison_distances:
                status = "matched_no_comparison"
            elif any(value >= ELECTRODE_MAT_OFF_DISTANCE for value in comparison_distances):
                status = "off"
            elif all(value <= ELECTRODE_MAT_CLOSE_DISTANCE for value in comparison_distances):
                status = "close"
            else:
                status = "review"
            row["mat_match_status"] = status
            if status in counts:
                counts[status] += 1
        return rows, counts

    @staticmethod
    def _row_fieldnames(rows):
        fieldnames = []
        for row in rows:
            for key in row.keys():
                if key not in fieldnames:
                    fieldnames.append(key)
        return fieldnames

    def _workbook_value(self, value):
        value = self._export_value(value)
        if isinstance(value, float) and not np.isfinite(value):
            return None
        if isinstance(value, (list, tuple, dict)):
            return str(value)
        return value

    def _save_electrode_workbook(self, xlsx_path, rows, mat_records=None):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Electrodes"

        fieldnames = self._row_fieldnames(rows)
        header_fill = PatternFill("solid", fgColor="D9EAF7")
        close_fill = PatternFill("solid", fgColor="C6EFCE")
        off_fill = PatternFill("solid", fgColor="FFC7CE")
        review_fill = PatternFill("solid", fgColor="FFEB9C")

        for col_idx, name in enumerate(fieldnames, start=1):
            cell = sheet.cell(row=1, column=col_idx, value=name)
            cell.font = Font(bold=True)
            cell.fill = header_fill

        for row_idx, row in enumerate(rows, start=2):
            status = row.get("mat_match_status")
            fill = None
            if status == "close":
                fill = close_fill
            elif status == "off":
                fill = off_fill
            elif status == "review":
                fill = review_fill
            for col_idx, name in enumerate(fieldnames, start=1):
                cell = sheet.cell(row=row_idx, column=col_idx, value=self._workbook_value(row.get(name)))
                if fill is not None:
                    cell.fill = fill

        if fieldnames:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for col_idx, name in enumerate(fieldnames, start=1):
                max_len = len(str(name))
                for row in rows[:200]:
                    value = row.get(name)
                    if value is not None:
                        max_len = max(max_len, len(str(value)))
                sheet.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 32)

        if mat_records:
            mat_sheet = workbook.create_sheet("Electrodes.mat")
            mat_fieldnames = self._row_fieldnames(mat_records)
            for col_idx, name in enumerate(mat_fieldnames, start=1):
                cell = mat_sheet.cell(row=1, column=col_idx, value=name)
                cell.font = Font(bold=True)
                cell.fill = header_fill
            for row_idx, record in enumerate(mat_records, start=2):
                for col_idx, name in enumerate(mat_fieldnames, start=1):
                    mat_sheet.cell(row=row_idx, column=col_idx, value=self._workbook_value(record.get(name)))
            if mat_fieldnames:
                mat_sheet.freeze_panes = "A2"
                mat_sheet.auto_filter.ref = mat_sheet.dimensions
                for col_idx, name in enumerate(mat_fieldnames, start=1):
                    max_len = len(str(name))
                    for record in mat_records[:200]:
                        value = record.get(name)
                        if value is not None:
                            max_len = max(max_len, len(str(value)))
                    mat_sheet.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 32)

        workbook.save(xlsx_path)

    def _transform_path(self, ct_path):
        if not ct_path:
            return None
        ct_path = Path(ct_path)
        base_name = ct_path.name
        if base_name.endswith(".nii.gz"):
            base_name = base_name[:-7]
        else:
            base_name = ct_path.stem
        return str(ct_path.parent / f"{base_name}_transform.npy")

    def _load_transform_matrix(self, ct_path):
        """Load a saved 4x4 transform matrix if present."""
        base_path = self._transform_path(ct_path)
        if not base_path:
            return None
        candidates = [base_path]
        txt_path = str(Path(base_path).with_suffix(".txt"))
        if txt_path not in candidates:
            candidates.append(txt_path)

        for path in candidates:
            if not os.path.exists(path):
                continue
            try:
                if path.lower().endswith(".npy"):
                    matrix = np.load(path)
                else:
                    matrix = np.loadtxt(path)
            except Exception as exc:
                self.log(f"Failed to load transform from {path}: {exc}")
                continue
            matrix = np.asarray(matrix)
            if matrix.shape != (4, 4):
                self.log(f"Transform matrix must be 4x4; got {matrix.shape} in {path}")
                continue
            return matrix
        return None

    def _electrode_mri_world_coords(self, electrode):
        """Return best available native MRI-world coordinates for an electrode."""
        coords = [electrode.get("x"), electrode.get("y"), electrode.get("z")]
        try:
            vox = np.array([float(coords[0]), float(coords[1]), float(coords[2]), 1.0], dtype=float)
        except (TypeError, ValueError):
            vox = None

        transforms = []
        if self.controller is not None and hasattr(self.controller, "_get_loaded_mri_voxel_to_subject_mri_trans"):
            try:
                transforms.append(self.controller._get_loaded_mri_voxel_to_subject_mri_trans())
            except Exception:
                pass
        if self.before_mri_img is not None:
            try:
                transforms.append(np.asarray(self.before_mri_img.affine, dtype=float))
            except Exception:
                pass
        if self.mri_path and os.path.exists(self.mri_path):
            try:
                transforms.append(np.asarray(nib.load(self.mri_path).affine, dtype=float))
            except Exception:
                pass

        if vox is not None:
            for transform in transforms:
                if transform is None:
                    continue
                try:
                    transform = np.asarray(transform, dtype=float)
                    if transform.shape != (4, 4):
                        continue
                    world = (vox @ transform.T)[:3]
                    return (float(world[0]), float(world[1]), float(world[2]))
                except Exception:
                    continue

        stored = [
            electrode.get("mri_world_x"),
            electrode.get("mri_world_y"),
            electrode.get("mri_world_z"),
        ]
        try:
            if all(value is not None and np.isfinite(float(value)) for value in stored):
                return tuple(float(value) for value in stored)
        except (TypeError, ValueError):
            pass
        return (None, None, None)

    def _electrode_parcellation_name(self, electrode):
        """Return the FreeSurfer parcellation name for an electrode, if available."""
        vol = getattr(self, "before_parcellation_volume", None)
        if vol is None or getattr(vol, "ndim", None) != 3:
            return electrode.get("parcellation")
        label_id, label_name = self._parcellation_for_electrode(electrode)
        if label_id is None:
            return None
        electrode["parcellation"] = label_name
        return label_name

    def _refresh_all_shaft_contact_numbers_for_export(self):
        """Normalize shaft contact order and labels before saving."""
        electrodes = getattr(self, "before_electrodes", []) or []
        shafts = getattr(self, "before_shafts", {}) or {}
        if not electrodes or not shafts:
            return {}

        ordered_shafts = {}
        contact_numbers = {}
        for shaft_name in self._sorted_shaft_names(shafts, include_unassigned=True):
            raw_indices = []
            for electrode_idx in shafts.get(shaft_name, []):
                try:
                    electrode_idx = int(electrode_idx)
                except (TypeError, ValueError):
                    continue
                if 0 <= electrode_idx < len(electrodes):
                    raw_indices.append(electrode_idx)

            if shaft_name == "Unassigned":
                indices = raw_indices
            else:
                indices = self._order_electrode_indices_distal_first(electrodes, raw_indices)
                for contact_number, electrode_idx in enumerate(indices, start=1):
                    electrodes[electrode_idx]["shaft_contact_number"] = contact_number
                    contact_numbers[electrode_idx] = contact_number
                self._refresh_contact_numbers_for_shaft(shaft_name, indices)

            ordered_shafts[shaft_name] = list(indices)

        self.before_shafts = ordered_shafts
        return contact_numbers

    def _electrode_export_indices(self):
        """Return electrode indices in shaft/contact order for stable exports."""
        electrodes = getattr(self, "before_electrodes", []) or []
        shafts = getattr(self, "before_shafts", {}) or {}
        if not electrodes:
            return []
        if not shafts:
            return list(range(len(electrodes)))

        ordered = []
        used = set()
        for shaft_name in self._sorted_shaft_names(shafts, include_unassigned=True):
            for electrode_idx in shafts.get(shaft_name, []):
                try:
                    electrode_idx = int(electrode_idx)
                except (TypeError, ValueError):
                    continue
                if 0 <= electrode_idx < len(electrodes) and electrode_idx not in used:
                    ordered.append(electrode_idx)
                    used.add(electrode_idx)

        for electrode_idx in range(len(electrodes)):
            if electrode_idx not in used:
                ordered.append(electrode_idx)
        return ordered

    def _electrode_export_row(self, electrode, index, source_index=None, shaft_contact_number=None):
        name = electrode.get("name") or f"Electrode {index + 1}"
        native_space = electrode.get("space") or self._current_electrode_space()
        mri_world_x, mri_world_y, mri_world_z = self._electrode_mri_world_coords(electrode)
        if all(value is not None for value in (mri_world_x, mri_world_y, mri_world_z)):
            electrode["mri_world_x"] = mri_world_x
            electrode["mri_world_y"] = mri_world_y
            electrode["mri_world_z"] = mri_world_z
        parcellation = self._electrode_parcellation_name(electrode)
        row = {
            "index": index,
            "source_index": source_index,
            "name": name,
            "shaft": electrode.get("shaft", ""),
            "shaft_contact_number": shaft_contact_number,
            "x": electrode.get("x"),
            "y": electrode.get("y"),
            "z": electrode.get("z"),
            "space": native_space,
            "display_space": self._current_display_electrode_space(),
            "tissue": electrode.get("tissue", "unknown"),
            "parcellation": parcellation,
            "detected": bool(electrode.get("detected", False)),
            "visible": bool(electrode.get("visible", True)),
            "view": electrode.get("view"),
            "slice_axial": electrode.get("slice_axial"),
            "slice_coronal": electrode.get("slice_coronal"),
            "slice_sagittal": electrode.get("slice_sagittal"),
            "mri_world_x": mri_world_x,
            "mri_world_y": mri_world_y,
            "mri_world_z": mri_world_z,
            "mni_x": electrode.get("mni_x"),
            "mni_y": electrode.get("mni_y"),
            "mni_z": electrode.get("mni_z"),
            "mni_space": electrode.get("mni_space"),
            "trajectory_name": electrode.get("trajectory_name"),
            "trajectory_label": electrode.get("trajectory_label"),
            "planned_contacts": electrode.get("planned_contacts"),
            "trajectory_order": electrode.get("trajectory_order"),
            "trajectory_hemisphere": electrode.get("trajectory_hemisphere"),
            "planning_assignment_method": electrode.get("planning_assignment_method"),
            "planning_assignment_cost": electrode.get("planning_assignment_cost"),
            "trajectory_schema_ap": electrode.get("trajectory_schema_ap"),
            "trajectory_schema_si": electrode.get("trajectory_schema_si"),
            "shaft_entry_ap": electrode.get("shaft_entry_ap"),
            "shaft_entry_si": electrode.get("shaft_entry_si"),
            "planning_worksheet": electrode.get("planning_worksheet"),
            "map_shaft": electrode.get("map_shaft"),
            "map_contact": electrode.get("map_contact"),
            "map_contact_number": electrode.get("map_contact_number"),
            "electrode_map": electrode.get("electrode_map"),
        }
        return {key: self._export_value(value) for key, value in row.items()}

    def _save_electrode_locations(self):
        """Save electrode locations and compare against Electrodes.mat when available."""
        electrodes = self.before_electrodes
        if not electrodes:
            QMessageBox.information(self, "No Electrodes", "No electrode locations to save.")
            return
        if not self.ct_path:
            QMessageBox.warning(
                self,
                "Missing CT Path",
                "CT file path is not set. Load a CT file before saving electrodes.",
            )
            return
        fname = self._electrode_info_path(self.ct_path)
        if fname is None:
            QMessageBox.warning(
                self,
                "Missing CT Path",
                "CT file path is not set. Load a CT file before saving electrodes.",
            )
            return

        if self.fs_subject_path and getattr(self, "before_parcellation_volume", None) is None:
            self._load_before_parcellation_volume()

        contact_numbers = self._refresh_all_shaft_contact_numbers_for_export()
        export_indices = self._electrode_export_indices()
        rows = [
            self._electrode_export_row(
                electrodes[electrode_idx],
                export_idx,
                source_index=electrode_idx,
                shaft_contact_number=contact_numbers.get(electrode_idx),
            )
            for export_idx, electrode_idx in enumerate(export_indices)
        ]
        rows, restored_source_rows, dropped_source_duplicates = self._clean_stale_micro_source_rows(rows)
        if dropped_source_duplicates:
            self.log(
                f"Ignored {dropped_source_duplicates} stale macro-source duplicate rows from old micro-contact exports."
            )
        if restored_source_rows:
            self.log(
                f"Restored {restored_source_rows} stale macro-source rows to their original map shafts."
            )
        base_row_count = len(rows)
        rows = [row for row in rows if not self._is_synthetic_micro_export_row(row)]
        removed_micro_rows = base_row_count - len(rows)
        if removed_micro_rows:
            self.log(
                f"Ignored {removed_micro_rows} previously exported micro-contact rows before regenerating from the .map."
            )
        map_path = self.electrode_map_path or find_electrode_map(self._planning_search_paths())
        synthetic_micro_rows, unresolved_micro_shafts = self._synthetic_micro_rows_from_map(rows, map_path)
        if synthetic_micro_rows:
            rows.extend(synthetic_micro_rows)
            self.log(
                f"Added {len(synthetic_micro_rows)} micro-contact export rows from "
                f"{os.path.basename(map_path or self.electrode_map_path or '')}."
            )
        if unresolved_micro_shafts:
            self.log(
                "Could not infer source macro locations for micro shafts: "
                + ", ".join(unresolved_micro_shafts)
            )
        mat_path = self._find_electrodes_mat_path()
        mat_records = []
        mat_counts = None
        if mat_path:
            try:
                mat_records, mat_summary = self._load_electrodes_mat_records(mat_path)
                rows, mat_counts = self._merge_electrodes_mat_info(rows, mat_records, mat_path)
                if mat_records:
                    native_source = mat_summary.get("native_source") or "none"
                    mni_source = mat_summary.get("mni_source") or "none"
                    label_source = mat_summary.get("label_source") or "none"
                    self.log(
                        "Merged Electrodes.mat comparison "
                        f"({len(mat_records)} rows; native={native_source}, mni={mni_source}, labels={label_source})."
                    )
                else:
                    self.log(f"Electrodes.mat found but no coordinate rows were recognized: {mat_path}")
            except Exception as exc:
                mat_counts = None
                self.log(f"Could not merge Electrodes.mat comparison from {mat_path}: {exc}")
        else:
            rows, mat_counts = self._merge_electrodes_mat_info(rows, [], None)
            self.log("No Electrodes.mat file found next to the loaded MRI/CT or patient registration folder.")

        xlsx_fname = self._electrode_info_xlsx_path(self.ct_path)
        saved_paths = []
        try:
            with open(fname, "w", newline="", encoding="utf-8") as handle:
                writer = csv.DictWriter(handle, fieldnames=self._row_fieldnames(rows))
                writer.writeheader()
                writer.writerows(rows)
            self.log(f"Saved electrodes to: {fname}")
            saved_paths.append(fname)
        except Exception as e:
            QMessageBox.critical(self, "Save Error", f"Error saving electrode locations: {e}")
            return

        if xlsx_fname:
            try:
                self._save_electrode_workbook(xlsx_fname, rows, mat_records)
                self.log(f"Saved electrode comparison workbook to: {xlsx_fname}")
                saved_paths.append(xlsx_fname)
            except Exception as exc:
                self.log(f"Could not save electrode comparison workbook: {exc}")

        if mat_counts:
            self.log(
                "Electrodes.mat comparison: "
                f"{mat_counts.get('matched', 0)} matched, "
                f"{mat_counts.get('close', 0)} close, "
                f"{mat_counts.get('off', 0)} very off, "
                f"{mat_counts.get('review', 0)} review, "
                f"{mat_counts.get('unmatched', 0)} unmatched."
            )
        QMessageBox.information(self, "Saved", "Electrode locations saved to:\n" + "\n".join(saved_paths))

    def _load_electrode_locations(self, ct_path):
        info_path = self._electrode_info_path(ct_path)
        if not info_path or not os.path.exists(info_path):
            return False

        def _parse_text(value):
            if value is None:
                return None
            text = str(value).strip()
            return text or None

        def _parse_float(value):
            text = _parse_text(value)
            if text is None:
                return None
            try:
                return float(text)
            except (TypeError, ValueError):
                return None

        def _parse_int(value):
            text = _parse_text(value)
            if text is None:
                return None
            try:
                return int(float(text))
            except (TypeError, ValueError):
                return None

        def _parse_bool(value, default=False):
            if isinstance(value, bool):
                return value
            text = _parse_text(value)
            if text is None:
                return default
            text = text.lower()
            if text in ("true", "1", "yes", "y", "t"):
                return True
            if text in ("false", "0", "no", "n", "f"):
                return False
            return default

        try:
            with open(info_path, newline="", encoding="utf-8") as handle:
                reader = csv.DictReader(handle)
                rows = list(reader)
        except Exception as e:
            self.log(f"Error reading electrode info: {e}")
            return False

        if not rows:
            self.log(f"No electrode entries found in: {info_path}")
            return False

        self._refresh_planning_trajectories()
        self._refresh_electrode_map()
        restored_micro_assignments = self._restore_micro_map_assignments_from_export_rows(rows)
        if restored_micro_assignments:
            self.log(
                f"Restored {restored_micro_assignments} manual micro-contact trajectory assignments from saved rows."
            )

        indices = []
        has_indices = True
        for row in rows:
            idx = _parse_int(row.get("index"))
            if idx is None:
                has_indices = False
                break
            indices.append(idx)
        if has_indices:
            rows = [row for _, row in sorted(zip(indices, rows), key=lambda item: item[0])]

        electrodes = []
        display_spaces = []
        skipped_synthetic_micro_rows = 0
        rows, restored_source_rows, dropped_source_duplicates = self._clean_stale_micro_source_rows(rows)
        if dropped_source_duplicates:
            self.log(
                f"Skipped {dropped_source_duplicates} stale macro-source duplicate rows from old micro-contact exports."
            )
        if restored_source_rows:
            self.log(
                f"Restored {restored_source_rows} stale macro-source rows to their original map shafts while loading."
            )
        for row in rows:
            if self._is_synthetic_micro_export_row(row):
                skipped_synthetic_micro_rows += 1
                continue
            x = _parse_float(row.get("x"))
            y = _parse_float(row.get("y"))
            z = _parse_float(row.get("z"))
            if x is None or y is None or z is None:
                continue
            shaft = _parse_text(row.get("shaft")) or "Unassigned"
            name = _parse_text(row.get("name"))
            space = _parse_text(row.get("space")) or self._current_electrode_space()
            tissue = _parse_text(row.get("tissue")) or "unknown"
            detected = _parse_bool(row.get("detected"), default=False)
            visible = _parse_bool(row.get("visible"), default=True)
            view = _parse_int(row.get("view"))
            slice_axial = _parse_int(row.get("slice_axial"))
            slice_coronal = _parse_int(row.get("slice_coronal"))
            slice_sagittal = _parse_int(row.get("slice_sagittal"))
            display_space = _parse_text(row.get("display_space"))
            mri_world_x = _parse_float(row.get("mri_world_x"))
            mri_world_y = _parse_float(row.get("mri_world_y"))
            mri_world_z = _parse_float(row.get("mri_world_z"))
            mni_x = _parse_float(row.get("mni_x"))
            mni_y = _parse_float(row.get("mni_y"))
            mni_z = _parse_float(row.get("mni_z"))
            mni_space = _parse_text(row.get("mni_space"))
            parcellation = _parse_text(row.get("parcellation"))
            trajectory_name = _parse_text(row.get("trajectory_name"))
            trajectory_label = _parse_text(row.get("trajectory_label"))
            planned_contacts = _parse_int(row.get("planned_contacts"))
            trajectory_order = _parse_int(row.get("trajectory_order"))
            trajectory_hemisphere = _parse_text(row.get("trajectory_hemisphere"))
            planning_assignment_method = _parse_text(row.get("planning_assignment_method"))
            planning_assignment_cost = _parse_float(row.get("planning_assignment_cost"))
            trajectory_schema_ap = _parse_float(row.get("trajectory_schema_ap"))
            trajectory_schema_si = _parse_float(row.get("trajectory_schema_si"))
            shaft_entry_ap = _parse_float(row.get("shaft_entry_ap"))
            shaft_entry_si = _parse_float(row.get("shaft_entry_si"))
            planning_worksheet = _parse_text(row.get("planning_worksheet"))
            map_shaft = _parse_text(row.get("map_shaft"))
            map_contact = _parse_text(row.get("map_contact"))
            map_contact_number = _parse_int(row.get("map_contact_number"))
            shaft_contact_number = _parse_int(row.get("shaft_contact_number"))
            if shaft_contact_number is None:
                shaft_contact_number = map_contact_number
            electrode_map = _parse_text(row.get("electrode_map"))
            electrodes.append(Electrode(
                x=x,
                y=y,
                z=z,
                shaft=shaft,
                view=view,
                name=name,
                visible=visible,
                detected=detected,
                space=space,
                tissue=tissue,
                slice_axial=slice_axial,
                slice_coronal=slice_coronal,
                slice_sagittal=slice_sagittal,
                mri_world_x=mri_world_x,
                mri_world_y=mri_world_y,
                mri_world_z=mri_world_z,
                mni_x=mni_x,
                mni_y=mni_y,
                mni_z=mni_z,
                mni_space=mni_space,
                parcellation=parcellation,
                trajectory_name=trajectory_name,
                trajectory_label=trajectory_label,
                planned_contacts=planned_contacts,
                trajectory_order=trajectory_order,
                trajectory_hemisphere=trajectory_hemisphere,
                planning_assignment_method=planning_assignment_method,
                planning_assignment_cost=planning_assignment_cost,
                trajectory_schema_ap=trajectory_schema_ap,
                trajectory_schema_si=trajectory_schema_si,
                shaft_entry_ap=shaft_entry_ap,
                shaft_entry_si=shaft_entry_si,
                planning_worksheet=planning_worksheet,
                shaft_contact_number=shaft_contact_number,
                map_shaft=map_shaft,
                map_contact=map_contact,
                map_contact_number=map_contact_number,
                electrode_map=electrode_map,
            ))
            if display_space:
                display_spaces.append(display_space)

        if skipped_synthetic_micro_rows:
            self.log(
                f"Skipped {skipped_synthetic_micro_rows} derived micro-contact rows while loading saved electrodes; "
                "they will be regenerated from the .map on save."
            )

        if not electrodes:
            self.log(f"No valid electrodes loaded from: {info_path}")
            return False

        shafts = {}
        for idx, electrode in enumerate(electrodes):
            shaft_name = electrode.shaft or "Unassigned"
            electrode.shaft = shaft_name
            shafts.setdefault(shaft_name, []).append(idx)
        shafts = self._order_shaft_dict_distal_first(electrodes, shafts)
        shaft_visibility = {name: True for name in shafts}
        self._refresh_planning_trajectories()
        self._refresh_electrode_map()
        shafts, shaft_visibility, _ = self._assign_planning_names_to_shafts(
            electrodes,
            shafts,
            shaft_visibility,
            log=False,
        )
        for electrode in electrodes:
            shaft_name = electrode.shaft
            order = electrode.get("trajectory_order")
            if shaft_name and order is not None:
                current = self.planning_shaft_order.get(shaft_name)
                if current is None or order < current:
                    self.planning_shaft_order[shaft_name] = order
        has_detected = any(electrode.detected for electrode in electrodes)
        all_visible = all(electrode.visible for electrode in electrodes)

        self.before_electrodes = electrodes
        self.electrode_display_space = self._current_electrode_space()
        if display_spaces:
            candidate = display_spaces[0]
            if all(space == candidate for space in display_spaces):
                self.electrode_display_space = candidate
        self.before_shafts = shafts
        self.before_shaft_visibility = shaft_visibility
        self.before_electrode_count_label.setText(f"Electrodes: {len(electrodes)}")
        self.before_all_electrodes_visible = all_visible
        self.before_detected_visible = all_visible
        try:
            self.before_showhide_all_btn.setText(
                "Hide All Electrodes" if all_visible else "Show All Electrodes"
            )
        except Exception:
            pass
        try:
            self.before_group_shafts_btn.setEnabled(has_detected)
        except Exception:
            pass
        self._update_shaft_list_with_colors(self.before_shaft_list, self.before_shafts)
        if self.before_shaft_list.count() > 0:
            self.before_shaft_list.setCurrentRow(0)
            self._on_before_shaft_selected()

        detected_positions = [
            (electrode.x, electrode.y, electrode.z)
            for electrode in electrodes
            if electrode.detected
        ]
        self.before_detected_positions = (
            np.asarray(detected_positions) if detected_positions else None
        )

        self._update_mni_button_state()
        self._schedule_detected_sync(electrodes)
        self.log(f"Loaded electrodes from: {info_path}")
        return True
    
    def _create_visualization_controls(self):
        """Create visualization control panel with overlay toggles, transparency, and electrode tools."""
        group = QGroupBox("Visualization Controls")
        main_layout = QVBoxLayout()

        # Panel for all checkboxes/buttons
        controls_panel = QWidget()
        controls_layout = QGridLayout()
        controls_layout.setContentsMargins(6, 6, 6, 6)
        controls_layout.setHorizontalSpacing(10)
        controls_layout.setVerticalSpacing(6)

        mri_checkbox = QCheckBox("Show MRI")
        mri_checkbox.setChecked(True)
        mri_checkbox.setStyleSheet("font-weight: bold;")

        ct_checkbox = QCheckBox("Show CT")
        ct_checkbox.setChecked(True)
        ct_checkbox.setStyleSheet("font-weight: bold;")

        overlay_widget = QWidget()
        overlay_layout = QVBoxLayout()
        overlay_layout.setContentsMargins(0, 0, 0, 0)
        overlay_layout.setSpacing(2)
        overlay_layout.addWidget(mri_checkbox)
        overlay_layout.addWidget(ct_checkbox)
        overlay_widget.setLayout(overlay_layout)
        controls_layout.addWidget(overlay_widget, 0, 0)

        mip_checkbox = QCheckBox("Show MIP")
        mip_checkbox.setToolTip("Show maximum intensity projection (MIP) of CT volume")
        mip_checkbox.setStyleSheet("font-weight: bold; color: #cc6600;")
        shaft_label_checkbox = QCheckBox("Show Shaft Labels")
        shaft_label_checkbox.setToolTip("Toggle shaft labels in 2D views")
        shaft_label_checkbox.setChecked(True)

        mip_widget = QWidget()
        mip_layout = QVBoxLayout()
        mip_layout.setContentsMargins(0, 0, 0, 0)
        mip_layout.setSpacing(2)
        mip_layout.addWidget(mip_checkbox)
        mip_layout.addWidget(shaft_label_checkbox)
        mip_widget.setLayout(mip_layout)
        controls_layout.addWidget(mip_widget, 0, 1)

        electrode_checkbox = QCheckBox("Mark Electrodes")
        electrode_checkbox.setToolTip("Click on images to mark SEEG electrode locations")
        electrode_checkbox.setStyleSheet("font-weight: bold; color: #0066cc;")

        parcellation_checkbox = QCheckBox("Parcellation")
        parcellation_checkbox.setToolTip("Show cortical parcellation overlay on slice views when FreeSurfer data is available")
        parcellation_checkbox.setStyleSheet("color: #0066cc;")
        parcellation_checkbox.setChecked(False)
        parcellation_checkbox.setEnabled(False)

        electrode_widget = QWidget()
        electrode_widget_layout = QVBoxLayout()
        electrode_widget_layout.setContentsMargins(0, 0, 0, 0)
        electrode_widget_layout.setSpacing(2)
        electrode_widget_layout.addWidget(electrode_checkbox)
        electrode_widget_layout.addWidget(parcellation_checkbox)
        electrode_widget.setLayout(electrode_widget_layout)
        controls_layout.addWidget(electrode_widget, 0, 2)

        electrode_count_label = QLabel("Electrodes: 0")
        electrode_count_label.setStyleSheet("color: #0066cc;")
        controls_layout.addWidget(electrode_count_label, 0, 3, Qt.AlignRight)

        alpha_label = QLabel("CT Transparency:")
        alpha_label.setStyleSheet("font-weight: bold;")
        controls_layout.addWidget(alpha_label, 1, 0)

        alpha_slider = QSlider(Qt.Horizontal)
        alpha_slider.setMinimum(0)
        alpha_slider.setMaximum(100)
        alpha_slider.setValue(80)
        alpha_slider.setToolTip("Adjust CT overlay transparency (0=transparent, 100=opaque)")
        controls_layout.addWidget(alpha_slider, 1, 1, 1, 2)

        alpha_value_label = QLabel("80%")
        controls_layout.addWidget(alpha_value_label, 1, 3)

        threshold_label = QLabel("Detection Threshold (HU):")
        threshold_label.setStyleSheet("font-weight: bold; color: #006600;")
        controls_layout.addWidget(threshold_label, 2, 0)

        threshold_lower_spinbox = QDoubleSpinBox()
        threshold_lower_spinbox.setRange(-1000, 5000)
        threshold_lower_spinbox.setValue(2000)
        threshold_lower_spinbox.setDecimals(0)
        threshold_lower_spinbox.setSuffix(" HU (min)")
        controls_layout.addWidget(threshold_lower_spinbox, 2, 1)

        threshold_upper_spinbox = QDoubleSpinBox()
        threshold_upper_spinbox.setRange(-1000, 10000)
        threshold_upper_spinbox.setValue(10000)
        threshold_upper_spinbox.setDecimals(0)
        threshold_upper_spinbox.setSuffix(" HU (max)")
        controls_layout.addWidget(threshold_upper_spinbox, 2, 2)

        mip_threshold_checkbox = QCheckBox("Mask MIP by Threshold")
        mip_threshold_checkbox.setToolTip("When MIP is enabled, hide voxels outside the detection threshold range")
        controls_layout.addWidget(mip_threshold_checkbox, 2, 3)

        axial_wiggle_label = QLabel("Line Wiggle (vox):")
        axial_wiggle_label.setStyleSheet("font-weight: bold; color: #006600;")
        controls_layout.addWidget(axial_wiggle_label, 3, 0)

        axial_wiggle_spinbox = QDoubleSpinBox()
        axial_wiggle_spinbox.setRange(1, 10)
        axial_wiggle_spinbox.setValue(DEFAULT_AXIAL_LINE_WIGGLE)
        axial_wiggle_spinbox.setDecimals(1)
        axial_wiggle_spinbox.setSingleStep(0.5)
        axial_wiggle_spinbox.setSuffix(" vox")
        axial_wiggle_spinbox.setToolTip("Max distance from a straight line when grouping in 3D")
        controls_layout.addWidget(axial_wiggle_spinbox, 3, 1)

        axial_gap_label = QLabel("Max Gap (vox):")
        axial_gap_label.setStyleSheet("font-weight: bold; color: #006600;")
        controls_layout.addWidget(axial_gap_label, 3, 2)

        axial_gap_spinbox = QDoubleSpinBox()
        axial_gap_spinbox.setRange(1, 50)
        axial_gap_spinbox.setValue(DEFAULT_AXIAL_MAX_GAP)
        axial_gap_spinbox.setDecimals(1)
        axial_gap_spinbox.setSingleStep(1.0)
        axial_gap_spinbox.setSuffix(" vox")
        axial_gap_spinbox.setToolTip("Max spacing along a shaft line before splitting into a new shaft")
        controls_layout.addWidget(axial_gap_spinbox, 3, 3)

        min_contacts_label = QLabel("Min Contacts:")
        min_contacts_label.setStyleSheet("font-weight: bold; color: #006600;")
        controls_layout.addWidget(min_contacts_label, 4, 0)

        min_contacts_spinbox = QtWidgets.QSpinBox()
        min_contacts_spinbox.setRange(2, 32)
        min_contacts_spinbox.setValue(DEFAULT_GROUP_MIN)
        min_contacts_spinbox.setSuffix(" contacts")
        min_contacts_spinbox.setToolTip("Minimum contacts required before a grouped line is accepted as a shaft")
        controls_layout.addWidget(min_contacts_spinbox, 4, 1)

        auto_detect_btn = QPushButton("Auto-Detect")
        auto_detect_btn.setToolTip("Automatically detect electrodes from CT")
        controls_layout.addWidget(auto_detect_btn, 5, 0)

        group_shafts_btn = QPushButton("Group Shafts")
        group_shafts_btn.setToolTip("Cluster detected electrodes into shafts")
        group_shafts_btn.setEnabled(False)
        controls_layout.addWidget(group_shafts_btn, 5, 1)

        classify_btn = QPushButton("Classify Tissue")
        classify_btn.setToolTip("Classify detected electrodes as gray/white using FreeSurfer aseg")
        classify_btn.setEnabled(False)
        if self.controller is not None and hasattr(self.controller, "_on_classify_tissue_clicked"):
            classify_btn.clicked.connect(self.controller._on_classify_tissue_clicked)
        controls_layout.addWidget(classify_btn, 5, 2)
        self.classify_tissue_btn = classify_btn
        if self.controller is not None:
            self.controller.classify_tissue_btn = classify_btn
            if hasattr(self.controller, "_update_classify_button_state"):
                self.controller._update_classify_button_state()

        showhide_all_btn = QPushButton("Hide All Electrodes")
        showhide_all_btn.setToolTip("Toggle visibility of every electrode marker")
        controls_layout.addWidget(showhide_all_btn, 5, 3)

        clear_electrodes_btn = QPushButton("Clear Electrodes")
        clear_electrodes_btn.setToolTip("Remove all electrode markers")
        clear_electrodes_btn.setSizePolicy(QtWidgets.QSizePolicy.Expanding, QtWidgets.QSizePolicy.Fixed)
        controls_layout.addWidget(clear_electrodes_btn, 6, 0)

        transform_mni_btn = QPushButton("MRI->MNI")
        transform_mni_btn.setToolTip("Transform registered MRI-space electrodes to MNI Talairach space")
        transform_mni_btn.setEnabled(False)
        transform_mni_btn.setSizePolicy(QtWidgets.QSizePolicy.Expanding, QtWidgets.QSizePolicy.Fixed)
        controls_layout.addWidget(transform_mni_btn, 6, 1)

        map_trajectories_btn = QPushButton("Map Trajectories")
        map_trajectories_btn.setToolTip("Match .map electrode labels such as LN to planning worksheet trajectories")
        map_trajectories_btn.setSizePolicy(QtWidgets.QSizePolicy.Expanding, QtWidgets.QSizePolicy.Fixed)
        controls_layout.addWidget(map_trajectories_btn, 6, 2)

        save_electrodes_btn = QPushButton("Save Electrodes")
        save_electrodes_btn.setToolTip("Export electrode locations and Electrodes.mat comparison workbook")
        save_electrodes_btn.setSizePolicy(QtWidgets.QSizePolicy.Expanding, QtWidgets.QSizePolicy.Fixed)
        controls_layout.addWidget(save_electrodes_btn, 6, 3)

        controls_layout.setColumnStretch(0, 0)
        controls_layout.setColumnStretch(1, 1)
        controls_layout.setColumnStretch(2, 1)
        controls_layout.setColumnStretch(3, 0)

        controls_panel.setLayout(controls_layout)
        main_layout.addWidget(controls_panel)
        group.setLayout(main_layout)

        # Shafts list group
        shaft_group = QGroupBox("Shafts")
        shaft_layout = QVBoxLayout()
        shaft_list = QListWidget()
        shaft_list.setToolTip("Select shaft to view electrodes")
        shaft_list.setSelectionMode(QtWidgets.QAbstractItemView.ExtendedSelection)
        shaft_layout.addWidget(shaft_list)
        shaft_group.setLayout(shaft_layout)
        shaft_group.setMinimumWidth(220)

        # Electrodes list group
        electrode_group = QGroupBox("Electrodes")
        electrode_layout = QVBoxLayout()
        electrode_list = QListWidget()
        electrode_list.setSelectionMode(QtWidgets.QAbstractItemView.ExtendedSelection)
        electrode_list.setToolTip("Select electrode to view location")
        electrode_layout.addWidget(electrode_list)
        electrode_group.setLayout(electrode_layout)
        electrode_group.setMinimumWidth(180)

        self.before_mri_checkbox = mri_checkbox
        self.before_ct_checkbox = ct_checkbox
        self.before_alpha_slider = alpha_slider
        self.before_alpha_label = alpha_value_label
        self.before_electrode_checkbox = electrode_checkbox
        self.before_axial_wiggle_spinbox = axial_wiggle_spinbox
        self.before_axial_gap_spinbox = axial_gap_spinbox
        self.before_group_min_spinbox = min_contacts_spinbox
        self.before_electrode_count_label = electrode_count_label
        self.before_auto_detect_btn = auto_detect_btn
        self.before_clear_electrodes_btn = clear_electrodes_btn
        self.before_showhide_all_btn = showhide_all_btn
        self.before_group_shafts_btn = group_shafts_btn
        self.before_map_trajectories_btn = map_trajectories_btn
        self.before_shaft_list = shaft_list
        self.before_electrode_list = electrode_list
        self.transform_mni_btn = transform_mni_btn
        self.before_save_electrodes_btn = save_electrodes_btn
        self.before_shaft_group = shaft_group
        self.before_electrode_group = electrode_group
        self.before_mip_checkbox = mip_checkbox
        self.before_shaft_label_checkbox = shaft_label_checkbox
        self.before_mip_threshold_checkbox = mip_threshold_checkbox
        self.before_threshold_lower = threshold_lower_spinbox
        self.before_threshold_upper = threshold_upper_spinbox
        self.before_parcellation_checkbox = parcellation_checkbox
        self.before_all_electrodes_visible = True

        mri_checkbox.stateChanged.connect(self._update_before_overlay)
        ct_checkbox.stateChanged.connect(self._update_before_overlay)
        alpha_slider.valueChanged.connect(self._update_before_transparency)
        electrode_checkbox.stateChanged.connect(self._toggle_before_electrode_marking)
        parcellation_checkbox.stateChanged.connect(self._toggle_before_parcellation_overlay)
        auto_detect_btn.clicked.connect(self._auto_detect_before_electrodes)
        group_shafts_btn.clicked.connect(self._group_before_shafts)
        clear_electrodes_btn.clicked.connect(self._clear_before_electrodes)
        transform_mni_btn.clicked.connect(self._transform_electrodes_to_mni)
        mip_checkbox.stateChanged.connect(self._toggle_before_mip)
        shaft_label_checkbox.stateChanged.connect(self._update_before_shaft_labels)
        mip_threshold_checkbox.stateChanged.connect(self._update_before_mip_mask)
        threshold_lower_spinbox.valueChanged.connect(self._on_before_threshold_changed)
        self._sync_before_parcellation_checkbox()
        threshold_upper_spinbox.valueChanged.connect(self._on_before_threshold_changed)
        shaft_list.itemSelectionChanged.connect(self._on_before_shaft_selected)
        shaft_list.itemPressed.connect(self._on_before_shaft_pressed)
        shaft_list.itemClicked.connect(self._on_before_shaft_clicked)
        electrode_list.itemSelectionChanged.connect(self._on_before_electrode_selected)
        map_trajectories_btn.clicked.connect(self._open_map_trajectory_dialog)
        showhide_all_btn.clicked.connect(self._toggle_all_electrodes)
        save_electrodes_btn.clicked.connect(self._save_electrode_locations)
        self._update_mni_button_state()

        return group, shaft_group, electrode_group
    
    def _create_slice_controls(self):
        """Create slice control sliders for axial, sagittal, and coronal views.
        """
        controls = {}

        def _make_control(label_text):
            widget = QWidget()
            layout = QVBoxLayout()
            layout.setContentsMargins(0, 0, 0, 0)
            layout.setSpacing(2)

            label = QLabel(f"{label_text} Slice: 0/0")
            label.setStyleSheet("font-weight: bold;")
            layout.addWidget(label)

            slider = QSlider(Qt.Horizontal)
            slider.setMinimum(0)
            slider.setMaximum(100)  # Will be updated when data is loaded
            slider.setValue(50)
            slider.setEnabled(False)
            layout.addWidget(slider)

            widget.setLayout(layout)
            widget.setMaximumHeight(70)
            return widget, label, slider

        ax_widget, ax_label, ax_slider = _make_control("Axial")
        cor_widget, cor_label, cor_slider = _make_control("Coronal")
        sag_widget, sag_label, sag_slider = _make_control("Sagittal")
        controls["axial"] = ax_widget
        controls["coronal"] = cor_widget
        controls["sagittal"] = sag_widget
        
        self.before_ax_slider = ax_slider
        self.before_ax_label = ax_label
        self.before_cor_slider = cor_slider
        self.before_cor_label = cor_label
        self.before_sag_slider = sag_slider
        self.before_sag_label = sag_label

        # Connect signals
        ax_slider.valueChanged.connect(lambda v: self._update_before_slice('z', v))
        cor_slider.valueChanged.connect(lambda v: self._update_before_slice('y', v))
        sag_slider.valueChanged.connect(lambda v: self._update_before_slice('x', v))
        
        return controls

    def _toggle_view_visibility(self, view_key, state):
        """Show or hide individual slice views and refresh the canvas."""
        if isinstance(state, bool):
            show = state
        else:
            show = state == Qt.Checked
        setattr(self, f"before_show_{view_key}", show)
        btn = self.before_view_buttons.get(view_key)
        if btn is not None and btn.isChecked() != show:
            btn.blockSignals(True)
            btn.setChecked(show)
            btn.blockSignals(False)
        container = self.before_pg_containers.get(view_key)
        if container is not None:
            container.setVisible(show)
            if show:
                container.updateGeometry()
        view = self.before_pg_views.get(view_key)
        if view is not None:
            view_box = view.get("view")
            if view_box is not None:
                view_box.setVisible(show)
                if show:
                    try:
                        view_box.autoRange(padding=0.0)
                    except Exception:
                        pass
        if show:
            QtCore.QTimer.singleShot(
                0, lambda key=view_key: self._refresh_before_view_range(key)
            )
        self._refresh_before_view_layout()
        if self.before_mri_data is not None:
            self._plot_slices(self.before_mri_data, self.before_ct_data, self.before_canvas,
                              "Unaligned CT Overlaid on MRI",
                              self.before_slice_x, self.before_slice_y, self.before_slice_z)

    def _refresh_before_view_layout(self):
        """Ensure visible views get equal stretch after toggling."""
        if self.before_canvas_layout is None:
            return
        for idx, view_key in enumerate(self.before_view_order):
            container = self.before_pg_containers.get(view_key)
            if container is None:
                continue
            stretch = 1 if container.isVisible() else 0
            self.before_canvas_layout.setStretch(idx, stretch)
        try:
            self.before_canvas_layout.invalidate()
            self.before_canvas_layout.activate()
        except Exception:
            pass
        if self.before_canvas is not None:
            self.before_canvas.updateGeometry()

    def _refresh_before_view_range(self, view_key):
        """Reset view range after a toggle to avoid collapsed scenes."""
        view = self.before_pg_views.get(view_key)
        if view is None:
            return
        view_box = view.get("view")
        if view_box is None:
            return
        mri_slice, _ = self.before_pg_slices.get(view_key, (None, None))
        if mri_slice is None:
            return
        mri_slice = np.asarray(mri_slice)
        if mri_slice.ndim != 2:
            return
        height, width = mri_slice.shape
        view_box.setRange(QtCore.QRectF(0, 0, width, height), padding=0.0)
        view_box.update()

    def _capture_before_pg_view_ranges(self):
        """Capture current PyQtGraph zoom ranges for visible slice views."""
        ranges = {}
        for view_key, view in getattr(self, "before_pg_views", {}).items():
            view_box = view.get("view")
            if view_box is None or not view_box.isVisible():
                continue
            mri_slice, _ = self.before_pg_slices.get(view_key, (None, None))
            if mri_slice is None:
                continue
            try:
                mri_slice = np.asarray(mri_slice)
                if mri_slice.ndim != 2:
                    continue
                x_range, y_range = view_box.viewRange()
                x_range = tuple(float(value) for value in x_range)
                y_range = tuple(float(value) for value in y_range)
            except Exception:
                continue
            values = x_range + y_range
            if not all(np.isfinite(value) for value in values):
                continue
            if abs(x_range[1] - x_range[0]) <= 1e-6 or abs(y_range[1] - y_range[0]) <= 1e-6:
                continue
            ranges[view_key] = {
                "shape": tuple(mri_slice.shape[:2]),
                "x_range": x_range,
                "y_range": y_range,
            }
        return ranges

    def _set_before_pg_view_range(self, view_key, view_box, width, height, saved_ranges=None):
        """Restore a saved zoom range when compatible, otherwise show the full image."""
        saved = (saved_ranges or {}).get(view_key)
        if saved and saved.get("shape") == (height, width):
            try:
                view_box.setRange(
                    xRange=saved["x_range"],
                    yRange=saved["y_range"],
                    padding=0.0,
                )
                return
            except Exception:
                pass
        view_box.setRange(QtCore.QRectF(0, 0, width, height), padding=0.0)
    
    def _update_before_slice(self, axis, value):
        """Update before view when slice changes."""
        if self.before_mri_data is None:
            return
            
        if axis == 'x':
            self.before_slice_x = value
            self.before_sag_label.setText(f"Sagittal Slice: {value}/{self.before_mri_data.shape[0]-1}")
        elif axis == 'y':
            self.before_slice_y = value
            self.before_cor_label.setText(f"Coronal Slice: {value}/{self.before_mri_data.shape[1]-1}")
        elif axis == 'z':
            self.before_slice_z = value
            self.before_ax_label.setText(f"Axial Slice: {value}/{self.before_mri_data.shape[2]-1}")
        
        self._plot_slices(self.before_mri_data, self.before_ct_data, self.before_canvas,
                         "Unaligned CT Overlaid on MRI",
                         self.before_slice_x, self.before_slice_y, self.before_slice_z)
    
    def _update_before_overlay(self):
        """Update before view when MRI/CT toggles change."""
        self.before_show_mri = self.before_mri_checkbox.isChecked()
        self.before_show_ct = self.before_ct_checkbox.isChecked()
        if self.before_mri_data is not None:
            self._plot_slices(self.before_mri_data, self.before_ct_data, self.before_canvas,
                            "Unaligned CT Overlaid on MRI",
                            self.before_slice_x, self.before_slice_y, self.before_slice_z)

    def _update_before_shaft_labels(self):
        """Toggle shaft labels in before view."""
        self.before_show_shaft_labels = bool(self.before_shaft_label_checkbox.isChecked())
        if self.before_mri_data is not None:
            self._plot_slices(self.before_mri_data, self.before_ct_data, self.before_canvas,
                            "Unaligned CT Overlaid on MRI",
                            self.before_slice_x, self.before_slice_y, self.before_slice_z)
    
    def _update_before_transparency(self, value):
        """Update before CT transparency."""
        self.before_ct_alpha = value / 100.0
        self.before_alpha_label.setText(f"{value}%")
        if self.before_mri_data is not None:
            self._plot_slices(self.before_mri_data, self.before_ct_data, self.before_canvas,
                            "Unaligned CT Overlaid on MRI",
                            self.before_slice_x, self.before_slice_y, self.before_slice_z)
    
    def _toggle_before_parcellation_overlay(self, state):
        """Enable or disable FreeSurfer parcellation overlay on slice views."""
        self.before_parcellation_overlay = bool(state)
        if self.before_mri_data is not None:
            self._plot_slices(self.before_mri_data, self.before_ct_data, self.before_canvas,
                            "Unaligned CT Overlaid on MRI",
                            self.before_slice_x, self.before_slice_y, self.before_slice_z)

    def _sync_before_parcellation_checkbox(self):
        """Sync the parcellation overlay checkbox with availability and current state."""
        checkbox = getattr(self, "before_parcellation_checkbox", None)
        if checkbox is None:
            return

        checkbox.blockSignals(True)
        checkbox.setChecked(bool(self.before_parcellation_overlay))
        checkbox.setEnabled(bool(self.before_parcellation_available))
        if self.before_parcellation_available:
            checkbox.setToolTip("Show cortical parcellation overlay on slice views")
        else:
            checkbox.setToolTip(
                "FreeSurfer parcellation overlay unavailable until a valid subject folder is loaded."
            )
        checkbox.blockSignals(False)

    def _load_before_parcellation_volume(self):
        """Load FreeSurfer parcellation volume for slice overlay if available."""
        self.before_parcellation_volume = None
        self.before_parcellation_available = False
        self.before_parcellation_mismatch_warned = False
        self.before_selected_parcellation_label = None
        self.before_selected_parcellation_name = None
        if not self.fs_subject_path:
            self._sync_before_parcellation_checkbox()
            return

        self.before_parcellation_lut = self._load_freesurfer_color_lut()
        candidate_names = ["aparc+aseg.mgz", "aseg.mgz", "aparc.DKTatlas+aseg.mgz"]
        mri_dir = os.path.join(self.fs_subject_path, "mri")
        if not os.path.isdir(mri_dir):
            self._sync_before_parcellation_checkbox()
            return

        for filename in candidate_names:
            path = os.path.join(mri_dir, filename)
            if not os.path.exists(path):
                continue
            try:
                img = nib.load(path)
                target_img = getattr(self, "before_mri_img", None)
                if target_img is None and self.mri_path and os.path.exists(self.mri_path):
                    target_img = nib.load(self.mri_path)
                    self.before_mri_img = target_img

                if target_img is not None:
                    from nibabel.processing import resample_from_to

                    target = (target_img.shape[:3], target_img.affine)
                    img = resample_from_to(img, target, order=0, mode="constant", cval=0)

                volume = np.rint(np.asarray(img.get_fdata(dtype=np.float32))).astype(np.int32)
                if volume.ndim == 3:
                    self.before_parcellation_volume = volume
                    self.before_parcellation_available = True
                    self.log(
                        f"Loaded FreeSurfer parcellation volume: {path}"
                        + (" (resampled to displayed MRI grid)" if target_img is not None else "")
                    )
                    break
            except Exception as exc:
                self.log(f"Could not load FreeSurfer parcellation volume from {path}: {exc}")

        self._sync_before_parcellation_checkbox()

    def _load_freesurfer_color_lut(self):
        """Load FreeSurfer label id -> name mappings from FreeSurferColorLUT.txt."""
        candidates = []
        fs_home = os.environ.get("FREESURFER_HOME")
        if fs_home:
            candidates.append(os.path.join(fs_home, "FreeSurferColorLUT.txt"))
        subject_path = getattr(self, "fs_subject_path", None)
        if subject_path:
            path = Path(subject_path)
            for parent in [path, *path.parents]:
                candidates.append(str(parent / "FreeSurferColorLUT.txt"))
        candidates.extend([
            "/Applications/freesurfer/8.1.0/FreeSurferColorLUT.txt",
            "/Applications/freesurfer/FreeSurferColorLUT.txt",
            "/usr/local/freesurfer/FreeSurferColorLUT.txt",
        ])

        for path in candidates:
            if not path or not os.path.exists(path):
                continue
            lut = {}
            try:
                with open(path, "r", encoding="utf-8", errors="replace") as handle:
                    for line in handle:
                        text = line.strip()
                        if not text or text.startswith("#"):
                            continue
                        parts = text.split()
                        if len(parts) < 2:
                            continue
                        try:
                            lut[int(parts[0])] = parts[1]
                        except ValueError:
                            continue
                if lut:
                    return lut
            except Exception as exc:
                self.log(f"Could not read FreeSurfer color LUT from {path}: {exc}")
        return {}

    def _get_before_parcellation_slice(self, view_key, slice_x, slice_y, slice_z):
        """Return a parcellation slice aligned to the requested view orientation."""
        vol = self.before_parcellation_volume
        if vol is None or vol.ndim != 3:
            return None
        if vol.shape != self.before_mri_data.shape:
            if not getattr(self, "before_parcellation_mismatch_warned", False):
                self.log(
                    "FreeSurfer parcellation grid does not match the displayed MRI grid; "
                    "reload the MRI or FreeSurfer subject to resample the overlay."
                )
                self.before_parcellation_mismatch_warned = True
            return None
        if view_key == "axial":
            return vol[:, :, slice_z].T
        if view_key == "coronal":
            return vol[:, slice_y, :].T
        if view_key == "sagittal":
            return vol[slice_x, :, :].T
        return None

    def _selected_parcellation_id(self):
        """Return the selected parcellation id as an int, if one is active."""
        label = getattr(self, "before_selected_parcellation_label", None)
        try:
            return int(label)
        except (TypeError, ValueError):
            return None

    def _parcellation_name_for_label(self, label_id):
        """Return a display name for a FreeSurfer parcellation label id."""
        try:
            label_id = int(label_id)
        except (TypeError, ValueError):
            return None
        lut = getattr(self, "before_parcellation_lut", {}) or {}
        return lut.get(label_id, f"FreeSurfer label {label_id}")

    def _parcellation_for_electrode(self, electrode):
        """Return (label_id, label_name) for an electrode in displayed MRI voxel space."""
        vol = getattr(self, "before_parcellation_volume", None)
        if vol is None and self.fs_subject_path:
            self._load_before_parcellation_volume()
            vol = getattr(self, "before_parcellation_volume", None)
        if vol is None or vol.ndim != 3:
            return None, None

        try:
            coords = np.array([electrode.x, electrode.y, electrode.z], dtype=float)
        except Exception:
            return None, None
        if not np.all(np.isfinite(coords)):
            return None, None

        ix, iy, iz = np.round(coords).astype(int)
        if ix < 0 or iy < 0 or iz < 0 or ix >= vol.shape[0] or iy >= vol.shape[1] or iz >= vol.shape[2]:
            return None, None

        label_id = int(vol[ix, iy, iz])
        if label_id == 0:
            return 0, "unlabeled/background"
        return label_id, self._parcellation_name_for_label(label_id)

    def _set_selected_parcellation_from_electrode(self, electrode_idx, electrode_name):
        """Update the active parcellation highlight from the selected electrode."""
        self.before_selected_parcellation_label = None
        self.before_selected_parcellation_name = None

        if electrode_idx is None or electrode_idx >= len(self.before_electrodes):
            return
        electrode = self.before_electrodes[electrode_idx]
        label_id, label_name = self._parcellation_for_electrode(electrode)

        if label_id is None:
            self.log(f"No FreeSurfer parcellation volume available for '{electrode_name}'.")
            return

        if label_id == 0:
            self.log(f"'{electrode_name}' is outside labeled FreeSurfer parcellations.")
            return

        self.before_selected_parcellation_label = label_id
        self.before_selected_parcellation_name = label_name
        self.before_parcellation_overlay = True
        self._sync_before_parcellation_checkbox()
        self.log(f"'{electrode_name}' is in FreeSurfer parcellation: {label_name} ({label_id}).")

    def _make_parcellation_rgba(self, label_slice):
        """Create an RGBA overlay image for a parcellation label slice."""
        if label_slice is None or label_slice.size == 0:
            return None
        labels = np.asarray(label_slice, dtype=np.int32)
        if labels.ndim != 2:
            return None

        filled = labels != 0
        if not filled.any():
            return None

        boundaries = np.zeros_like(filled)
        boundaries[1:, :] |= labels[1:, :] != labels[:-1, :]
        boundaries[:-1, :] |= labels[:-1, :] != labels[1:, :]
        boundaries[:, 1:] |= labels[:, 1:] != labels[:, :-1]
        boundaries[:, :-1] |= labels[:, :-1] != labels[:, 1:]
        boundaries &= filled

        selected_label = self._selected_parcellation_id()
        selected = labels == selected_label if selected_label is not None else np.zeros_like(filled)

        if not boundaries.any() and not selected.any():
            return None

        rgba = np.zeros(labels.shape + (4,), dtype=np.uint8)
        rgba[..., 0] = 255
        rgba[..., 1] = 215
        rgba[..., 2] = 0
        rgba[..., 3] = np.where(boundaries, 180, 0).astype(np.uint8)

        if selected_label is not None:
            if selected.any():
                rgba[selected, 0] = 0
                rgba[selected, 1] = 210
                rgba[selected, 2] = 255
                rgba[selected, 3] = 95
                selected_boundary = boundaries & selected
                rgba[selected_boundary, 0] = 255
                rgba[selected_boundary, 1] = 255
                rgba[selected_boundary, 2] = 255
                rgba[selected_boundary, 3] = 230
        return rgba

    def _update_before_mip_mask(self):
        """Update before view when MIP threshold mask toggles."""
        if self.before_mri_data is not None:
            self._plot_slices(self.before_mri_data, self.before_ct_data, self.before_canvas,
                            "Unaligned CT Overlaid on MRI",
                            self.before_slice_x, self.before_slice_y, self.before_slice_z)

    def _populate_detected_electrodes(self, positions_3d):
        """Populate detected electrodes without grouping into shafts."""
        if positions_3d is None or len(positions_3d) == 0:
            self.log("No detected electrodes to populate.")
            return

        electrodes = []
        shafts = {"Detected": []}
        shaft_visibility = {"Detected": True}

        for pos_3d in positions_3d:
            x_idx, y_idx, z_idx = [int(round(v)) for v in pos_3d]
            electrode_idx = len(electrodes)
            electrodes.append(Electrode(
                x=x_idx,
                y=y_idx,
                z=z_idx,
                shaft="Detected",
                view=None,
                visible=True,
                detected=True,
                space=self.electrode_space
            ))
            shafts["Detected"].append(electrode_idx)

        self.before_electrodes = electrodes
        self.electrode_display_space = self.electrode_space
        self.before_shafts = shafts
        self.before_shaft_visibility = shaft_visibility
        self.before_electrode_count_label.setText(f"Electrodes: {len(electrodes)}")
        self.before_all_electrodes_visible = True
        self.before_detected_visible = True
        try:
            self.before_showhide_all_btn.setText("Hide All Electrodes")
        except Exception:
            pass
        try:
            self.before_group_shafts_btn.setEnabled(True)
        except Exception:
            pass
        self._update_shaft_list_with_colors(self.before_shaft_list, self.before_shafts)
        if self.before_shaft_list.count() > 0:
            self.before_shaft_list.setCurrentRow(0)
            self._on_before_shaft_selected()

        self._update_mni_button_state()
        self._schedule_detected_sync(electrodes)

    def _schedule_detected_sync(self, electrodes):
        """Coalesce 3D viewer updates to keep the UI responsive."""
        self._pending_detected_sync = list(electrodes)
        if self._detected_sync_timer is None:
            self._detected_sync_timer = QtCore.QTimer(self)
            self._detected_sync_timer.setSingleShot(True)
            self._detected_sync_timer.timeout.connect(self._flush_detected_sync)
        self._detected_sync_timer.start(DETECTED_SYNC_DELAY_MS)

    def _flush_detected_sync(self):
        electrodes = self._pending_detected_sync
        self._pending_detected_sync = None
        if electrodes is None:
            return
        self._sync_detected_to_controller(electrodes)

    def _apply_shaft_visibility_to_electrodes(self, electrodes=None):
        """Store per-shaft visibility on electrodes for the 3D viewer."""
        target = list(electrodes if electrodes is not None else self.before_electrodes)
        shaft_visibility = getattr(self, "before_shaft_visibility", {}) or {}
        for electrode in target:
            try:
                shaft_name = electrode.get("shaft") or "Unassigned"
                electrode["shaft_visible"] = bool(shaft_visibility.get(shaft_name, True))
            except Exception:
                pass
        return target

    def _sync_detected_to_controller(self, electrodes):
        """Update the main 3D viewer with the current detected electrodes."""
        if self.controller is None:
            self._update_mni_button_state()
            return
        t0 = time.perf_counter() if PROFILE_MERGE else None
        synced_electrodes = self._apply_shaft_visibility_to_electrodes(electrodes)
        self.controller.detected_electrodes = synced_electrodes
        self.controller.detected_electrode_space = self._current_display_electrode_space()
        self.controller.detected_electrode_native_space = self._current_electrode_space()
        if not synced_electrodes and hasattr(self.controller, "clear_detected_electrodes"):
            self.controller.detected_selected_shaft = None
            self.controller.clear_detected_electrodes()
        elif hasattr(self.controller, "_display_detected_electrodes"):
            self.controller._display_detected_electrodes(
                self.controller.detected_electrodes,
                self.controller.detected_electrode_space,
            )
        if hasattr(self.controller, "_update_classify_button_state"):
            self.controller._update_classify_button_state()
        self._update_mni_button_state()
        if PROFILE_MERGE and t0 is not None:
            elapsed = time.perf_counter() - t0
            self.log(f"[profile] 3D update: {elapsed:.3f}s for {len(synced_electrodes)} electrodes")

    def _force_3d_shaft_refresh(self, electrodes):
        """Force a shaft-style refresh in the 3D viewer after grouping."""
        if self.controller is None:
            return
        if not hasattr(self.controller, "_display_detected_shafts"):
            return
        self.controller.detected_electrodes = self._apply_shaft_visibility_to_electrodes(electrodes)
        self.controller.detected_electrode_space = self._current_display_electrode_space()
        self.controller.detected_electrode_native_space = self._current_electrode_space()
        self.controller._display_detected_shafts(
            self.controller.detected_electrodes,
            self.controller.detected_electrode_space,
        )
        self._update_mni_button_state()

    def _group_before_shafts(self):
        """Group detected electrodes into shafts (before view)."""
        self._group_detected_shafts()

    def _groupable_electrode_positions(self):
        """Return current electrode positions and matching source objects for shaft grouping."""
        electrodes = list(getattr(self, "before_electrodes", []) or [])
        positions = []
        sources = []
        for electrode in electrodes:
            try:
                coords = (
                    float(electrode.x),
                    float(electrode.y),
                    float(electrode.z),
                )
            except (TypeError, ValueError, AttributeError):
                continue
            if not np.all(np.isfinite(coords)):
                continue
            positions.append(coords)
            sources.append(electrode)

        if positions:
            return np.asarray(positions, dtype=float), sources

        positions = getattr(self, "before_detected_positions", None)
        if positions is None or len(positions) == 0:
            return np.empty((0, 3), dtype=float), []
        return np.asarray(positions, dtype=float), None

    def _group_detected_shafts(self):
        """Group detected electrodes into shafts using deterministic line candidates."""
        pos_arr, source_electrodes = self._groupable_electrode_positions()
        if len(pos_arr) == 0:
            self.log("No electrodes to group. Run Auto-Detect or mark electrodes first.")
            return
        self.shaft_counter = 0

        min_samples = getattr(self, "before_group_min_spinbox", None)
        min_samples = min_samples.value() if min_samples is not None else DEFAULT_GROUP_MIN

        wiggle_spinbox = getattr(self, "before_axial_wiggle_spinbox", None)
        gap_spinbox = getattr(self, "before_axial_gap_spinbox", None)
        axial_wiggle = wiggle_spinbox.value() if wiggle_spinbox is not None else DEFAULT_AXIAL_LINE_WIGGLE
        axial_gap = gap_spinbox.value() if gap_spinbox is not None else DEFAULT_AXIAL_MAX_GAP
        manual_count = sum(
            1 for electrode in (source_electrodes or [])
            if not getattr(electrode, "detected", False)
        )

        labels = self._cluster_shafts_ransac(
            pos_arr,
            dist_thresh=axial_wiggle,
            max_gap=axial_gap,
            min_inliers=int(min_samples),
            n_iters=300
        )
        if labels is not None and len(labels) > 0 and np.any(labels >= 0):
            labels = self._split_clusters_by_gap(pos_arr, labels, axial_gap, min_contacts=int(min_samples))
            labels = self._split_clusters_by_world_gap(pos_arr, labels, axial_gap, min_contacts=int(min_samples))
            labels = self._split_clusters_by_hemisphere(pos_arr, labels, min_contacts=int(min_samples))
            labels = self._filter_small_shaft_clusters(labels, int(min_samples))
            recovery_dist = max(1.0, float(axial_wiggle) * 0.75)
            labels, recovered_count = self._recover_unassigned_to_nearest_shaft(
                pos_arr,
                labels,
                max_dist=recovery_dist,
                max_gap=float(axial_gap),
                min_contacts=int(min_samples),
            )
            labels = self._split_clusters_by_world_gap(pos_arr, labels, axial_gap, min_contacts=int(min_samples))
            labels = self._split_clusters_by_hemisphere(pos_arr, labels, min_contacts=int(min_samples))
            labels = self._filter_small_shaft_clusters(labels, int(min_samples))
        else:
            recovered_count = 0

        if labels is None or len(labels) == 0:
            self.log("Shaft grouping failed; leaving electrodes unassigned.")
            labels = np.full(len(pos_arr), -1, dtype=int)
        unassigned_count = int(np.sum(labels == -1))

        self._populate_grouped_shafts(
            pos_arr,
            labels,
            min_contacts=int(min_samples),
            projection_threshold=float(axial_wiggle),
            source_electrodes=source_electrodes,
        )
        shaft_count = len([name for name in self.before_shafts if name != "Unassigned"])
        recovered_text = f" ({recovered_count} recovered)" if recovered_count else ""
        manual_text = f", including {manual_count} manual" if manual_count else ""
        self.log(
            f"Grouped {len(pos_arr) - unassigned_count}/{len(pos_arr)} contacts{manual_text} "
            f"into {shaft_count} shafts; {unassigned_count} unassigned"
            f"{recovered_text}."
        )
        self._flush_detected_sync()
        electrodes = self.before_electrodes
        self._force_3d_shaft_refresh(electrodes)
        self._update_before_shaft_labels()

    def _split_clusters_by_gap(self, positions, labels, max_gap, min_contacts=DEFAULT_GROUP_MIN):
        """Split line clusters on large gaps and reject undersized fragments."""
        return split_line_clusters_by_gap(
            positions,
            labels,
            max_gap=max_gap,
            min_contacts=min_contacts,
        )

    def _positions_to_world(self, positions):
        """Convert voxel electrode positions to world/RAS coordinates when possible."""
        affine = self._electrode_voxel_affine()
        if affine is None:
            return None
        pts = np.asarray(positions, dtype=float)
        if pts.size == 0:
            return np.empty((0, 3), dtype=float)
        hom = np.column_stack((pts, np.ones(len(pts), dtype=float)))
        return (hom @ np.asarray(affine, dtype=float).T)[:, :3]

    def _split_clusters_by_world_gap(self, positions, labels, max_gap, min_contacts=DEFAULT_GROUP_MIN):
        """Apply the same gap rule in physical/world space when an affine is available."""
        world = self._positions_to_world(positions)
        if world is None:
            return labels
        return split_line_clusters_by_gap(
            world,
            labels,
            max_gap=max_gap,
            min_contacts=min_contacts,
        )

    def _split_clusters_by_hemisphere(self, positions, labels, min_contacts=DEFAULT_GROUP_MIN):
        """Prevent one shaft label from spanning left and right hemispheres."""
        world = self._positions_to_world(positions)
        if world is None or len(world) == 0:
            return labels

        # Leave a small midline dead band so contacts near x=0 do not force
        # a split by numerical jitter; mixed concrete sides are still split.
        midline_margin_mm = 1.0
        groups = np.full(len(world), None, dtype=object)
        groups[world[:, 0] < -midline_margin_mm] = "left"
        groups[world[:, 0] > midline_margin_mm] = "right"
        split = split_labels_by_group(labels, groups, min_contacts=min_contacts)
        if np.any(split != np.asarray(labels, dtype=int)):
            self.log("Split detected shaft labels at the left/right hemisphere boundary.")
        return split

    def _filter_small_shaft_clusters(self, labels, min_contacts):
        """Move clusters below the minimum contact count back to Unassigned."""
        labels = np.asarray(labels, dtype=int).copy()
        if labels.size == 0 or min_contacts <= 1:
            return labels
        for label in sorted(set(labels)):
            if label < 0:
                continue
            if int(np.sum(labels == label)) < int(min_contacts):
                labels[labels == label] = -1
        return labels

    def _recover_unassigned_to_nearest_shaft(
        self,
        positions,
        labels,
        max_dist,
        max_gap,
        min_contacts=DEFAULT_GROUP_MIN,
    ):
        """Attach unassigned contacts to nearby accepted shaft axes."""
        labels = np.asarray(labels, dtype=int).copy()
        positions = np.asarray(positions, dtype=float)
        if positions.size == 0 or not np.any(labels == -1):
            return labels, 0

        models = []
        for label in sorted(set(labels)):
            if label < 0:
                continue
            idxs = np.where(labels == label)[0]
            if len(idxs) < int(min_contacts):
                continue
            pts = positions[idxs]
            centroid, direction = fit_shaft_axis_pca(pts)
            if centroid is None or direction is None:
                continue
            norm = np.linalg.norm(direction)
            if norm < 1e-6:
                continue
            direction = direction / norm
            ts = (pts - centroid) @ direction
            models.append({
                "label": label,
                "centroid": centroid,
                "direction": direction,
                "t_min": float(np.min(ts)),
                "t_max": float(np.max(ts)),
            })

        if not models:
            return labels, 0

        recovered = 0
        for idx in np.where(labels == -1)[0]:
            point = positions[idx]
            candidates = []
            for model in models:
                _proj, t_scalar, perp_dist = project_point_to_line(
                    point,
                    model["centroid"],
                    model["direction"],
                )
                if perp_dist > float(max_dist):
                    continue
                if t_scalar < model["t_min"]:
                    end_gap = model["t_min"] - t_scalar
                elif t_scalar > model["t_max"]:
                    end_gap = t_scalar - model["t_max"]
                else:
                    end_gap = 0.0
                if max_gap > 0 and end_gap > float(max_gap):
                    continue
                score = perp_dist + (0.25 * end_gap)
                candidates.append((score, model["label"]))

            if candidates:
                candidates.sort(key=lambda item: item[0])
                if len(candidates) > 1:
                    ambiguity_margin = max(1.0, float(max_dist) * 0.25)
                    if (candidates[1][0] - candidates[0][0]) < ambiguity_margin:
                        continue
                labels[idx] = candidates[0][1]
                recovered += 1

        return labels, recovered

    def _populate_grouped_shafts(
        self,
        electrode_array,
        labels,
        manual_electrodes=None,
        min_contacts=DEFAULT_GROUP_MIN,
        projection_threshold=DEFAULT_AXIAL_LINE_WIGGLE,
        source_electrodes=None,
    ):
        """Populate grouped shafts and electrodes from clustered positions."""
        electrodes = []
        shafts = {}
        shaft_visibility = {}
        unassigned_positions = []
        unassigned_contacts = []
        source_electrodes = list(source_electrodes) if source_electrodes is not None else None

        def make_electrode(pos_3d, shaft_name, source_idx=None):
            x_idx, y_idx, z_idx = [int(round(v)) for v in np.asarray(pos_3d, dtype=float)]
            if source_electrodes is not None and source_idx is not None and source_idx < len(source_electrodes):
                electrode = source_electrodes[source_idx]
                electrode.x = x_idx
                electrode.y = y_idx
                electrode.z = z_idx
                electrode.shaft = shaft_name
                electrode.slice_axial = z_idx
                electrode.slice_coronal = y_idx
                electrode.slice_sagittal = x_idx
                electrode["3d_pos"] = (x_idx, y_idx, z_idx)
                if electrode.get("visible") is None:
                    electrode.visible = True
                if not electrode.get("space"):
                    electrode.space = self.electrode_space
                return electrode

            return Electrode(
                x=x_idx,
                y=y_idx,
                z=z_idx,
                shaft=shaft_name,
                view=None,
                visible=True,
                detected=True,
                space=self.electrode_space
            )

        for shaft_idx in sorted(set(labels)):
            if shaft_idx == -1:
                continue

            shaft_name = f"S_{self.shaft_counter}"
            self.shaft_counter += 1
            shafts[shaft_name] = []
            shaft_visibility[shaft_name] = True

            shaft_mask = labels == shaft_idx
            shaft_indices = np.where(shaft_mask)[0]
            shaft_electrodes = electrode_array[shaft_mask]
            if len(shaft_electrodes) < int(min_contacts):
                if source_electrodes is not None:
                    unassigned_contacts.extend((int(idx), electrode_array[int(idx)]) for idx in shaft_indices)
                else:
                    unassigned_positions.extend(shaft_electrodes)
                shafts.pop(shaft_name, None)
                shaft_visibility.pop(shaft_name, None)
                continue

            try:
                pts = np.asarray(shaft_electrodes, dtype=float)
                centroid, direction = fit_shaft_axis_pca(pts)
                if centroid is None or direction is None:
                    raise ValueError("Could not fit shaft axis")
                perp_threshold = max(float(projection_threshold), 0.1)
                ordered_contacts = []
                off_axis_count = 0
                for source_idx, p in zip(shaft_indices, pts):
                    _proj_pt, t_scalar, perp_dist = project_point_to_line(p, centroid, direction)
                    ordered_contacts.append((int(source_idx), p, t_scalar, perp_dist))
                    if perp_dist > perp_threshold:
                        off_axis_count += 1

                if len(ordered_contacts) == 0:
                    ordered_contacts = [
                        (int(source_idx), p, float(np.dot((p - centroid), direction)), 0.0)
                        for source_idx, p in zip(shaft_indices, pts)
                    ]

                if len(ordered_contacts) < int(min_contacts):
                    if source_electrodes is not None:
                        unassigned_contacts.extend((int(source_idx), p) for source_idx, p, *_ in ordered_contacts)
                    else:
                        unassigned_positions.extend(pts)
                    shafts.pop(shaft_name, None)
                    shaft_visibility.pop(shaft_name, None)
                    continue

                order = self._distal_first_order_for_positions(
                    np.asarray([contact[1] for contact in ordered_contacts], dtype=float)
                )
                ordered_contacts = [ordered_contacts[int(order_idx)] for order_idx in order]

                for source_idx, contact_pt, t_scalar, perp_dist in ordered_contacts:
                    electrode_idx = len(electrodes)
                    electrodes.append(make_electrode(contact_pt, shaft_name, source_idx))
                    shafts[shaft_name].append(electrode_idx)
                if off_axis_count:
                    self.log(
                        f"{shaft_name}: kept {off_axis_count} contacts beyond "
                        f"{perp_threshold:.1f} vox after shaft grouping."
                    )
            except Exception as e:
                self.log(f"Axis fit failed for {shaft_name}, adding raw points: {e}")
                for source_idx, pos_3d in zip(shaft_indices, shaft_electrodes):
                    electrode_idx = len(electrodes)
                    electrodes.append(make_electrode(pos_3d, shaft_name, int(source_idx)))
                    shafts[shaft_name].append(electrode_idx)

        if np.any(labels == -1):
            unassigned_indices = np.where(labels == -1)[0]
            if source_electrodes is not None:
                unassigned_contacts.extend((int(idx), electrode_array[int(idx)]) for idx in unassigned_indices)
            else:
                unassigned_positions.extend(electrode_array[labels == -1])

        if unassigned_positions or unassigned_contacts:
            shaft_name = "Unassigned"
            shafts[shaft_name] = []
            shaft_visibility[shaft_name] = True
            for source_idx, pos_3d in unassigned_contacts:
                electrode_idx = len(electrodes)
                electrodes.append(make_electrode(pos_3d, shaft_name, source_idx))
                shafts[shaft_name].append(electrode_idx)
            for pos_3d in unassigned_positions:
                electrode_idx = len(electrodes)
                electrodes.append(make_electrode(pos_3d, shaft_name))
                shafts[shaft_name].append(electrode_idx)

        if manual_electrodes and source_electrodes is None:
            for elec in manual_electrodes:
                electrode_idx = len(electrodes)
                electrodes.append(elec)
                shaft_name = elec.shaft or "Manual"
                if shaft_name not in shafts:
                    shafts[shaft_name] = []
                    shaft_visibility[shaft_name] = True
                shafts[shaft_name].append(electrode_idx)

        shafts, shaft_visibility, _ = self._assign_planning_names_to_shafts(
            electrodes,
            shafts,
            shaft_visibility,
        )

        self.before_electrodes = electrodes
        self.electrode_display_space = self.electrode_space
        self.before_shafts = shafts
        self.before_shaft_visibility = shaft_visibility
        self.before_electrode_count_label.setText(f"Electrodes: {len(electrodes)}")
        self.before_all_electrodes_visible = True
        self.before_detected_visible = True
        self._update_shaft_list_with_colors(self.before_shaft_list, self.before_shafts)
        if self.before_shaft_list.count() > 0:
            self.before_shaft_list.setCurrentRow(0)
            self._on_before_shaft_selected()

        self._update_mni_button_state()
        self._schedule_detected_sync(electrodes)

    def _cluster_shafts_ransac(self, positions, dist_thresh=3.0, max_gap=DEFAULT_AXIAL_MAX_GAP, min_inliers=3, n_iters=300):
        """Cluster electrodes into shafts using deterministic line candidates.

        The method name is kept for compatibility with older UI callbacks.
        """
        pts = np.asarray(positions, dtype=float)
        n_points = len(pts)
        if n_points == 0:
            return np.array([], dtype=int)

        labels = cluster_shaft_contacts(
            pts,
            dist_thresh=float(dist_thresh),
            max_gap=float(max_gap),
            min_contacts=int(min_inliers),
        )
        if not np.any(labels >= 0):
            return np.full(n_points, -1, dtype=int)

        return labels

    def _build_brain_mask(self, ct_data, min_fraction=0.02, dilate_iters=3):
        """Create a conservative brain-region mask from CT to suppress extracranial detections."""
        ct = np.asarray(ct_data)
        finite = np.isfinite(ct)
        if finite.sum() == 0:
            return np.ones(ct.shape, dtype=bool)

        ranges = [(-200, 200), (-300, 300), (-400, 400)]
        for soft_min, soft_max in ranges:
            mask = (ct > soft_min) & (ct < soft_max) & finite
            mask = ndi.binary_closing(mask, iterations=2)
            mask = ndi.binary_fill_holes(mask)
            labeled, num = ndi.label(mask)
            if num > 1:
                counts = np.bincount(labeled.ravel())
                counts[0] = 0
                mask = labeled == counts.argmax()
            if dilate_iters > 0:
                mask = ndi.binary_dilation(mask, iterations=dilate_iters)
            if mask.mean() >= min_fraction:
                return mask

        return np.ones(ct.shape, dtype=bool)

    def _safe_nanmax(self, data, axis):
        """Compute nanmax without warning on all-NaN slices."""
        with warnings.catch_warnings():
            warnings.filterwarnings("ignore", message="All-NaN slice encountered")
            return np.nanmax(data, axis=axis)

    def _project_shaft_coords_to_view(self, coords, view_key):
        """Project 3D shaft coordinates into the requested 2D slice/MIP view."""
        pts_3d = np.asarray(coords, dtype=float)
        if pts_3d.size == 0:
            return np.empty((0, 2), dtype=float)
        pts_3d = np.reshape(pts_3d, (-1, 3))
        if view_key == "axial":
            return pts_3d[:, [0, 1]]
        if view_key == "coronal":
            return pts_3d[:, [0, 2]]
        return pts_3d[:, [1, 2]]

    def _view_display_size(self, volume_shape, view_key):
        """Return display width/height in data coordinates for a 2D view."""
        if view_key == "axial":
            return float(volume_shape[0]), float(volume_shape[1])
        if view_key == "coronal":
            return float(volume_shape[0]), float(volume_shape[2])
        return float(volume_shape[1]), float(volume_shape[2])

    @staticmethod
    def _estimate_label_half_size(label_text, volume_shape):
        """Approximate a label's data-space footprint for collision avoidance."""
        max_dim = float(max(volume_shape[:2])) if volume_shape else 256.0
        text_len = len(str(label_text or ""))
        half_width = min(max(10.0, text_len * 1.7), max_dim * 0.22)
        half_height = 5.0
        return half_width, half_height

    @staticmethod
    def _label_box_center(label_pos, half_width, anchor_x=0.5):
        """Return the center of a label box from a text anchor point."""
        pos = np.asarray(label_pos, dtype=float)
        return np.array([
            pos[0] + ((0.5 - float(anchor_x)) * 2.0 * float(half_width)),
            pos[1],
        ])

    @staticmethod
    def _label_collision_count(label_pos, half_width, half_height, points, anchor_x=0.5):
        """Count points that would sit under the estimated label rectangle."""
        if points is None:
            return 0
        pts = np.asarray(points, dtype=float)
        if pts.size == 0:
            return 0
        pts = np.reshape(pts, (-1, 2))
        center = RegistrationDialog._label_box_center(label_pos, half_width, anchor_x)
        dx = np.abs(pts[:, 0] - float(center[0]))
        dy = np.abs(pts[:, 1] - float(center[1]))
        return int(np.sum((dx <= half_width) & (dy <= half_height)))

    def _compute_shaft_label_position(
        self,
        coords,
        view_key,
        volume_shape,
        label_offset=4.0,
        avoid_points=None,
        placed_labels=None,
        label_text="",
    ):
        """Compute a lateral-end label position away from visible contacts."""
        pts = self._project_shaft_coords_to_view(coords, view_key)
        if pts.size == 0:
            return None

        width, height = self._view_display_size(volume_shape, view_key)
        mean_2d = np.mean(pts, axis=0)
        avoid = np.asarray(avoid_points, dtype=float) if avoid_points is not None else pts
        if avoid.size:
            avoid = np.reshape(avoid, (-1, 2))
        base_offset = max(float(label_offset), 8.0)
        half_width, half_height = self._estimate_label_half_size(label_text, volume_shape)
        margin = 3.0
        min_y = margin + half_height
        max_y = max(min_y, height - margin - half_height)

        lateral_sign = -1.0 if float(mean_2d[0]) < (width / 2.0) else 1.0
        endpoint_idx = int(np.argmin(pts[:, 0])) if lateral_sign < 0 else int(np.argmax(pts[:, 0]))
        endpoint = pts[endpoint_idx]
        anchor_x = 1.0 if lateral_sign < 0 else 0.0
        x_anchor = endpoint[0] + lateral_sign * base_offset
        x_anchor = min(max(float(x_anchor), margin), width - margin)

        y_offsets = [0.0, -10.0, 10.0, -20.0, 20.0, -32.0, 32.0, -46.0, 46.0]
        candidates = [
            np.array([
                x_anchor,
                min(max(float(endpoint[1] + y_offset), min_y), max_y),
            ])
            for y_offset in y_offsets
        ]

        placed = placed_labels or []
        best = None
        for candidate in candidates:
            center = self._label_box_center(candidate, half_width, anchor_x)
            if avoid.size:
                dists = np.linalg.norm(avoid - center, axis=1)
                nearest_contact = float(np.min(dists))
                contact_collisions = self._label_collision_count(
                    candidate,
                    half_width,
                    half_height,
                    avoid,
                    anchor_x=anchor_x,
                )
            else:
                nearest_contact = 100.0
                contact_collisions = 0

            label_penalty = 0.0
            for placed_label in placed:
                other_pos = np.asarray(placed_label.get("pos"), dtype=float)
                other_half_width = float(placed_label.get("half_width", half_width))
                other_half_height = float(placed_label.get("half_height", half_height))
                other_anchor_x = float(placed_label.get("anchor_x", 0.5))
                other_center = self._label_box_center(
                    other_pos,
                    other_half_width,
                    other_anchor_x,
                )
                overlap_x = (half_width + other_half_width) - abs(center[0] - other_center[0])
                overlap_y = (half_height + other_half_height) - abs(center[1] - other_center[1])
                if overlap_x > 0 and overlap_y > 0:
                    label_penalty += 150.0 + overlap_x + overlap_y

            score = (
                nearest_contact
                - 200.0 * contact_collisions
                - label_penalty
                - 0.5 * abs(float(candidate[1] - endpoint[1]))
            )
            if best is None or score > best[0]:
                best = (score, candidate)

        if best is None:
            return np.array([x_anchor, min(max(float(endpoint[1]), min_y), max_y)])
        return best[1]

    def _shaft_label_horizontal_alignment(self, label_pos, view_key, slice_shape):
        """Return text alignment that points labels away from lateral contacts."""
        if label_pos is None:
            return "center", 0.5
        try:
            width = float(slice_shape[1])
            x_pos = float(label_pos[0])
        except (TypeError, ValueError, IndexError):
            return "center", 0.5
        if not np.isfinite(width) or width <= 0 or not np.isfinite(x_pos):
            return "center", 0.5
        if x_pos < width / 2.0:
            return "right", 1.0
        return "left", 0.0

    def _find_mip_blob_centroids(self, mip, threshold_min, threshold_max, min_size=3):
        """Find intensity-weighted blob centroids in a thresholded MIP."""
        mip = np.asarray(mip)
        mask = np.isfinite(mip)
        if threshold_min is not None:
            mask &= mip >= threshold_min
        if threshold_max is not None:
            mask &= mip <= threshold_max
        if not mask.any():
            return []

        mask = ndi.binary_opening(mask, iterations=1)
        labeled, num = ndi.label(mask)
        if num == 0:
            return []

        weights = np.nan_to_num(mip, nan=0.0)
        centroids = []
        for label in range(1, num + 1):
            coords = np.argwhere(labeled == label)
            if coords.shape[0] < min_size:
                continue
            try:
                c = ndi.center_of_mass(weights, labeled, label)
            except Exception:
                c = coords.mean(axis=0)
            if np.any(np.isnan(c)):
                continue
            centroids.append((float(c[0]), float(c[1])))
        return centroids

    def _expand_axial_candidates_to_3d(self, ct_data, axial_positions, threshold_min, threshold_max, roi_radius=2):
        """Expand axial MIP blobs into 3D candidates by scanning Z in a local ROI."""
        if ct_data is None or len(axial_positions) == 0:
            return np.empty((0, 3), dtype=int), []

        ct = np.asarray(ct_data)
        candidates = []
        failed_positions = []
        for x_ax, y_ax in axial_positions:
            x_idx = int(round(x_ax))
            y_idx = int(round(y_ax))
            if x_idx < 0 or y_idx < 0 or x_idx >= ct.shape[0] or y_idx >= ct.shape[1]:
                continue

            x0 = max(0, x_idx - roi_radius)
            x1 = min(ct.shape[0] - 1, x_idx + roi_radius)
            y0 = max(0, y_idx - roi_radius)
            y1 = min(ct.shape[1] - 1, y_idx + roi_radius)

            roi = ct[x0:x1 + 1, y0:y1 + 1, :]
            if threshold_min is not None or threshold_max is not None:
                roi = np.where(
                    (roi >= threshold_min) & (roi <= threshold_max),
                    roi,
                    np.nan
                )

            slice_has_voxel = np.isfinite(roi).any(axis=(0, 1))
            if not slice_has_voxel.any():
                failed_positions.append((x_ax, y_ax))
                continue

            labeled, num = ndi.label(slice_has_voxel)
            had_candidate = False
            for label in range(1, num + 1):
                z_idxs = np.where(labeled == label)[0]
                if z_idxs.size == 0:
                    continue
                # Choose the Z with the strongest in-threshold voxel within this segment.
                best_z = None
                best_val = -np.inf
                for z in z_idxs:
                    slice_roi = roi[:, :, z]
                    if not np.isfinite(slice_roi).any():
                        continue
                    val = np.nanmax(slice_roi)
                    if val > best_val:
                        best_val = val
                        best_z = int(z)

                if best_z is None:
                    continue

                slice_roi = roi[:, :, best_z]
                flat_idx = np.nanargmax(slice_roi)
                dx, dy = np.unravel_index(flat_idx, slice_roi.shape)
                x_peak = x0 + dx
                y_peak = y0 + dy

                candidates.append([x_peak, y_peak, best_z])
                had_candidate = True

            if not had_candidate:
                failed_positions.append((x_ax, y_ax))

        if len(candidates) == 0:
            return np.empty((0, 3), dtype=int), failed_positions
        return np.asarray(candidates, dtype=int), failed_positions

    def _save_mip_debug(self, ct_mip_axial, ct_mip_coronal,
                        axial_positions, coronal_positions, threshold_min, threshold_max,
                        mri_data=None, slice_y=None, slice_z=None, ct_alpha=0.6,
                        axial_failed_positions=None):
        """Save thresholded MIP images with detected blob overlays for debugging."""
        debug_root = os.path.join(os.getcwd(), "debug", "electrode_detection")
        os.makedirs(debug_root, exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        prefix = f"ct_{timestamp}"

        def _save_mip(path, mip, positions, title, mri_slice=None, failed_positions=None):
            fig = Figure(figsize=(6, 6), dpi=120)
            ax = fig.add_subplot(1, 1, 1)
            if mri_slice is not None:
                ax.imshow(mri_slice, cmap='gray', origin='lower')
                ax.imshow(mip, cmap='hot', origin='lower', alpha=ct_alpha)
            else:
                ax.imshow(mip, cmap='hot', origin='lower')
            if positions:
                # positions are (row, col) from the MIP array; MIP is transposed for display
                xs = [p[0] for p in positions]
                ys = [p[1] for p in positions]
                ax.scatter(xs, ys, s=30, c='cyan', marker='o', edgecolors='black', linewidths=0.5)
            if failed_positions:
                xs = [p[0] for p in failed_positions]
                ys = [p[1] for p in failed_positions]
                ax.scatter(xs, ys, s=40, c='red', marker='x', linewidths=1.5)
            ax.set_title(f"{title} (thr {threshold_min:.0f}-{threshold_max:.0f})", fontsize=9)
            ax.set_axis_off()
            fig.savefig(path, bbox_inches='tight')
            fig.clear()

        mri_axial = None
        mri_coronal = None
        if mri_data is not None:
            mri_arr = np.asarray(mri_data)
            if slice_z is not None and 0 <= slice_z < mri_arr.shape[2]:
                mri_axial = mri_arr[:, :, slice_z].T
            if slice_y is not None and 0 <= slice_y < mri_arr.shape[1]:
                mri_coronal = mri_arr[:, slice_y, :].T

        _save_mip(
            os.path.join(debug_root, f"{prefix}_axial_mip.png"),
            ct_mip_axial.T,
            axial_positions,
            "Axial MIP",
            mri_slice=mri_axial,
            failed_positions=axial_failed_positions
        )
        _save_mip(
            os.path.join(debug_root, f"{prefix}_coronal_mip.png"),
            ct_mip_coronal.T,
            coronal_positions,
            "Coronal MIP",
            mri_slice=mri_coronal
        )
        self.log(f"Saved MIP debug images to: {debug_root}")

    def _find_nearest_electrode(self, view_idx, x, y, pick_radius=6.0):
        """Find the nearest visible electrode on a given view."""
        electrodes = self.before_electrodes
        shaft_visibility = self.before_shaft_visibility
        show_mip = self.before_show_mip
        slice_x = self.before_slice_x
        slice_y = self.before_slice_y
        slice_z = self.before_slice_z

        if x is None or y is None or slice_x is None or slice_y is None or slice_z is None:
            return None

        best_idx = None
        best_dist2 = pick_radius * pick_radius

        for elec_idx, electrode in enumerate(electrodes):
            if not electrode.visible:
                continue

            shaft_name = electrode.shaft
            if shaft_visibility and not shaft_visibility.get(shaft_name, True):
                continue

            show_electrode = False
            x_display = y_display = None

            if hasattr(electrode, "slice_axial"):  # Auto-detected electrode with slice info
                if view_idx == 0:  # Axial view
                    if show_mip or abs(electrode.slice_axial - slice_z) <= 1:
                        show_electrode = True
                        x_display = electrode.x
                        y_display = electrode.y
                elif view_idx == 1:  # Coronal view
                    if show_mip or abs(electrode.slice_coronal - slice_y) <= 1:
                        show_electrode = True
                        x_display = electrode.x
                        y_display = electrode.z
                elif view_idx == 2:  # Sagittal view
                    if show_mip or abs(electrode.slice_sagittal - slice_x) <= 1:
                        show_electrode = True
                        x_display = electrode.y
                        y_display = electrode.z
            else:  # Manually marked electrode (old format)
                if electrode.view == view_idx:
                    show_electrode = True
                    x_display = electrode.x
                    y_display = electrode.y

            if not show_electrode:
                continue

            dx = x - x_display
            dy = y - y_display
            dist2 = (dx * dx) + (dy * dy)
            if dist2 <= best_dist2:
                best_dist2 = dist2
                best_idx = elec_idx

        return best_idx

    def _get_electrode_display_positions(self, view_idx):
        """Return list of (electrode_idx, x_display, y_display) for a view."""
        electrodes = self.before_electrodes
        shaft_visibility = self.before_shaft_visibility
        show_mip = self.before_show_mip
        slice_x = self.before_slice_x
        slice_y = self.before_slice_y
        slice_z = self.before_slice_z

        if slice_x is None or slice_y is None or slice_z is None:
            return []

        positions = []
        for elec_idx, electrode in enumerate(electrodes):
            if not electrode.visible:
                continue

            shaft_name = electrode.shaft
            if shaft_visibility and not shaft_visibility.get(shaft_name, True):
                continue

            show_electrode = False
            x_display = y_display = None

            if hasattr(electrode, "slice_axial"):
                if view_idx == 0:  # Axial
                    if show_mip or abs(electrode.slice_axial - slice_z) <= 1:
                        show_electrode = True
                        x_display = electrode.x
                        y_display = electrode.y
                elif view_idx == 1:  # Coronal
                    if show_mip or abs(electrode.slice_coronal - slice_y) <= 1:
                        show_electrode = True
                        x_display = electrode.x
                        y_display = electrode.z
                elif view_idx == 2:  # Sagittal
                    if show_mip or abs(electrode.slice_sagittal - slice_x) <= 1:
                        show_electrode = True
                        x_display = electrode.y
                        y_display = electrode.z
            else:
                if electrode.view == view_idx:
                    show_electrode = True
                    x_display = electrode.x
                    y_display = electrode.y

            if show_electrode and x_display is not None and y_display is not None:
                positions.append((elec_idx, x_display, y_display))

        return positions

    def _select_electrodes_in_list(self, elec_indices):
        """Select multiple electrodes in the list, scoped to one shaft."""
        if not elec_indices:
            return

        electrodes = self.before_electrodes
        shaft_list = self.before_shaft_list
        electrode_list = self.before_electrode_list

        by_shaft = {}
        for idx in elec_indices:
            if idx is None or idx >= len(electrodes):
                continue
            shaft_name = electrodes[idx].shaft
            by_shaft.setdefault(shaft_name, []).append(idx)

        if not by_shaft:
            return

        selected_shaft = None
        if shaft_list is not None:
            selected_items = shaft_list.selectedItems()
            if selected_items:
                item = selected_items[0]
                item_name = item.data(Qt.UserRole)
                if not item_name:
                    text = item.text()
                    if text:
                        item_name = text.split(' (')[0]
                if item_name in by_shaft:
                    selected_shaft = item_name

        if selected_shaft is None:
            selected_shaft = max(by_shaft.keys(), key=lambda k: len(by_shaft[k]))

        if shaft_list is not None:
            for i in range(shaft_list.count()):
                item = shaft_list.item(i)
                item_name = item.data(Qt.UserRole)
                if not item_name:
                    text = item.text()
                    if text:
                        item_name = text.split(' (')[0]
                if item_name == selected_shaft:
                    shaft_list.setCurrentItem(item)
                    self._on_before_shaft_selected()
                    break

        if electrode_list is None:
            return

        target_indices = set(by_shaft.get(selected_shaft, []))
        electrode_list.clearSelection()
        for i in range(electrode_list.count()):
            item = electrode_list.item(i)
            if item.data(Qt.UserRole) in target_indices:
                item.setSelected(True)

        if electrode_list.selectedItems():
            electrode_list.scrollToItem(electrode_list.selectedItems()[0])

    def _select_electrode_in_list(self, elec_idx):
        """Select the electrode in the shaft and electrode lists."""
        electrodes = self.before_electrodes
        shaft_list = self.before_shaft_list
        electrode_list = self.before_electrode_list

        if elec_idx is None or elec_idx >= len(electrodes):
            return

        shaft_name = electrodes[elec_idx].shaft
        if shaft_name and shaft_list is not None:
            shaft_list.clearSelection()
            for i in range(shaft_list.count()):
                item = shaft_list.item(i)
                item_name = item.data(Qt.UserRole)
                if not item_name:
                    text = item.text()
                    if text:
                        item_name = text.split(' (')[0]
                if item_name == shaft_name:
                    shaft_list.setCurrentItem(item)
                    item.setSelected(True)
                    self._on_before_shaft_selected()
                    break

        if electrode_list is not None:
            def _select_electrode_item():
                if electrode_list is None:
                    return
                electrode_list.clearSelection()
                for i in range(electrode_list.count()):
                    item = electrode_list.item(i)
                    if item.data(Qt.UserRole) == elec_idx:
                        electrode_list.setCurrentItem(item)
                        item.setSelected(True)
                        electrode_list.scrollToItem(item)
                        break

            QtCore.QTimer.singleShot(0, _select_electrode_item)

    def _on_before_select(self, event):
        """Select electrode from axial view on click."""
        if event.inaxes is None:
            return
        if event.button == 3:
            return
        axes_list = getattr(self.before_canvas, "_view_axes", None)
        if axes_list is None:
            axes_list = [ax for ax in self.before_canvas.figure.get_axes()]
        if event.inaxes not in axes_list:
            return

        view_idx = axes_list.index(event.inaxes)

        elec_idx = self._find_nearest_electrode(view_idx, event.xdata, event.ydata)
        if elec_idx is not None:
            self._select_electrode_in_list(elec_idx)

    def _on_before_pg_click(self, view_key, pos):
        """Handle click event in PyQtGraph view."""
        if pos is None:
            return
        view_idx = self._view_key_to_index(view_key)
        if getattr(self, "before_electrode_checkbox", None) and self.before_electrode_checkbox.isChecked():
            self._add_manual_before_electrode(view_idx, pos.x(), pos.y())
            return
        elec_idx = self._find_nearest_electrode(view_idx, pos.x(), pos.y())
        if elec_idx is not None:
            self._select_electrode_in_list(elec_idx)

    def _on_before_pg_drag_select(self, view_key, rect):
        """Select multiple electrodes within a rectangle (PyQtGraph)."""
        if rect is None:
            return
        if getattr(self, "before_electrode_checkbox", None) and self.before_electrode_checkbox.isChecked():
            return
        view_idx = self._view_key_to_index(view_key)
        positions = self._get_electrode_display_positions(view_idx)
        selected = [
            elec_idx for elec_idx, x_disp, y_disp in positions
            if rect.contains(QtCore.QPointF(x_disp, y_disp))
        ]
        self._select_electrodes_in_list(selected)

    def _on_before_pg_hover(self, view_key, pos):
        """Handle hover for PyQtGraph views."""
        marking_active = bool(
            getattr(self, "before_electrode_checkbox", None)
            and self.before_electrode_checkbox.isChecked()
        )
        if pos is None or marking_active:
            return

        view_idx = self._view_key_to_index(view_key)

        hovered_idx = self._find_nearest_electrode(
            view_idx, pos.x(), pos.y(), pick_radius=6.0
        )
        if hovered_idx != self.before_hover_elec_idx:
            self.before_hover_elec_idx = hovered_idx
            self._update_pg_markers()

    def _on_before_rect_select(self, eclick, erelease, view_idx):
        """Select multiple electrodes within a rectangle (left-drag)."""
        if eclick.xdata is None or eclick.ydata is None or erelease.xdata is None or erelease.ydata is None:
            return
        x0, x1 = sorted([eclick.xdata, erelease.xdata])
        y0, y1 = sorted([eclick.ydata, erelease.ydata])
        if (x1 - x0) < 2 or (y1 - y0) < 2:
            return

        positions = self._get_electrode_display_positions(view_idx)
        selected = [
            elec_idx for elec_idx, x_disp, y_disp in positions
            if x0 <= x_disp <= x1 and y0 <= y_disp <= y1
        ]
        self._select_electrodes_in_list(selected)

    def _add_manual_before_electrode(self, view_idx, x, y):
        """Add a manual electrode for the before view."""
        if x is None or y is None:
            return

        data_shape = None
        if self.before_ct_data is not None:
            data_shape = self.before_ct_data.shape
        elif self.before_mri_data is not None:
            data_shape = self.before_mri_data.shape
        if data_shape is None or len(data_shape) < 3:
            return

        def clamp(val, max_val):
            return max(0, min(int(round(val)), max_val - 1))

        if view_idx == 0:  # Axial
            x_3d = clamp(x, data_shape[0])
            y_3d = clamp(y, data_shape[1])
            z_3d = self._infer_mip_hidden_coordinate(
                view_idx, x_3d, y_3d, self.before_slice_z
            )
        elif view_idx == 1:  # Coronal
            x_3d = clamp(x, data_shape[0])
            z_3d = clamp(y, data_shape[2])
            y_3d = self._infer_mip_hidden_coordinate(
                view_idx, x_3d, z_3d, self.before_slice_y
            )
        else:  # Sagittal
            y_3d = clamp(x, data_shape[1])
            z_3d = clamp(y, data_shape[2])
            x_3d = self._infer_mip_hidden_coordinate(
                view_idx, y_3d, z_3d, self.before_slice_x
            )

        existing_idx = self._find_nearest_electrode(view_idx, x, y)
        if existing_idx is not None:
            self._select_electrode_in_list(existing_idx)
            return

        electrode_idx = len(self.before_electrodes)
        shaft_name = f"Manual_{electrode_idx}"
        new_electrode = Electrode(
            x=x_3d,
            y=y_3d,
            z=z_3d,
            shaft=shaft_name,
            view=view_idx,
            visible=True,
            detected=False,
            space=self._current_electrode_space(),
            slice_axial=z_3d,
            slice_coronal=y_3d,
            slice_sagittal=x_3d,
        )
        self.before_electrodes.append(new_electrode)
        self.electrode_display_space = self._current_electrode_space()

        if shaft_name not in self.before_shafts:
            self.before_shafts[shaft_name] = []
        self.before_shafts[shaft_name].append(electrode_idx)
        self.before_shaft_visibility[shaft_name] = True

        self.before_electrode_count_label.setText(f"Electrodes: {len(self.before_electrodes)}")
        try:
            self.before_group_shafts_btn.setEnabled(True)
        except Exception:
            pass
        self._update_shaft_list_with_colors(self.before_shaft_list, self.before_shafts)
        if self.before_pg_views:
            self._update_pg_markers()
        elif self.before_mri_data is not None:
            self._plot_slices(self.before_mri_data, self.before_ct_data, self.before_canvas,
                              "Unaligned CT Overlaid on MRI",
                              self.before_slice_x, self.before_slice_y, self.before_slice_z)
        self._schedule_detected_sync(self.before_electrodes)
        view_names = ("Axial", "Coronal", "Sagittal")
        source = "MIP" if getattr(self, "before_show_mip", False) else "slice"
        self.log(
            f"Electrode marked at ({x_3d}, {y_3d}, {z_3d}) from {view_names[view_idx]} {source}"
        )

    def _infer_mip_hidden_coordinate(self, view_idx, display_a, display_b, fallback):
        """Infer the hidden 3D axis for a manual click on a MIP view."""
        ct_data = getattr(self, "before_ct_data", None)
        if ct_data is None or not getattr(self, "before_show_mip", False):
            return fallback

        ct = np.asarray(ct_data)
        if ct.ndim != 3:
            return fallback

        def clamp(val, max_val):
            return max(0, min(int(round(val)), max_val - 1))

        use_range_threshold = bool(
            getattr(self, "before_mip_threshold_checkbox", None)
            and self.before_mip_threshold_checkbox.isChecked()
            and getattr(self, "before_threshold_lower", None)
            and getattr(self, "before_threshold_upper", None)
        )

        global_threshold = None
        if not use_range_threshold:
            finite_ct = ct[np.isfinite(ct)]
            if finite_ct.size == 0:
                return fallback
            global_threshold = float(np.quantile(finite_ct, 0.95))

        def filtered_scores(roi, axes):
            roi = np.asarray(roi, dtype=float)
            valid = np.isfinite(roi)
            if use_range_threshold:
                threshold_min = self.before_threshold_lower.value()
                threshold_max = self.before_threshold_upper.value()
                valid &= (roi >= threshold_min) & (roi <= threshold_max)
            else:
                valid &= roi >= global_threshold
            if not np.any(valid):
                return None
            return self._safe_nanmax(np.where(valid, roi, np.nan), axis=axes)

        def best_index(scores, fallback_idx):
            if scores is None:
                return fallback_idx
            finite_scores = np.isfinite(scores)
            if not np.any(finite_scores):
                return fallback_idx
            best_value = np.nanmax(scores)
            if not np.isfinite(best_value):
                return fallback_idx
            best_candidates = np.flatnonzero(scores == best_value)
            if best_candidates.size == 0:
                return fallback_idx
            return int(best_candidates[np.argmin(np.abs(best_candidates - fallback_idx))])

        radius = 2
        if view_idx == 0:  # Axial MIP: clicked x/y, infer z.
            x_idx = clamp(display_a, ct.shape[0])
            y_idx = clamp(display_b, ct.shape[1])
            fallback_idx = clamp(fallback, ct.shape[2])
            x0 = max(0, x_idx - radius)
            x1 = min(ct.shape[0], x_idx + radius + 1)
            y0 = max(0, y_idx - radius)
            y1 = min(ct.shape[1], y_idx + radius + 1)
            scores = filtered_scores(ct[x0:x1, y0:y1, :], axes=(0, 1))
            return best_index(scores, fallback_idx)

        if view_idx == 1:  # Coronal MIP: clicked x/z, infer y.
            x_idx = clamp(display_a, ct.shape[0])
            z_idx = clamp(display_b, ct.shape[2])
            fallback_idx = clamp(fallback, ct.shape[1])
            x0 = max(0, x_idx - radius)
            x1 = min(ct.shape[0], x_idx + radius + 1)
            z0 = max(0, z_idx - radius)
            z1 = min(ct.shape[2], z_idx + radius + 1)
            scores = filtered_scores(ct[x0:x1, :, z0:z1], axes=(0, 2))
            return best_index(scores, fallback_idx)

        # Sagittal MIP: clicked y/z, infer x.
        y_idx = clamp(display_a, ct.shape[1])
        z_idx = clamp(display_b, ct.shape[2])
        fallback_idx = clamp(fallback, ct.shape[0])
        y0 = max(0, y_idx - radius)
        y1 = min(ct.shape[1], y_idx + radius + 1)
        z0 = max(0, z_idx - radius)
        z1 = min(ct.shape[2], z_idx + radius + 1)
        scores = filtered_scores(ct[:, y0:y1, z0:z1], axes=(1, 2))
        return best_index(scores, fallback_idx)

    def _on_before_hover(self, event):
        """Highlight electrode marker on hover."""
        if event.button == 1:
            return
        marking_active = bool(
            getattr(self, "before_electrode_checkbox", None)
            and self.before_electrode_checkbox.isChecked()
        )
        if event.inaxes is None or event.xdata is None or event.ydata is None:
            if not marking_active and self.before_hover_elec_idx is not None:
                self.before_hover_elec_idx = None
                if self.before_mri_data is not None:
                    self._plot_slices(self.before_mri_data, self.before_ct_data, self.before_canvas,
                                      "Unaligned CT Overlaid on MRI",
                                      self.before_slice_x, self.before_slice_y, self.before_slice_z)
            return

        axes_list = getattr(self.before_canvas, "_view_axes", None)
        if axes_list is None:
            axes_list = [ax for ax in self.before_canvas.figure.get_axes()]
        if event.inaxes not in axes_list:
            return

        view_idx = axes_list.index(event.inaxes)
        if marking_active:
            return
        hovered_idx = self._find_nearest_electrode(
            view_idx, event.xdata, event.ydata, pick_radius=6.0
        )

        if hovered_idx != self.before_hover_elec_idx:
            self.before_hover_elec_idx = hovered_idx
            if self.before_mri_data is not None:
                self._plot_slices(self.before_mri_data, self.before_ct_data, self.before_canvas,
                                  "Unaligned CT Overlaid on MRI",
                                  self.before_slice_x, self.before_slice_y, self.before_slice_z)

    def _clear_before_zoom(self):
        """Remove the zoom inset from the before canvas."""
        if self.before_zoom_ax is not None:
            try:
                self.before_zoom_ax.remove()
            except Exception:
                pass
            self.before_zoom_ax = None
            try:
                self.before_canvas.draw()
            except Exception:
                pass
        for zoom_item in self.before_zoom_items.values():
            try:
                zoom_item.setVisible(False)
            except Exception:
                pass

    def _update_before_zoom(self, event, view_idx):
        """Update the zoom inset around the cursor for manual marking."""
        ax = event.inaxes
        if ax is None:
            return
        images = list(getattr(ax, "images", []))
        if not images:
            return

        mri_img = None
        ct_img = None
        for img in images:
            cmap = getattr(img, "get_cmap", lambda: None)()
            cmap_name = getattr(cmap, "name", "")
            if cmap_name == "hot":
                ct_img = img
            else:
                mri_img = img
        if ct_img is None and mri_img is None:
            return

        data_ct = ct_img.get_array() if ct_img is not None else None
        data_mri = mri_img.get_array() if mri_img is not None else None
        base = data_ct if data_ct is not None else data_mri
        if base is None:
            return

        height, width = base.shape[:2]
        cx = int(round(event.xdata))
        cy = int(round(event.ydata))
        if cx < 0 or cy < 0 or cx >= width or cy >= height:
            return

        half = 10
        x0 = max(0, cx - half)
        x1 = min(width, cx + half + 1)
        y0 = max(0, cy - half)
        y1 = min(height, cy + half + 1)

        ax_pos = ax.get_position()
        xfrac = (event.xdata - ax.get_xlim()[0]) / max(ax.get_xlim()[1] - ax.get_xlim()[0], 1e-6)
        yfrac = (event.ydata - ax.get_ylim()[0]) / max(ax.get_ylim()[1] - ax.get_ylim()[0], 1e-6)
        w = 0.18 * ax_pos.width
        h = 0.18 * ax_pos.height
        x0_fig = ax_pos.x0 + xfrac * ax_pos.width - w * 0.5
        y0_fig = ax_pos.y0 + yfrac * ax_pos.height - h * 0.5
        x0_fig = max(ax_pos.x0, min(ax_pos.x1 - w, x0_fig))
        y0_fig = max(ax_pos.y0, min(ax_pos.y1 - h, y0_fig))

        if self.before_zoom_ax is not None and self.before_zoom_ax not in ax.figure.axes:
            self.before_zoom_ax = None
        if self.before_zoom_ax is None:
            self.before_zoom_ax = ax.figure.add_axes([x0_fig, y0_fig, w, h])
        else:
            self.before_zoom_ax.set_position([x0_fig, y0_fig, w, h])
            self.before_zoom_ax.clear()

        if data_mri is not None:
            self.before_zoom_ax.imshow(
                data_mri[y0:y1, x0:x1],
                cmap='gray',
                origin='lower'
            )
        if data_ct is not None:
            self.before_zoom_ax.imshow(
                data_ct[y0:y1, x0:x1],
                cmap='hot',
                origin='lower',
                alpha=0.7
            )
        self.before_zoom_ax.set_xticks([])
        self.before_zoom_ax.set_yticks([])
        self.before_zoom_ax.set_title("Zoom", fontsize=8, pad=2)
        try:
            self.before_canvas.draw()
        except Exception:
            pass

    def _update_before_zoom_pg(self, view_key, pos):
        """Update the zoom inset around the cursor for PyQtGraph."""
        if view_key not in self.before_pg_views:
            return
        mri_slice, ct_slice = self.before_pg_slices.get(view_key, (None, None))
        if mri_slice is None and ct_slice is None:
            return

        base = ct_slice if ct_slice is not None else mri_slice
        height, width = base.shape[:2]
        cx = int(round(pos.x()))
        cy = int(round(pos.y()))
        if cx < 0 or cy < 0 or cx >= width or cy >= height:
            self._clear_before_zoom()
            return

        half = 12
        x0 = max(0, cx - half)
        x1 = min(width, cx + half + 1)
        y0 = max(0, cy - half)
        y1 = min(height, cy + half + 1)

        mri_patch = mri_slice[y0:y1, x0:x1] if mri_slice is not None else None
        ct_patch = ct_slice[y0:y1, x0:x1] if ct_slice is not None else None
        rgba = self._build_zoom_rgba(mri_patch, ct_patch)
        if rgba is None:
            return

        # Overlay electrode markers in the zoom patch.
        view_idx = self._view_key_to_index(view_key)
        spots = self._collect_marker_spots(view_idx)
        for spot in spots:
            x_disp, y_disp = spot["pos"]
            if x0 <= x_disp < x1 and y0 <= y_disp < y1:
                px = int(round(x_disp - x0))
                py = int(round(y_disp - y0))
                self._draw_zoom_marker(rgba, px, py, spot["color"])

        qimg = pg.makeQImage(rgba, alpha=True)
        pixmap = QtGui.QPixmap.fromImage(qimg)

        zoom_item = self.before_zoom_items.get(view_key)
        if zoom_item is None:
            zoom_item = QtWidgets.QGraphicsPixmapItem()
            zoom_item.setZValue(2000)
            self.before_pg_widgets[view_key].scene().addItem(zoom_item)
            self.before_zoom_items[view_key] = zoom_item

        zoom_item.setPixmap(pixmap)
        view_box = self.before_pg_views[view_key]["view"]
        scene_pos = view_box.mapViewToScene(pos)
        offset = QtCore.QPointF(12, -pixmap.height() - 12)
        zoom_item.setPos(scene_pos + offset)
        zoom_item.setVisible(True)

    def _on_before_threshold_changed(self, _value=None):
        """Enable auto-detect when before thresholds change."""
        try:
            self.before_auto_detect_btn.setEnabled(True)
        except Exception:
            pass
        if self.before_show_mip and getattr(self, "before_mip_threshold_checkbox", None):
            if self.before_mip_threshold_checkbox.isChecked() and self.before_mri_data is not None:
                self._plot_slices(self.before_mri_data, self.before_ct_data, self.before_canvas,
                                "Unaligned CT Overlaid on MRI",
                                self.before_slice_x, self.before_slice_y, self.before_slice_z)

    def _toggle_before_electrode_marking(self, state):
        """Toggle electrode marking mode for before view."""
        if state == 2:  # Qt.Checked
            if self.before_hover_elec_idx is not None:
                self.before_hover_elec_idx = None
                if self.before_mri_data is not None:
                    self._plot_slices(self.before_mri_data, self.before_ct_data, self.before_canvas,
                                      "Unaligned CT Overlaid on MRI",
                                      self.before_slice_x, self.before_slice_y, self.before_slice_z)
            self.log("Electrode marking enabled for Before view - click on images to mark electrodes")
        else:
            self.log("Electrode marking disabled for Before view")
    
    def _on_before_click(self, event):
        """Handle click event on before canvas for electrode marking."""
        if not getattr(self, "before_electrode_checkbox", None) or not self.before_electrode_checkbox.isChecked():
            return
        if event.inaxes is None:
            return
        if event.button == 3:
            return
        
        # Determine which view was clicked (0=axial, 1=coronal, 2=sagittal)
        axes_list = getattr(self.before_canvas, "_view_axes", None)
        if axes_list is None:
            axes_list = [ax for ax in self.before_canvas.figure.get_axes()]
        if event.inaxes not in axes_list:
            return
        
        view_idx = axes_list.index(event.inaxes)
        self._add_manual_before_electrode(view_idx, event.xdata, event.ydata)

    def _view_key_to_index(self, view_key):
        """Map view key to index ordering."""
        return {"axial": 0, "coronal": 1, "sagittal": 2}.get(view_key, 0)

    def _pick_pg_electrode(self, view_key, pos):
        """Pick the nearest electrode for dragging in a PyQtGraph view."""
        if pos is None:
            return None
        view_idx = self._view_key_to_index(view_key)
        return self._find_nearest_electrode(
            view_idx, pos.x(), pos.y(), pick_radius=6.0
        )

    def _refresh_detected_positions(self):
        """Rebuild detected position arrays from current electrodes."""
        electrodes = self.before_electrodes
        positions = [
            (float(e.x), float(e.y), float(e.z))
            for e in electrodes
            if getattr(e, "detected", False)
        ]
        arr = np.asarray(positions, dtype=float) if positions else np.empty((0, 3), dtype=float)
        self.before_detected_positions = arr

    def _drag_pg_electrode(self, view_key, elec_idx, pos, finished):
        """Drag an electrode marker in the PyQtGraph view."""
        if pos is None:
            return
        if elec_idx is None or elec_idx < 0 or elec_idx >= len(self.before_electrodes):
            return
        view_idx = self._view_key_to_index(view_key)
        elec = self.before_electrodes[elec_idx]

        data_shape = None
        if self.before_mri_data is not None:
            data_shape = self.before_mri_data.shape
        elif self.before_ct_data is not None:
            data_shape = self.before_ct_data.shape
        if data_shape is None or len(data_shape) < 3:
            return

        def clamp(val, max_val):
            return max(0, min(int(round(val)), max_val - 1))

        use_mip_depth = bool(getattr(self, "before_show_mip", False))
        if view_idx == 0:  # Axial
            x_3d = clamp(pos.x(), data_shape[0])
            y_3d = clamp(pos.y(), data_shape[1])
            z_3d = clamp(elec.z if use_mip_depth else self.before_slice_z, data_shape[2])
        elif view_idx == 1:  # Coronal
            x_3d = clamp(pos.x(), data_shape[0])
            y_3d = clamp(elec.y if use_mip_depth else self.before_slice_y, data_shape[1])
            z_3d = clamp(pos.y(), data_shape[2])
        else:  # Sagittal
            x_3d = clamp(elec.x if use_mip_depth else self.before_slice_x, data_shape[0])
            y_3d = clamp(pos.x(), data_shape[1])
            z_3d = clamp(pos.y(), data_shape[2])

        elec.x = x_3d
        elec.y = y_3d
        elec.z = z_3d
        elec.slice_axial = z_3d
        elec.slice_coronal = y_3d
        elec.slice_sagittal = x_3d

        self._update_pg_markers()

        if finished:
            self._refresh_detected_positions()
            self._schedule_detected_sync(self.before_electrodes)
    
    def _make_ct_rgba(self, ct_slice, alpha, levels=None):
        """Create an RGBA image for CT overlay."""
        if ct_slice is None:
            return None
        data = np.array(ct_slice, copy=True)
        mask = ~np.isfinite(data)
        if mask.all():
            return None
        if levels is None:
            vmin = np.nanmin(data)
            vmax = np.nanmax(data)
        else:
            vmin, vmax = levels
        if vmax <= vmin:
            vmax = vmin + 1.0
        norm = (data - vmin) / (vmax - vmin)
        norm = np.clip(norm, 0.0, 1.0)
        norm[~np.isfinite(norm)] = 0.0
        cmap = self._get_ct_colormap()
        rgba = cmap.map(norm, mode='byte')
        rgba[..., 3] = (rgba[..., 3].astype(float) * alpha).astype(np.uint8)
        rgba[mask, 3] = 0
        return rgba

    def _get_ct_colormap(self):
        """Return a colormap for CT overlays with safe fallbacks."""
        cmap = getattr(self, "_ct_cmap", None)
        if cmap is not None:
            return cmap
        try:
            cmap = pg.colormap.get("hot")
        except Exception:
            try:
                cmap = pg.colormap.getFromMatplotlib("hot")
            except Exception:
                pos = np.array([0.0, 0.4, 0.75, 1.0], dtype=float)
                colors = np.array(
                    [
                        [0, 0, 0, 255],
                        [255, 0, 0, 255],
                        [255, 255, 0, 255],
                        [255, 255, 255, 255],
                    ],
                    dtype=np.ubyte,
                )
                cmap = pg.ColorMap(pos, colors)
        self._ct_cmap = cmap
        return cmap

    def _get_ct_colorbar_rgba(self, width=256, height=6):
        """Build or return a cached CT colorbar RGBA image."""
        key = (int(width), int(height))
        cache = getattr(self, "_ct_colorbar_cache", {})
        if key in cache:
            return cache[key]
        cmap = self._get_ct_colormap()
        ramp = np.linspace(0.0, 1.0, width)
        colors = cmap.map(ramp, mode='byte')
        rgba = np.repeat(colors[np.newaxis, :, :], height, axis=0)
        cache[key] = rgba
        self._ct_colorbar_cache = cache
        return rgba

    @staticmethod
    def _prepare_ct_slice_overlay(ct_slice, air_threshold=-900.0):
        """Mask air/background while keeping the full in-head CT slice visible."""
        data = np.asarray(ct_slice, dtype=np.float32)
        return np.where(data > float(air_threshold), data, np.nan)

    @staticmethod
    def _ct_slice_overlay_levels():
        """Return a stable CT display window for non-MIP slice overlays."""
        return -1000.0, 3000.0

    def _safe_mri_levels(self, mri_slice):
        """Return stable grayscale levels for MRI display."""
        data = np.asarray(mri_slice, dtype=np.float32)
        finite = np.isfinite(data)
        if not finite.any():
            return 0.0, 1.0
        vmin = np.nanpercentile(data[finite], 1.0)
        vmax = np.nanpercentile(data[finite], 99.0)
        if vmax <= vmin:
            vmax = vmin + 1.0
        return float(vmin), float(vmax)

    def _build_zoom_rgba(self, mri_patch, ct_patch):
        """Build zoom RGBA image with MRI base and CT overlay."""
        if mri_patch is None and ct_patch is None:
            return None
        if mri_patch is None:
            base = np.zeros_like(ct_patch, dtype=float)
        else:
            base = np.asarray(mri_patch, dtype=float)
        vmin = np.nanmin(base) if np.isfinite(base).any() else 0.0
        vmax = np.nanmax(base) if np.isfinite(base).any() else 1.0
        if vmax <= vmin:
            vmax = vmin + 1.0
        norm = (base - vmin) / (vmax - vmin)
        norm = np.clip(norm, 0.0, 1.0)
        gray = (norm * 255).astype(np.uint8)
        rgba = np.stack([gray, gray, gray, np.full_like(gray, 255)], axis=-1)

        if ct_patch is not None:
            ct_rgba = self._make_ct_rgba(ct_patch, self.before_ct_alpha)
            if ct_rgba is not None:
                alpha = ct_rgba[..., 3:4].astype(float) / 255.0
                rgba[..., :3] = (rgba[..., :3] * (1.0 - alpha) + ct_rgba[..., :3] * alpha).astype(np.uint8)
        return rgba

    def _draw_zoom_marker(self, rgba, x, y, color):
        """Draw a small marker on the zoom RGBA image."""
        if rgba is None:
            return
        height, width = rgba.shape[:2]
        radius = 2
        r, g, b, _ = color
        for dy in range(-radius, radius + 1):
            for dx in range(-radius, radius + 1):
                if dx * dx + dy * dy > radius * radius:
                    continue
                px = x + dx
                py = y + dy
                if 0 <= px < width and 0 <= py < height:
                    rgba[py, px, 0] = r
                    rgba[py, px, 1] = g
                    rgba[py, px, 2] = b
                    rgba[py, px, 3] = 255

    def _collect_marker_spots(self, view_idx):
        """Collect scatter spots for the given view index."""
        show_mip = self.before_show_mip
        electrodes = self.before_electrodes
        shaft_visibility = self.before_shaft_visibility
        shafts = self.before_shafts

        shaft_colors = ['red', 'blue', 'green', 'orange', 'purple', 'cyan', 'magenta', 'yellow']
        shaft_names = self._sorted_shaft_names(shafts, include_unassigned=True)
        shaft_color_map = {"Unassigned": "#777777"}
        color_idx = 0
        for name in shaft_names:
            if name == "Unassigned":
                continue
            shaft_color_map[name] = shaft_colors[color_idx % len(shaft_colors)]
            color_idx += 1

        selected_shaft = None
        shaft_list = self.before_shaft_list
        if shaft_list is not None:
            selected_items = shaft_list.selectedItems()
            if selected_items:
                selected_shaft = selected_items[0].data(Qt.UserRole)
        if selected_shaft == "Detected":
            only_detected = all(name in ("Detected", "Unassigned") for name in shafts)
            if only_detected:
                selected_shaft = None

        spots = []
        hover_idx = None
        if not (getattr(self, "before_electrode_checkbox", None) and self.before_electrode_checkbox.isChecked()):
            hover_idx = self.before_hover_elec_idx

        slice_x = self.before_slice_x
        slice_y = self.before_slice_y
        slice_z = self.before_slice_z
        for elec_idx, electrode in enumerate(electrodes):
            if not electrode.visible:
                continue
            show_electrode = False
            x_display, y_display = 0, 0
            if hasattr(electrode, "slice_axial"):
                if view_idx == 0:
                    if show_mip or abs(electrode.slice_axial - slice_z) <= 1:
                        show_electrode = True
                        x_display = electrode.x
                        y_display = electrode.y
                elif view_idx == 1:
                    if show_mip or abs(electrode.slice_coronal - slice_y) <= 1:
                        show_electrode = True
                        x_display = electrode.x
                        y_display = electrode.z
                elif view_idx == 2:
                    if show_mip or abs(electrode.slice_sagittal - slice_x) <= 1:
                        show_electrode = True
                        x_display = electrode.y
                        y_display = electrode.z
            else:
                if electrode.view == view_idx:
                    show_electrode = True
                    x_display = electrode.x
                    y_display = electrode.y

            if not show_electrode:
                continue
            shaft_name = electrode.shaft
            if not shaft_visibility.get(shaft_name, True):
                continue

            if shaft_name == "Detected":
                color = '#00aa00'
            else:
                color = shaft_color_map.get(shaft_name, 'red')

            is_selected = selected_shaft is not None and shaft_name == selected_shaft
            marker_style = 'o' if shaft_name == "Detected" else (
                'd' if getattr(electrode, "tissue", "unknown") == "gray" else 'o'
            )
            if hover_idx is not None and elec_idx == hover_idx:
                size = DEFAULT_MARKER_SIZE + 4
                pen = pg.mkPen('black', width=2)
                brush = pg.mkBrush('yellow')
            elif is_selected:
                size = DEFAULT_MARKER_SIZE + 3
                pen = pg.mkPen(color, width=2.5)
                brush = pg.mkBrush(0, 0, 0, 0)
            else:
                size = DEFAULT_MARKER_SIZE
                pen = pg.mkPen(color, width=1.5)
                brush = pg.mkBrush(0, 0, 0, 0)

            spots.append({
                "pos": (x_display, y_display),
                "size": size,
                "pen": pen,
                "brush": brush,
                "symbol": marker_style
            })
        return spots

    def _update_pg_markers(self):
        """Refresh scatter markers for all PyQtGraph views."""
        if not self.before_pg_views:
            return
        for view_key, view in self.before_pg_views.items():
            view_idx = self._view_key_to_index(view_key)
            spots = self._collect_marker_spots(view_idx)
            scatter = view["scatter"]
            scatter.setData(spots, pxMode=False)

    @staticmethod
    def _coerce_positive_int(value):
        """Return value as a positive int when possible."""
        if value is None or isinstance(value, bool):
            return None
        try:
            parsed = int(float(value))
        except (TypeError, ValueError):
            return None
        return parsed if parsed > 0 else None

    def _planned_contact_counts_for_shaft(self, shaft_name, shafts_dict):
        """Collect planned contact counts stored on electrodes in a shaft."""
        if shaft_name == "Unassigned":
            return []
        electrodes = getattr(self, "before_electrodes", [])
        counts = []
        for electrode_idx in shafts_dict.get(shaft_name, []):
            try:
                electrode_idx = int(electrode_idx)
            except (TypeError, ValueError):
                continue
            if electrode_idx is None or electrode_idx < 0 or electrode_idx >= len(electrodes):
                continue
            count = self._coerce_positive_int(electrodes[electrode_idx].get("planned_contacts"))
            if count is not None:
                counts.append(count)
        return counts

    def _shaft_contact_label_status(self, shaft_name, shafts_dict):
        """Build shaft list text and mismatch tooltip from worksheet metadata."""
        detected_count = len(shafts_dict.get(shaft_name, []))
        default_text = f"{shaft_name} ({detected_count} electrodes)"
        planned_counts = self._planned_contact_counts_for_shaft(shaft_name, shafts_dict)
        if not planned_counts:
            return default_text, None, False

        unique_counts = sorted(set(planned_counts))
        if len(unique_counts) > 1:
            planned_text = ", ".join(str(count) for count in unique_counts)
            tooltip = (
                f"{shaft_name} contains electrodes assigned to multiple planned contact counts "
                f"({planned_text}); detected {detected_count} contacts."
            )
            return f"{shaft_name} ({detected_count}/mixed contacts)", tooltip, True

        planned_count = unique_counts[0]
        if detected_count == planned_count:
            return default_text, None, False

        tooltip = (
            f"Detected {detected_count} contacts for {shaft_name}; "
            f"planning worksheet expects {planned_count}."
        )
        return f"{shaft_name} ({detected_count}/{planned_count} contacts)", tooltip, True

    def _shaft_trajectory_label(self, shaft_name, indices=None):
        """Return a stable worksheet letter label for a shaft when available."""
        if shaft_name == "Unassigned":
            return None
        if indices is None:
            indices = (getattr(self, "before_shafts", {}) or {}).get(shaft_name, [])
        electrodes = getattr(self, "before_electrodes", [])
        labels = []
        for electrode_idx in indices or []:
            try:
                electrode_idx = int(electrode_idx)
            except (TypeError, ValueError):
                continue
            if electrode_idx < 0 or electrode_idx >= len(electrodes):
                continue
            label = str(electrodes[electrode_idx].get("trajectory_label") or "").strip()
            if label and label not in labels:
                labels.append(label)
        if len(labels) == 1:
            return labels[0]
        return None

    def _shaft_display_label(self, shaft_name, indices=None):
        """Return shaft label text with worksheet letter label appended."""
        label = self._shaft_trajectory_label(shaft_name, indices)
        if not label:
            return shaft_name
        if self._normalize_planning_name(label) in self._normalize_planning_name(shaft_name):
            return shaft_name
        return f"{shaft_name} ({label})"
    
    def _update_shaft_list_with_colors(self, shaft_list, shafts_dict):
        """Update shaft list with colored text and embedded action icons."""
        shaft_list.setUpdatesEnabled(False)
        blocker = QtCore.QSignalBlocker(shaft_list)
        shaft_list.clear()
        shaft_colors = ['red', 'blue', 'green', 'orange', 'purple', 'cyan', 'magenta', 'yellow']
        shaft_names = self._sorted_shaft_names(shafts_dict, include_unassigned=True)
        color_idx = 0
        for shaft_name in shaft_names:
            # Create list item
            item = QtWidgets.QListWidgetItem()
            item.setFlags(Qt.ItemIsSelectable | Qt.ItemIsEnabled)
            # Store original shaft name for reference
            item.setData(Qt.UserRole, shaft_name)
            
            # Create custom widget with text and icons
            if shaft_name == "Unassigned":
                color = "#777777"
            else:
                color = shaft_colors[color_idx % len(shaft_colors)]
                color_idx += 1
            label_text, tooltip, has_contact_warning = self._shaft_contact_label_status(shaft_name, shafts_dict)
            widget = ListItemWidget(label_text, color=color)
            widget.edit_btn.setToolTip("Edit name/trajectory")
            widget.set_hidden_state(not self.before_shaft_visibility.get(shaft_name, True))
            if has_contact_warning:
                widget.text_label.setStyleSheet(f"color: {color}; font-weight: bold; text-decoration: underline;")
            if tooltip:
                item.setToolTip(tooltip)
                widget.setToolTip(tooltip)
                widget.text_label.setToolTip(tooltip)
            
            # Connect widget signals to handlers
            widget.edit_clicked.connect(
                lambda w, sl=shaft_list: self._on_shaft_action_clicked(sl, self._edit_shaft_from_widget, w)
            )
            widget.hide_clicked.connect(
                lambda w, sl=shaft_list: self._on_shaft_action_clicked(sl, self._toggle_shaft_visibility_from_widget, w)
            )
            widget.delete_clicked.connect(
                lambda w, sl=shaft_list: self._on_shaft_action_clicked(sl, self._delete_shaft_from_widget, w)
            )
            
            # Add item and set custom widget
            shaft_list.addItem(item)
            shaft_list.setItemWidget(item, widget)
        self._update_groupbox_counts_for_list(shaft_list, shafts_dict)
        del blocker
        shaft_list.setUpdatesEnabled(True)

    def _on_shaft_action_clicked(self, shaft_list, handler, widget):
        """Suppress toggle behavior when clicking list-item action buttons."""
        if shaft_list == self.before_shaft_list:
            self._before_shaft_action_click = True
        handler(widget, shaft_list)

    def _consume_shaft_action_flag(self, shaft_list):
        if shaft_list == self.before_shaft_list:
            flag = self._before_shaft_action_click
            self._before_shaft_action_click = False
            return flag
        return False

    def _sync_selected_shaft_to_controller(self, shaft_name):
        """Highlight a selected shaft in the 3D viewer if available."""
        if self.controller is None:
            return
        self.controller.detected_selected_shaft = shaft_name
        self.controller.detected_electrodes = self._apply_shaft_visibility_to_electrodes(
            getattr(self.controller, "detected_electrodes", None) or self.before_electrodes
        )
        if hasattr(self.controller, "_display_detected_electrodes"):
            self.controller._display_detected_electrodes(
                self.controller.detected_electrodes,
                self.controller.detected_electrode_space,
            )

    def _get_shaft_color_map(self, shafts_dict):
        """Return mapping of shaft name -> color."""
        shaft_colors = ['red', 'blue', 'green', 'orange', 'purple', 'cyan', 'magenta', 'yellow']
        names = self._sorted_shaft_names(shafts_dict, include_unassigned=False)
        color_map = {}
        color_map["Unassigned"] = "#777777"
        for idx, name in enumerate(names):
            if name == "Unassigned":
                continue
            color_map[name] = shaft_colors[idx % len(shaft_colors)]
        return color_map

    def _update_groupbox_counts_for_list(self, shaft_list, shafts_dict):
        """Update shaft/electrode group titles with counts."""
        shaft_count = len([name for name in shafts_dict.keys() if name != "Unassigned"])
        if shaft_list == getattr(self, "before_shaft_list", None):
            total_electrodes = len(getattr(self, "before_electrodes", []))
            if getattr(self, "before_shaft_group", None) is not None:
                self.before_shaft_group.setTitle(f"Shafts ({shaft_count})")
            if getattr(self, "before_electrode_group", None) is not None:
                self.before_electrode_group.setTitle(f"Electrodes ({total_electrodes})")

    def _sorted_shaft_names(self, shafts_dict, include_unassigned=True):
        """Sort shaft names by numeric suffix when using S_# convention."""
        names = list(shafts_dict.keys())
        unassigned = [name for name in names if name == "Unassigned"]
        others = [name for name in names if name != "Unassigned"]
        planned_order = getattr(self, "planning_shaft_order", {}) or {}
        planned = []
        numeric = []
        alpha = []
        for name in others:
            if name in planned_order:
                planned.append((planned_order[name], name))
                continue
            match = re.match(r"^[Ss][-_]?(\d+)$", name)
            if match:
                numeric.append((int(match.group(1)), name))
            else:
                alpha.append(name)
        planned.sort(key=lambda x: (x[0], x[1]))
        numeric.sort(key=lambda x: x[0])
        alpha.sort()
        ordered = [name for _, name in planned] + [name for _, name in numeric] + alpha
        if include_unassigned and unassigned:
            ordered.append("Unassigned")
        return ordered

    def _prompt_shaft_choice(self, title, current_shaft, shafts_dict, include_new=False):
        """Show a colored shaft dropdown and return (choice, ok)."""
        choices = self._sorted_shaft_names(shafts_dict, include_unassigned=True)
        unused_planned = self._unused_planning_trajectory_choices(shafts_dict) if include_new else []
        unused_planned_names = [name for name, _trajectory in unused_planned]
        if unused_planned_names:
            choices = choices + unused_planned_names
        if include_new:
            choices = choices + ["<New Shaft>"]

        dialog = QDialog(self)
        dialog.setWindowTitle(title)
        layout = QVBoxLayout(dialog)
        label = QLabel("Select target shaft:")
        combo = QtWidgets.QComboBox()
        combo.setEditable(True)
        combo.setInsertPolicy(QtWidgets.QComboBox.NoInsert)
        line_edit = combo.lineEdit()
        if line_edit is not None:
            line_edit.setReadOnly(True)

        model = QtGui.QStandardItemModel()
        color_map = self._get_shaft_color_map(shafts_dict)
        unused_planned_set = set(unused_planned_names)
        for name in choices:
            item = QtGui.QStandardItem(name)
            if name in unused_planned_set:
                item.setForeground(QtGui.QBrush(QtGui.QColor("#0b7a3b")))
                item.setBackground(QtGui.QBrush(QtGui.QColor("#e8f5ec")))
                font = item.font()
                font.setBold(True)
                item.setFont(font)
                item.setToolTip("Unused trajectory from the SEEG planning worksheet")
            elif name == "<New Shaft>":
                item.setForeground(QtGui.QBrush(QtGui.QColor("#333333")))
            else:
                color = color_map.get(name, "#333333")
                item.setForeground(QtGui.QBrush(QtGui.QColor(color)))
            model.appendRow(item)
        combo.setModel(model)

        def _update_combo_color():
            name = combo.currentText()
            if name in unused_planned_set:
                color = "#0b7a3b"
            elif name == "<New Shaft>":
                color = "#333333"
            else:
                color = color_map.get(name, "#333333")
            if line_edit is not None:
                weight = "bold" if name in unused_planned_set else "normal"
                line_edit.setStyleSheet(f"color: {color}; font-weight: {weight};")

        if current_shaft in choices:
            combo.setCurrentIndex(choices.index(current_shaft))
        elif unused_planned_names:
            combo.setCurrentIndex(choices.index(unused_planned_names[0]))
        _update_combo_color()
        combo.currentIndexChanged.connect(lambda _idx: _update_combo_color())
        layout.addWidget(label)
        layout.addWidget(combo)

        buttons = QtWidgets.QDialogButtonBox(QtWidgets.QDialogButtonBox.Ok | QtWidgets.QDialogButtonBox.Cancel)
        buttons.accepted.connect(dialog.accept)
        buttons.rejected.connect(dialog.reject)
        layout.addWidget(buttons)

        ok = dialog.exec_() == QDialog.Accepted
        return (combo.currentText(), ok)

    def _select_shaft_in_list(self, shaft_name, *, redraw_2d=True, sync_3d=True):
        """Select a shaft item by stored shaft name."""
        shaft_list = self.before_shaft_list
        if shaft_list is None:
            return
        blocker = QtCore.QSignalBlocker(shaft_list)
        shaft_list.clearSelection()
        target_item = None
        for i in range(shaft_list.count()):
            item = shaft_list.item(i)
            item_name = item.data(Qt.UserRole)
            if not item_name:
                text = item.text()
                if text:
                    item_name = text.split(' (')[0]
            if item_name == shaft_name:
                target_item = item
                shaft_list.setCurrentItem(item)
                item.setSelected(True)
                break
        del blocker
        if target_item is not None:
            shaft_list.scrollToItem(target_item)
            self._on_before_shaft_selected(redraw_2d=redraw_2d, sync_3d=sync_3d)

    def _refresh_before_view_after_shaft_metadata_change(self):
        """Refresh only the parts of the before view affected by shaft metadata."""
        if self.before_mri_data is None:
            return
        if self.before_pg_views:
            marking_active = bool(
                getattr(self, "before_electrode_checkbox", None)
                and self.before_electrode_checkbox.isChecked()
            )
            labels_need_redraw = bool(
                getattr(self, "before_show_mip", False)
                and getattr(self, "before_show_shaft_labels", False)
                and not marking_active
            )
            if labels_need_redraw:
                self._plot_slices(self.before_mri_data, self.before_ct_data, self.before_canvas,
                                "Unaligned CT Overlaid on MRI",
                                self.before_slice_x, self.before_slice_y, self.before_slice_z)
            else:
                self._update_pg_markers()
            return

        self._plot_slices(self.before_mri_data, self.before_ct_data, self.before_canvas,
                        "Unaligned CT Overlaid on MRI",
                        self.before_slice_x, self.before_slice_y, self.before_slice_z)

    def _current_trajectory_key_for_shaft(self, shaft_name, indices):
        """Return the planning trajectory key currently assigned to a shaft."""
        trajectory = self._trajectory_for_existing_shaft(shaft_name, indices)
        return self._trajectory_primary_key(trajectory) if trajectory is not None else None

    def _suggested_name_for_trajectory(self, trajectory, current_name):
        """Return a unique shaft name for a trajectory, preserving current name as available."""
        if trajectory is None:
            return current_name
        used_names = set(getattr(self, "before_shafts", {}) or {})
        used_names.discard(current_name)
        return self._unique_planned_shaft_name(
            self._trajectory_base_name(trajectory),
            trajectory.get("label"),
            used_names,
        )

    def _next_generated_shaft_name(self, used_names=None):
        """Return an unused generated shaft name."""
        used_names = set(used_names or set())
        max_existing = -1
        for name in used_names:
            match = re.match(r"^[Ss][-_]?(\d+)$", str(name or ""))
            if match:
                max_existing = max(max_existing, int(match.group(1)))
        counter = max(int(getattr(self, "shaft_counter", 0) or 0), max_existing + 1)
        while f"S_{counter}" in used_names:
            counter += 1
        self.shaft_counter = counter + 1
        return f"S_{counter}"

    def _replacement_trajectory_for_displaced_shaft(
        self,
        shaft_name,
        indices,
        used_keys,
        preferred_key=None,
    ):
        """Pick a conservative unused trajectory for a displaced duplicate assignment."""
        trajectories = self.planning_trajectories or self._refresh_planning_trajectories()
        used_keys = set(used_keys or set())

        if preferred_key and preferred_key not in used_keys:
            preferred = self._trajectory_by_primary_key(preferred_key)
            if preferred is not None:
                return preferred

        old_trajectory = self._trajectory_for_existing_shaft(shaft_name, indices)
        old_hemi = self._trajectory_hemisphere(old_trajectory) if old_trajectory is not None else None
        if old_hemi is None:
            old_hemi = self._shaft_hemisphere(getattr(self, "before_electrodes", []), indices)
        detected_count = len(indices or [])

        candidates = []
        for trajectory in trajectories:
            key = self._trajectory_primary_key(trajectory)
            if not key or key in used_keys:
                continue
            if old_hemi is not None and self._trajectory_hemisphere(trajectory) != old_hemi:
                continue
            candidates.append(trajectory)

        exact_count = [
            trajectory for trajectory in candidates
            if self._coerce_positive_int(trajectory.get("contacts")) == detected_count
        ]
        if len(exact_count) == 1:
            return exact_count[0]
        return None

    def _prompt_shaft_edit(self, shaft_name):
        """Show a shaft edit dialog with rename and trajectory reassignment controls."""
        indices = list((getattr(self, "before_shafts", {}) or {}).get(shaft_name, []))
        trajectories = self._refresh_planning_trajectories(log_missing=True)
        current_key = self._current_trajectory_key_for_shaft(shaft_name, indices)
        used_keys = self._used_planning_trajectory_keys(getattr(self, "before_shafts", {}))

        dialog = QDialog(self)
        dialog.setWindowTitle("Edit Shaft")
        layout = QVBoxLayout(dialog)
        form = QGridLayout()

        name_edit = QtWidgets.QLineEdit(shaft_name)
        form.addWidget(QLabel("Shaft name:"), 0, 0)
        form.addWidget(name_edit, 0, 1)

        trajectory_combo = QtWidgets.QComboBox()
        trajectory_combo.addItem("<No planning trajectory>", None)
        trajectory_by_key = {}
        for trajectory in trajectories:
            primary_key = self._trajectory_primary_key(trajectory)
            if not primary_key:
                continue
            trajectory_by_key[primary_key] = trajectory
            text = self._trajectory_display_name(trajectory)
            if primary_key in used_keys and primary_key != current_key:
                text = f"{text} (used)"
            trajectory_combo.addItem(text, primary_key)

        if trajectory_combo.count() == 1 and not trajectories:
            trajectory_combo.setItemText(0, "No planning worksheet trajectories loaded")
            trajectory_combo.setEnabled(False)

        if current_key:
            for combo_idx in range(trajectory_combo.count()):
                if trajectory_combo.itemData(combo_idx) == current_key:
                    trajectory_combo.setCurrentIndex(combo_idx)
                    break

        user_edited_name = {"value": False}
        last_suggested_name = {"value": shaft_name}

        def _mark_name_edited(_text):
            user_edited_name["value"] = True

        def _on_trajectory_changed(_idx):
            if user_edited_name["value"]:
                return
            primary_key = trajectory_combo.currentData()
            trajectory = trajectory_by_key.get(primary_key)
            suggested = self._suggested_name_for_trajectory(trajectory, shaft_name)
            if name_edit.text() in (shaft_name, last_suggested_name["value"]):
                name_edit.setText(suggested)
                last_suggested_name["value"] = suggested

        name_edit.textEdited.connect(_mark_name_edited)
        trajectory_combo.currentIndexChanged.connect(_on_trajectory_changed)

        form.addWidget(QLabel("Planning trajectory:"), 1, 0)
        form.addWidget(trajectory_combo, 1, 1)
        layout.addLayout(form)

        buttons = QtWidgets.QDialogButtonBox(
            QtWidgets.QDialogButtonBox.Ok | QtWidgets.QDialogButtonBox.Cancel
        )
        buttons.accepted.connect(dialog.accept)
        buttons.rejected.connect(dialog.reject)
        layout.addWidget(buttons)

        if dialog.exec_() != QDialog.Accepted:
            return None

        return {
            "name": name_edit.text().strip(),
            "trajectory_key": trajectory_combo.currentData(),
            "current_trajectory_key": current_key,
            "name_edited": user_edited_name["value"],
        }

    def _apply_shaft_edit(self, shaft_name, new_name, trajectory_key, current_trajectory_key, name_edited=False):
        """Apply shaft rename and optional planning trajectory reassignment."""
        shafts_dict = self.before_shafts
        visibility_dict = self.before_shaft_visibility
        electrodes = self.before_electrodes
        if shaft_name not in shafts_dict:
            self.log(f"Shaft '{shaft_name}' no longer exists.")
            return

        trajectory = self._trajectory_by_primary_key(trajectory_key) if trajectory_key else None
        trajectory_changed = trajectory_key != current_trajectory_key
        new_name = str(new_name or "").strip()
        shaft_name_trajectory = self._planning_trajectory_for_name(shaft_name)
        shaft_name_trajectory_key = (
            self._trajectory_primary_key(shaft_name_trajectory)
            if shaft_name_trajectory is not None
            else None
        )
        if (
            trajectory is not None
            and not name_edited
            and (trajectory_changed or shaft_name_trajectory_key != trajectory_key)
        ):
            new_name = self._suggested_name_for_trajectory(trajectory, shaft_name)
        if not new_name:
            QMessageBox.information(self, "Edit Shaft", "Shaft name cannot be blank.")
            return

        existing_trajectory_keys = {}
        for existing_name, existing_indices in shafts_dict.items():
            existing_trajectory = self._trajectory_for_existing_shaft(existing_name, existing_indices)
            existing_key = self._trajectory_primary_key(existing_trajectory) if existing_trajectory is not None else None
            if existing_key:
                existing_trajectory_keys[existing_name] = existing_key

        conflict_names = []
        if trajectory_key:
            conflict_names = [
                existing_name
                for existing_name, existing_key in existing_trajectory_keys.items()
                if existing_name != shaft_name and existing_key == trajectory_key
            ]

        if (
            trajectory is not None
            and not name_edited
            and (trajectory_changed or shaft_name_trajectory_key != trajectory_key)
        ):
            used_names = set(shafts_dict.keys())
            used_names.discard(shaft_name)
            for conflict_name in conflict_names:
                used_names.discard(conflict_name)
            new_name = self._unique_planned_shaft_name(
                self._trajectory_base_name(trajectory),
                trajectory.get("label"),
                used_names,
            )

        allowed_name_collisions = set(conflict_names)
        if new_name != shaft_name and new_name in shafts_dict and new_name not in allowed_name_collisions:
            self.log(f"Shaft name '{new_name}' already exists!")
            return

        original_shaft_name = shaft_name
        target_old_order = self.planning_shaft_order.get(shaft_name)
        old_indices = list(shafts_dict.get(shaft_name, []))
        target_name = new_name
        displacement_messages = []

        def _rename_existing_shaft(old_name, replacement_name):
            if old_name == replacement_name:
                return old_name
            indices = list(shafts_dict.pop(old_name, []))
            shafts_dict[replacement_name] = indices
            if old_name in visibility_dict:
                visibility_dict[replacement_name] = visibility_dict.pop(old_name)
            elif replacement_name not in visibility_dict:
                visibility_dict[replacement_name] = True
            self.planning_shaft_order.pop(old_name, None)
            for idx in indices:
                if idx < len(electrodes):
                    electrodes[idx].shaft = replacement_name
            return replacement_name

        replacement_keys = set()
        for conflict_name in conflict_names:
            if conflict_name not in shafts_dict:
                continue
            conflict_indices = list(shafts_dict.get(conflict_name, []))
            used_keys = {trajectory_key}
            for existing_name, existing_indices in shafts_dict.items():
                if existing_name in {shaft_name, conflict_name}:
                    continue
                existing_trajectory = self._trajectory_for_existing_shaft(existing_name, existing_indices)
                existing_key = self._trajectory_primary_key(existing_trajectory) if existing_trajectory is not None else None
                if existing_key:
                    used_keys.add(existing_key)
            used_keys.update(replacement_keys)

            replacement = self._replacement_trajectory_for_displaced_shaft(
                conflict_name,
                conflict_indices,
                used_keys,
                preferred_key=current_trajectory_key,
            )

            used_names = set(shafts_dict.keys())
            used_names.discard(conflict_name)
            if target_name != shaft_name:
                used_names.discard(shaft_name)
            used_names.add(target_name)

            if replacement is not None:
                replacement_key = self._trajectory_primary_key(replacement)
                replacement_name = self._unique_planned_shaft_name(
                    self._trajectory_base_name(replacement),
                    replacement.get("label"),
                    used_names,
                )
                if replacement_name == shaft_name and target_name != shaft_name:
                    temp_name = self._next_generated_shaft_name(set(shafts_dict.keys()) | {target_name})
                    _rename_existing_shaft(shaft_name, temp_name)
                    shaft_name = temp_name
                replacement_name = _rename_existing_shaft(conflict_name, replacement_name)
                self._apply_planning_trajectory_to_indices(
                    replacement_name,
                    replacement,
                    shafts_dict.get(replacement_name, conflict_indices),
                    "manual_conflict_reassignment",
                )
                if replacement_key:
                    replacement_keys.add(replacement_key)
                displacement_messages.append(
                    f"moved previous '{self._trajectory_display_name(trajectory)}' assignment on '{conflict_name}' "
                    f"to '{self._trajectory_display_name(replacement)}'"
                )
            else:
                replacement_name = self._next_generated_shaft_name(used_names)
                replacement_name = _rename_existing_shaft(conflict_name, replacement_name)
                self._clear_planning_metadata_for_indices(shafts_dict.get(replacement_name, conflict_indices))
                self.planning_shaft_order.pop(replacement_name, None)
                displacement_messages.append(
                    f"cleared previous '{self._trajectory_display_name(trajectory)}' assignment on '{conflict_name}'"
                )

        renamed = target_name != shaft_name

        if renamed:
            old_order = self.planning_shaft_order.get(shaft_name)
            _rename_existing_shaft(shaft_name, target_name)
            if trajectory_key is None and current_trajectory_key is None:
                preserved_order = old_order if old_order is not None else target_old_order
                if preserved_order is not None:
                    self.planning_shaft_order[target_name] = preserved_order

        target_indices = list(shafts_dict.get(target_name, old_indices))

        if trajectory is not None:
            self._apply_planning_trajectory_to_indices(
                target_name,
                trajectory,
                target_indices,
                "manual_shaft_edit",
            )
        elif current_trajectory_key is not None:
            self._clear_planning_metadata_for_indices(target_indices)
            self.planning_shaft_order.pop(target_name, None)

        self._update_shaft_list_with_colors(self.before_shaft_list, shafts_dict)
        self._select_shaft_in_list(target_name, redraw_2d=False, sync_3d=False)
        self._refresh_before_view_after_shaft_metadata_change()

        if self.controller is not None and hasattr(self.controller, "_display_detected_shafts"):
            self.controller.detected_selected_shaft = target_name
            self._force_3d_shaft_refresh(electrodes)
        else:
            self._schedule_detected_sync(electrodes)

        messages = []
        if target_name != original_shaft_name:
            messages.append(f"renamed from '{original_shaft_name}' to '{target_name}'")
        if trajectory is not None:
            messages.append(f"assigned to trajectory '{self._trajectory_display_name(trajectory)}'")
        elif trajectory_changed:
            messages.append("cleared planning trajectory assignment")
        messages.extend(displacement_messages)
        if messages:
            self.log(f"Shaft '{target_name}' " + " and ".join(messages) + ".")

    def _apply_bulk_shaft_trajectory_assignments(self, shaft_assignments, current_keys):
        """Apply detected shaft -> planning trajectory assignments as one batch."""
        shafts_dict = self.before_shafts
        visibility_dict = self.before_shaft_visibility
        electrodes = self.before_electrodes
        changed = [
            (shaft_name, trajectory_key)
            for shaft_name, trajectory_key in shaft_assignments.items()
            if shaft_name in shafts_dict and trajectory_key != current_keys.get(shaft_name)
        ]
        if not changed:
            return 0

        temp_names = {}
        used_names = set(shafts_dict.keys())

        def _rename_shaft(old_name, new_name):
            if old_name == new_name:
                return new_name
            indices = list(shafts_dict.pop(old_name, []))
            shafts_dict[new_name] = indices
            if old_name in visibility_dict:
                visibility_dict[new_name] = visibility_dict.pop(old_name)
            elif new_name not in visibility_dict:
                visibility_dict[new_name] = True
            self.planning_shaft_order.pop(old_name, None)
            for electrode_idx in indices:
                if 0 <= electrode_idx < len(electrodes):
                    electrodes[electrode_idx].shaft = new_name
            return new_name

        for shaft_name, _trajectory_key in changed:
            if shaft_name not in shafts_dict:
                continue
            temp_name = self._next_generated_shaft_name(used_names)
            used_names.add(temp_name)
            indices = self._order_electrode_indices_distal_first(
                electrodes,
                shafts_dict.get(shaft_name, []),
            )
            shafts_dict[shaft_name] = list(indices)
            self._clear_planning_metadata_for_indices(indices)
            temp_names[shaft_name] = _rename_shaft(shaft_name, temp_name)
            used_names.discard(shaft_name)

        applied = 0
        target_trajectory_keys = {
            trajectory_key
            for _shaft_name, trajectory_key in changed
            if trajectory_key
        }
        used_names = set(shafts_dict.keys())
        for original_name, trajectory_key in changed:
            temp_name = temp_names.get(original_name)
            if not temp_name or temp_name not in shafts_dict:
                continue
            indices = self._order_electrode_indices_distal_first(
                electrodes,
                shafts_dict.get(temp_name, []),
            )
            used_names.discard(temp_name)
            trajectory = self._trajectory_by_primary_key(trajectory_key) if trajectory_key else None

            if trajectory is not None:
                target_name = self._unique_planned_shaft_name(
                    self._trajectory_base_name(trajectory),
                    trajectory.get("label"),
                    used_names,
                )
            elif (
                original_name
                and original_name not in used_names
                and current_keys.get(original_name) not in target_trajectory_keys
            ):
                target_name = original_name
            else:
                target_name = self._next_generated_shaft_name(used_names)

            target_name = _rename_shaft(temp_name, target_name)
            used_names.add(target_name)
            shafts_dict[target_name] = list(indices)
            if trajectory is not None:
                self._apply_planning_trajectory_to_indices(
                    target_name,
                    trajectory,
                    indices,
                    "bulk_map_trajectory_dialog",
                )
            else:
                self._clear_planning_metadata_for_indices(indices)
                self.planning_shaft_order.pop(target_name, None)
            applied += 1

        self.before_shafts = self._order_shaft_dict_distal_first(electrodes, shafts_dict)
        return applied

    def _open_map_trajectory_dialog(self):
        """Open UI for matching .map shaft prefixes to worksheet trajectories."""
        trajectories = self._refresh_planning_trajectories(log_missing=True)
        map_shafts = self._refresh_electrode_map(log_missing=True)
        shaft_rows = [
            (name, list(indices))
            for name, indices in (getattr(self, "before_shafts", {}) or {}).items()
            if name != "Unassigned" and not self._is_micro_like_label(name)
        ]
        shaft_sort_names = self._sorted_shaft_names(
            {name: indices for name, indices in shaft_rows},
            include_unassigned=False,
        )
        shaft_rows_by_name = {name: indices for name, indices in shaft_rows}
        shaft_rows = [(name, shaft_rows_by_name[name]) for name in shaft_sort_names]
        if not trajectories:
            QMessageBox.information(
                self,
                "Map Trajectories",
                "No planning worksheet trajectories are loaded.",
            )
            return
        micro_map_shafts = getattr(self, "micro_map_shafts", {}) or {}
        if not map_shafts and not shaft_rows and not micro_map_shafts:
            QMessageBox.information(
                self,
                "Map Trajectories",
                "No usable electrode .map shafts or detected shafts were found.",
            )
            return

        self._infer_map_trajectory_assignments()
        rows = sorted(map_shafts.values(), key=lambda item: item.get("order", 0))
        micro_rows = sorted(
            micro_map_shafts.values(),
            key=lambda item: item.get("order", 0),
        )
        dialog = QDialog(self)
        dialog.setWindowTitle("Map Trajectories")
        dialog.resize(920, 640)
        layout = QVBoxLayout(dialog)

        trajectory_options = [
            (self._trajectory_display_name(trajectory), self._trajectory_primary_key(trajectory))
            for trajectory in trajectories
        ]

        def _trajectory_combo(current_key=None, none_label="<Unassigned>"):
            combo = QtWidgets.QComboBox()
            combo.addItem(none_label, None)
            for text, primary_key in trajectory_options:
                combo.addItem(text, primary_key)
            if current_key:
                for combo_idx in range(combo.count()):
                    if combo.itemData(combo_idx) == current_key:
                        combo.setCurrentIndex(combo_idx)
                        break
            return combo

        tabs = QtWidgets.QTabWidget()
        combos = {}
        micro_combos = {}
        if rows:
            map_tab = QWidget()
            map_layout = QVBoxLayout(map_tab)
            source_label = QLabel(
                f"Map file: {os.path.basename(self.electrode_map_path or '')}"
            )
            source_label.setStyleSheet("font-weight: bold;")
            map_layout.addWidget(source_label)
            rule_label = QLabel("Auto-match rule: first letter is side; remaining code matches worksheet Label, ignoring symbols like *.")
            rule_label.setWordWrap(True)
            map_layout.addWidget(rule_label)

            table = QtWidgets.QTableWidget(len(rows), 4)
            table.setHorizontalHeaderLabels(["Map Shaft", "Contacts", "First Contact", "Planning Trajectory"])
            table.verticalHeader().setVisible(False)
            table.setSelectionBehavior(QtWidgets.QAbstractItemView.SelectRows)
            table.setEditTriggers(QtWidgets.QAbstractItemView.NoEditTriggers)

            for row_idx, map_entry in enumerate(rows):
                prefix = map_entry.get("shaft")
                contacts = map_entry.get("contact_count")
                first_contact = map_entry.get("first_contact") or ""
                table.setItem(row_idx, 0, QtWidgets.QTableWidgetItem(str(prefix)))
                table.setItem(row_idx, 1, QtWidgets.QTableWidgetItem(str(contacts or "")))
                table.setItem(row_idx, 2, QtWidgets.QTableWidgetItem(str(first_contact)))

                combo = _trajectory_combo(self.map_trajectory_assignments.get(prefix))
                combos[prefix] = combo
                table.setCellWidget(row_idx, 3, combo)

            table.resizeColumnsToContents()
            header = table.horizontalHeader()
            header.setStretchLastSection(True)
            map_layout.addWidget(table)
            tabs.addTab(map_tab, "Map File")

        if micro_rows:
            micro_tab = QWidget()
            micro_layout = QVBoxLayout(micro_tab)
            micro_label = QLabel(
                f"Micro contacts from: {os.path.basename(self.electrode_map_path or '')}"
            )
            micro_label.setStyleSheet("font-weight: bold;")
            micro_layout.addWidget(micro_label)
            micro_rule_label = QLabel(
                "Assign micro bundles to planning trajectories when target-name matching is ambiguous or missing."
            )
            micro_rule_label.setWordWrap(True)
            micro_layout.addWidget(micro_rule_label)

            micro_table = QtWidgets.QTableWidget(len(micro_rows), 4)
            micro_table.setHorizontalHeaderLabels(["Micro Shaft", "Contacts", "First Contact", "Planning Trajectory"])
            micro_table.verticalHeader().setVisible(False)
            micro_table.setSelectionBehavior(QtWidgets.QAbstractItemView.SelectRows)
            micro_table.setEditTriggers(QtWidgets.QAbstractItemView.NoEditTriggers)

            for row_idx, map_entry in enumerate(micro_rows):
                prefix = map_entry.get("shaft")
                contacts = map_entry.get("contact_count")
                first_contact = map_entry.get("first_contact") or ""
                micro_table.setItem(row_idx, 0, QtWidgets.QTableWidgetItem(str(prefix)))
                micro_table.setItem(row_idx, 1, QtWidgets.QTableWidgetItem(str(contacts or "")))
                micro_table.setItem(row_idx, 2, QtWidgets.QTableWidgetItem(str(first_contact)))

                current_key = self.micro_map_trajectory_assignments.get(prefix)
                if not current_key:
                    trajectory = self._planning_trajectory_for_micro_target(prefix)
                    current_key = self._trajectory_primary_key(trajectory) if trajectory is not None else None
                combo = _trajectory_combo(current_key)
                micro_combos[prefix] = combo
                micro_table.setCellWidget(row_idx, 3, combo)

            micro_table.resizeColumnsToContents()
            micro_header = micro_table.horizontalHeader()
            micro_header.setStretchLastSection(True)
            micro_layout.addWidget(micro_table)
            tabs.addTab(micro_tab, "Micro Contacts")

        shaft_combos = {}
        shaft_current_keys = {}
        if shaft_rows:
            shaft_tab = QWidget()
            shaft_layout = QVBoxLayout(shaft_tab)
            info_label = QLabel("Assign detected shaft groups to planning worksheet trajectories in bulk.")
            info_label.setWordWrap(True)
            shaft_layout.addWidget(info_label)

            shaft_table = QtWidgets.QTableWidget(len(shaft_rows), 4)
            shaft_table.setHorizontalHeaderLabels(["Detected Shaft", "Contacts", "Current Trajectory", "Assign To"])
            shaft_table.verticalHeader().setVisible(False)
            shaft_table.setSelectionBehavior(QtWidgets.QAbstractItemView.SelectRows)
            shaft_table.setEditTriggers(QtWidgets.QAbstractItemView.NoEditTriggers)

            for row_idx, (shaft_name, indices) in enumerate(shaft_rows):
                current_key = self._current_trajectory_key_for_shaft(shaft_name, indices)
                shaft_current_keys[shaft_name] = current_key
                current_trajectory = self._trajectory_by_primary_key(current_key) if current_key else None
                current_text = self._trajectory_display_name(current_trajectory) if current_trajectory else ""
                shaft_table.setItem(row_idx, 0, QtWidgets.QTableWidgetItem(str(shaft_name)))
                shaft_table.setItem(row_idx, 1, QtWidgets.QTableWidgetItem(str(len(indices))))
                shaft_table.setItem(row_idx, 2, QtWidgets.QTableWidgetItem(current_text))
                combo = _trajectory_combo(current_key, none_label="<No planning trajectory>")
                shaft_combos[shaft_name] = combo
                shaft_table.setCellWidget(row_idx, 3, combo)

            shaft_table.resizeColumnsToContents()
            shaft_header = shaft_table.horizontalHeader()
            shaft_header.setStretchLastSection(True)
            shaft_layout.addWidget(shaft_table)
            tabs.addTab(shaft_tab, "Detected Shafts")
            tabs.setCurrentWidget(shaft_tab)

        layout.addWidget(tabs)

        def _duplicate_shaft_assignment_text(assignments):
            key_to_shafts = {}
            for shaft_name, primary_key in assignments.items():
                if primary_key:
                    key_to_shafts.setdefault(primary_key, []).append(shaft_name)
            duplicates = {
                primary_key: names
                for primary_key, names in key_to_shafts.items()
                if len(names) > 1
            }
            if not duplicates:
                return None
            lines = []
            for primary_key, names in duplicates.items():
                trajectory = self._trajectory_by_primary_key(primary_key)
                label = self._trajectory_display_name(trajectory) if trajectory else primary_key
                lines.append(f"{label}: {', '.join(names)}")
            return "Each planning trajectory can be assigned to one detected shaft.\n\n" + "\n".join(lines)

        def _validate_and_accept():
            shaft_assignments = {
                shaft_name: combo.currentData()
                for shaft_name, combo in shaft_combos.items()
            }
            shaft_changed = any(
                shaft_assignments.get(shaft_name) != shaft_current_keys.get(shaft_name)
                for shaft_name in shaft_assignments
            )
            duplicate_text = _duplicate_shaft_assignment_text(shaft_assignments) if shaft_changed else None
            if duplicate_text:
                QMessageBox.information(self, "Map Trajectories", duplicate_text)
                return
            dialog.accept()

        buttons = QtWidgets.QDialogButtonBox(
            QtWidgets.QDialogButtonBox.Ok | QtWidgets.QDialogButtonBox.Cancel
        )
        buttons.accepted.connect(_validate_and_accept)
        buttons.rejected.connect(dialog.reject)
        layout.addWidget(buttons)

        if dialog.exec_() != QDialog.Accepted:
            return

        shaft_assignments = {
            shaft_name: combo.currentData()
            for shaft_name, combo in shaft_combos.items()
        }
        shaft_changed = any(
            shaft_assignments.get(shaft_name) != shaft_current_keys.get(shaft_name)
            for shaft_name in shaft_assignments
        )
        duplicate_text = _duplicate_shaft_assignment_text(shaft_assignments) if shaft_changed else None
        if duplicate_text:
            QMessageBox.information(self, "Map Trajectories", duplicate_text)
            return

        assignments = {}
        for prefix, combo in combos.items():
            primary_key = combo.currentData()
            if primary_key:
                assignments[prefix] = primary_key
        self.map_trajectory_assignments = assignments
        micro_assignments = {}
        for prefix, combo in micro_combos.items():
            primary_key = combo.currentData()
            if primary_key:
                micro_assignments[prefix] = primary_key
        self.micro_map_trajectory_assignments = micro_assignments
        shaft_updated = self._apply_bulk_shaft_trajectory_assignments(
            shaft_assignments,
            shaft_current_keys,
        ) if shaft_assignments else 0
        updated = self._apply_map_assignments_to_existing_electrodes()
        self._update_shaft_list_with_colors(self.before_shaft_list, self.before_shafts)
        self._on_before_shaft_selected(redraw_2d=False, sync_3d=False)
        self._refresh_before_view_after_shaft_metadata_change()
        self._schedule_detected_sync(self.before_electrodes)
        self._force_3d_shaft_refresh(self.before_electrodes)
        parts = []
        if combos:
            parts.append(f"mapped {len(assignments)} .map shafts")
        if micro_combos:
            parts.append(f"mapped {len(micro_assignments)} micro shafts")
        if shaft_assignments:
            parts.append(f"updated {shaft_updated} detected shaft assignments")
        if updated:
            parts.append(f"refreshed {updated} contact labels")
        self.log("Map Trajectories: " + "; ".join(parts) + ".")
    
    def _edit_shaft_from_widget(self, widget, shaft_list):
        """Handle edit button click from ListItemWidget."""
        if shaft_list != self.before_shaft_list:
            return
        # Find the item for this widget
        for i in range(shaft_list.count()):
            item = shaft_list.item(i)
            if shaft_list.itemWidget(item) == widget:
                shaft_name = item.data(Qt.UserRole)
                selected_names = []
                for selected_item in shaft_list.selectedItems():
                    name = selected_item.data(Qt.UserRole)
                    if not name:
                        text = selected_item.text()
                        if text:
                            name = text.split(' (')[0]
                    if name and name not in selected_names:
                        selected_names.append(name)

                if len(selected_names) > 1 and shaft_name in selected_names:
                    merge_btn = QMessageBox()
                    merge_btn.setWindowTitle("Edit Shaft")
                    merge_btn.setText("Multiple shafts selected. What would you like to do?")
                    merge_selected = merge_btn.addButton("Merge Selected", QMessageBox.AcceptRole)
                    edit_btn = merge_btn.addButton("Edit This Shaft", QMessageBox.ActionRole)
                    cancel_btn = merge_btn.addButton("Cancel", QMessageBox.RejectRole)
                    merge_btn.setDefaultButton(merge_selected)
                    merge_btn.exec_()

                    if merge_btn.clickedButton() == merge_selected:
                        self._merge_selected_shafts(selected_names)
                        return
                    if merge_btn.clickedButton() != edit_btn:
                        return

                edit = self._prompt_shaft_edit(shaft_name)
                if edit is not None:
                    self._apply_shaft_edit(
                        shaft_name,
                        edit["name"],
                        edit["trajectory_key"],
                        edit["current_trajectory_key"],
                        edit.get("name_edited", False),
                    )
                break
    
    def _toggle_shaft_visibility_from_widget(self, widget, shaft_list):
        """Handle hide button click from ListItemWidget."""
        if shaft_list != self.before_shaft_list:
            return
        # Find the item for this widget
        for i in range(shaft_list.count()):
            item = shaft_list.item(i)
            if shaft_list.itemWidget(item) == widget:
                shaft_name = item.data(Qt.UserRole)

                visibility_dict = self.before_shaft_visibility
                # Toggle visibility
                is_visible = visibility_dict.get(shaft_name, True)
                visibility_dict[shaft_name] = not is_visible
                
                # Update widget button state
                widget.set_hidden_state(not visibility_dict[shaft_name])
                
                # Redraw
                if self.before_mri_data is not None:
                    self._plot_slices(self.before_mri_data, self.before_ct_data, self.before_canvas,
                                    "Unaligned CT Overlaid on MRI",
                                    self.before_slice_x, self.before_slice_y, self.before_slice_z)
                self._schedule_detected_sync(self.before_electrodes)
                
                status = "hidden" if not visibility_dict.get(shaft_name, True) else "shown"
                self.log(f"Shaft '{shaft_name}' {status}")
                break
    
    def _delete_shaft_from_widget(self, widget, shaft_list):
        """Handle delete button click from ListItemWidget."""
        if shaft_list != self.before_shaft_list:
            return
        # Find the item for this widget
        for i in range(shaft_list.count()):
            item = shaft_list.item(i)
            if shaft_list.itemWidget(item) == widget:
                shaft_name = item.data(Qt.UserRole)
                
                # Confirm deletion
                reply = QMessageBox.question(
                    self, "Delete Shaft",
                    f"Are you sure you want to delete shaft '{shaft_name}' and all its electrodes?",
                    QMessageBox.Yes | QMessageBox.No
                )
                
                if reply == QMessageBox.Yes:
                    shafts_dict = self.before_shafts
                    visibility_dict = self.before_shaft_visibility
                    electrodes = self.before_electrodes
                    
                    # Get electrode indices to delete
                    indices_to_delete = shafts_dict[shaft_name]
                    
                    # Remove electrodes (in reverse order to maintain indices)
                    for idx in sorted(indices_to_delete, reverse=True):
                        if idx < len(electrodes):
                            electrodes.pop(idx)
                    
                    # Remove shaft
                    del shafts_dict[shaft_name]
                    if shaft_name in visibility_dict:
                        del visibility_dict[shaft_name]
                    
                    # Rebuild shaft indices
                    shafts_dict.clear()
                    for idx, electrode in enumerate(electrodes):
                        shaft = electrode.shaft
                        if shaft not in shafts_dict:
                            shafts_dict[shaft] = []
                        shafts_dict[shaft].append(idx)
                    
                    self.log(f"Deleted shaft '{shaft_name}'")
                    self._update_shaft_list_with_colors(shaft_list, shafts_dict)
                    
                    # Update electrode count
                    self.before_electrode_count_label.setText(f"Electrodes: {len(self.before_electrodes)}")
                    if self.before_mri_data is not None:
                        self._plot_slices(self.before_mri_data, self.before_ct_data, self.before_canvas,
                                        "Unaligned CT Overlaid on MRI",
                                        self.before_slice_x, self.before_slice_y, self.before_slice_z)
                break

    def _merge_before_shafts(self):
        """Merge two shafts in the before view."""
        self._merge_shafts()

    def _merge_shafts(self):
        """Merge a source shaft into a target shaft."""
        shafts_dict = self.before_shafts
        shaft_list = self.before_shaft_list

        if len(shafts_dict) < 2:
            QMessageBox.information(self, "Merge Shafts", "At least two shafts are required to merge.")
            return

        choices = self._sorted_shaft_names(shafts_dict, include_unassigned=True)
        if len(choices) < 2:
            QMessageBox.information(self, "Merge Shafts", "At least two shafts are required to merge.")
            return

        selected_name = None
        if shaft_list is not None:
            selected_items = shaft_list.selectedItems()
            if selected_items:
                selected_name = selected_items[0].data(Qt.UserRole)

        default_target = selected_name if selected_name in choices else choices[0]
        default_source = next((name for name in choices if name != default_target), choices[0])

        dialog = QDialog(self)
        dialog.setWindowTitle("Merge Shafts")
        layout = QVBoxLayout(dialog)
        layout.addWidget(QLabel("Merge source shaft into target shaft:"))

        target_combo = QtWidgets.QComboBox()
        source_combo = QtWidgets.QComboBox()
        target_combo.addItems(choices)
        source_combo.addItems(choices)

        if default_target in choices:
            target_combo.setCurrentIndex(choices.index(default_target))
        if default_source in choices:
            source_combo.setCurrentIndex(choices.index(default_source))

        form = QGridLayout()
        form.addWidget(QLabel("Target:"), 0, 0)
        form.addWidget(target_combo, 0, 1)
        form.addWidget(QLabel("Source:"), 1, 0)
        form.addWidget(source_combo, 1, 1)
        layout.addLayout(form)

        buttons = QtWidgets.QDialogButtonBox(
            QtWidgets.QDialogButtonBox.Ok | QtWidgets.QDialogButtonBox.Cancel
        )
        buttons.accepted.connect(dialog.accept)
        buttons.rejected.connect(dialog.reject)
        layout.addWidget(buttons)

        if dialog.exec_() != QDialog.Accepted:
            return

        target = target_combo.currentText()
        source = source_combo.currentText()
        if target == source:
            QMessageBox.information(self, "Merge Shafts", "Target and source must be different.")
            return

        self._merge_into_target(target, [source])
        self.log(f"Merged shaft '{source}' into '{target}'")

    def _merge_selected_shafts(self, selected_names):
        """Merge multiple selected shafts into a chosen target."""
        shafts_dict = self.before_shafts
        if len(selected_names) < 2:
            QMessageBox.information(self, "Merge Shafts", "Select at least two shafts to merge.")
            return

        candidates = [name for name in selected_names if name in shafts_dict]
        if len(candidates) < 2:
            QMessageBox.information(self, "Merge Shafts", "Selected shafts are no longer available.")
            return

        target, ok = QInputDialog.getItem(
            self,
            "Merge Shafts",
            "Select target shaft:",
            candidates,
            0,
            False
        )
        if not ok or not target:
            return
        sources = [name for name in candidates if name != target]
        if not sources:
            QMessageBox.information(self, "Merge Shafts", "Target and source must be different.")
            return

        self._merge_into_target(target, sources)
        self.log(f"Merged shafts {', '.join(sources)} into '{target}'")

    def _merge_into_target(self, target, sources):
        """Merge source shafts into target shaft (shared logic)."""
        t_overall = time.perf_counter() if PROFILE_MERGE else None
        shafts_dict = self.before_shafts
        visibility_dict = self.before_shaft_visibility
        electrodes = self.before_electrodes
        shaft_list = self.before_shaft_list

        if target not in shafts_dict:
            QMessageBox.information(self, "Merge Shafts", "Target shaft no longer exists.")
            return
        valid_sources = [name for name in sources if name in shafts_dict and name != target]
        if not valid_sources:
            QMessageBox.information(self, "Merge Shafts", "No valid source shafts to merge.")
            return

        merged_indices = list(shafts_dict.get(target, []))
        for source in valid_sources:
            source_indices = list(shafts_dict.get(source, []))
            for idx in source_indices:
                if idx < len(electrodes):
                    electrodes[idx].shaft = target
            merged_indices.extend(source_indices)
            shafts_dict.pop(source, None)
            visibility_dict[target] = visibility_dict.get(target, True) or visibility_dict.get(source, True)
            visibility_dict.pop(source, None)

        shafts_dict[target] = list(dict.fromkeys(merged_indices))

        t_list = time.perf_counter() if PROFILE_MERGE else None
        self._update_shaft_list_with_colors(shaft_list, shafts_dict)
        if shaft_list is not None:
            blocker = QtCore.QSignalBlocker(shaft_list)
            for i in range(shaft_list.count()):
                item = shaft_list.item(i)
                item_name = item.data(Qt.UserRole)
                if not item_name:
                    text = item.text()
                    if text:
                        item_name = text.split(' (')[0]
                if item_name == target:
                    shaft_list.setCurrentItem(item)
                    break
            del blocker
        if PROFILE_MERGE and t_list is not None:
            self.log(f"[profile] Shaft list rebuild: {time.perf_counter() - t_list:.3f}s")

        t_select = time.perf_counter() if PROFILE_MERGE else None
        self._on_before_shaft_selected()
        if PROFILE_MERGE and t_select is not None:
            self.log(f"[profile] Electrode list populate: {time.perf_counter() - t_select:.3f}s")

        self._schedule_detected_sync(electrodes)
        if PROFILE_MERGE and t_overall is not None:
            self.log(f"[profile] Merge total: {time.perf_counter() - t_overall:.3f}s")
    
    def _edit_electrode_from_widget(self, widget, electrode_list):
        """Handle edit button click from electrode ListItemWidget."""
        if electrode_list != self.before_electrode_list:
            return
        # Find the item for this widget
        for i in range(electrode_list.count()):
            item = electrode_list.item(i)
            if electrode_list.itemWidget(item) == widget:
                electrode_idx = item.data(Qt.UserRole)
                selected_items = electrode_list.selectedItems()
                selected_indices = [
                    it.data(Qt.UserRole) for it in selected_items
                    if it.data(Qt.UserRole) is not None
                ]
                if electrode_idx not in selected_indices:
                    selected_indices = [electrode_idx]
                
                electrodes = self.before_electrodes
                shafts_dict = self.before_shafts
                shaft_list = self.before_shaft_list
                shaft_visibility = self.before_shaft_visibility
                
                if electrode_idx < len(electrodes):
                    elec = electrodes[electrode_idx]
                    action, ok = QInputDialog.getItem(
                        self, "Edit Electrode", "Action:", ["Rename", "Assign to Shaft"], 0, False
                    )
                    if not ok or not action:
                        break

                    if action == "Rename":
                        if len(selected_indices) > 1:
                            QMessageBox.information(
                                self,
                                "Rename Electrode",
                                "Rename works for a single electrode. Use the Assign button for bulk actions."
                            )
                            break
                        old_name = elec.name or f"Electrode {electrode_idx+1}"
                        new_name, ok = QInputDialog.getText(
                            self, "Rename Electrode", f"Enter new name for '{old_name}':", 
                            text=old_name
                        )
                        if ok and new_name:
                            elec.name = new_name
                            widget.set_text(new_name)
                            self.log(f"Renamed electrode to '{new_name}'")
                    else:
                        current_shaft = elec.shaft
                        choice, ok = self._prompt_shaft_choice(
                            "Assign to Shaft", current_shaft, shafts_dict, include_new=True
                        )
                        if not ok or not choice:
                            break
                        if choice == "<New Shaft>":
                            default_name = f"S_{self.shaft_counter}"
                            new_name, ok = QInputDialog.getText(
                                self, "New Shaft", "Enter new shaft name:", text=default_name
                            )
                            if not ok or not new_name:
                                break
                            if new_name in shafts_dict:
                                self.log(f"Shaft name '{new_name}' already exists!")
                                break
                            choice = new_name
                            self.shaft_counter += 1
                        target_trajectory = self._resolve_target_planning_trajectory(
                            choice,
                            shafts_dict.get(choice, []),
                        )
                        if all(electrodes[idx].shaft == choice for idx in selected_indices):
                            break

                        for idx in selected_indices:
                            if idx is not None and idx < len(electrodes):
                                electrodes[idx].shaft = choice
                        shafts_dict.clear()
                        for idx, electrode in enumerate(electrodes):
                            shaft = electrode.shaft
                            shafts_dict.setdefault(shaft, []).append(idx)

                        for name in list(shaft_visibility.keys()):
                            if name not in shafts_dict:
                                del shaft_visibility[name]
                        for name in shafts_dict:
                            if name not in shaft_visibility:
                                shaft_visibility[name] = True

                        target_indices = shafts_dict.get(choice, selected_indices)
                        if target_trajectory is not None:
                            self._apply_planning_trajectory_to_indices(
                                choice,
                                target_trajectory,
                                target_indices,
                                "manual_assignment",
                            )
                        else:
                            self._clear_planning_metadata_for_indices(selected_indices)

                        self._update_shaft_list_with_colors(shaft_list, shafts_dict)
                        for j in range(shaft_list.count()):
                            shaft_item = shaft_list.item(j)
                            item_name = shaft_item.data(Qt.UserRole)
                            if not item_name:
                                text = shaft_item.text()
                                if text:
                                    item_name = text.split(' (')[0]
                            if item_name == choice:
                                shaft_list.setCurrentItem(shaft_item)
                                self._on_before_shaft_selected()
                                break
                        self._select_electrodes_in_list(selected_indices)

                        if self.before_mri_data is not None:
                            self._plot_slices(self.before_mri_data, self.before_ct_data, self.before_canvas,
                                            "Unaligned CT Overlaid on MRI",
                                            self.before_slice_x, self.before_slice_y, self.before_slice_z)
                break
    
    def _toggle_electrode_visibility_from_widget(self, widget, electrode_list):
        """Handle hide button click from electrode ListItemWidget."""
        if electrode_list != self.before_electrode_list:
            return
        # Find the item for this widget
        for i in range(electrode_list.count()):
            item = electrode_list.item(i)
            if electrode_list.itemWidget(item) == widget:
                electrode_idx = item.data(Qt.UserRole)

                electrodes = self.before_electrodes

                selected_items = electrode_list.selectedItems()
                selected_indices = [it.data(Qt.UserRole) for it in selected_items if it.data(Qt.UserRole) is not None]
                apply_multi = electrode_idx in selected_indices and len(selected_indices) > 1

                if apply_multi:
                    any_visible = any(
                        idx < len(electrodes) and electrodes[idx].visible
                        for idx in selected_indices
                    )
                    new_visible = False if any_visible else True
                    for idx in selected_indices:
                        if idx < len(electrodes):
                            electrodes[idx].visible = new_visible
                    for j in range(electrode_list.count()):
                        list_item = electrode_list.item(j)
                        idx = list_item.data(Qt.UserRole)
                        if idx in selected_indices:
                            list_widget = electrode_list.itemWidget(list_item)
                            if list_widget is not None:
                                list_widget.set_hidden_state(not new_visible)
                    status = "hidden" if not new_visible else "shown"
                    self.log(f"Electrodes {status} (n={len(selected_indices)})")
                elif electrode_idx < len(electrodes):
                    elec = electrodes[electrode_idx]
                    is_visible = elec.visible
                    elec.visible = not is_visible
                    widget.set_hidden_state(not elec.visible)
                    status = "hidden" if not elec.visible else "shown"
                    self.log(f"Electrode {electrode_idx+1} {status}")
                
                # Redraw
                if self.before_mri_data is not None:
                    self._plot_slices(self.before_mri_data, self.before_ct_data, self.before_canvas,
                                    "Unaligned CT Overlaid on MRI",
                                    self.before_slice_x, self.before_slice_y, self.before_slice_z)
                break
    
    def _delete_electrode_from_widget(self, widget, electrode_list):
        """Handle delete button click from electrode ListItemWidget."""
        if electrode_list != self.before_electrode_list:
            return
        # Find the item for this widget
        for i in range(electrode_list.count()):
            item = electrode_list.item(i)
            if electrode_list.itemWidget(item) == widget:
                electrode_idx = item.data(Qt.UserRole)

                electrodes = self.before_electrodes
                shafts_dict = self.before_shafts
                shaft_list = self.before_shaft_list

                selected_items = electrode_list.selectedItems()
                selected_indices = [it.data(Qt.UserRole) for it in selected_items if it.data(Qt.UserRole) is not None]
                apply_multi = electrode_idx in selected_indices and len(selected_indices) > 1

                if apply_multi:
                    reply = QMessageBox.question(
                        self, "Delete Electrodes",
                        f"Are you sure you want to delete {len(selected_indices)} electrodes?",
                        QMessageBox.Yes | QMessageBox.No
                    )
                    if reply != QMessageBox.Yes:
                        break
                    for idx in sorted(selected_indices, reverse=True):
                        if idx < len(electrodes):
                            electrodes.pop(idx)
                    self.log(f"Deleted {len(selected_indices)} electrodes")
                elif electrode_idx < len(electrodes):
                    elec = electrodes[electrode_idx]
                    elec_name = elec.name or f"Electrode {electrode_idx+1}"
                    reply = QMessageBox.question(
                        self, "Delete Electrode",
                        f"Are you sure you want to delete '{elec_name}'?",
                        QMessageBox.Yes | QMessageBox.No
                    )
                    if reply != QMessageBox.Yes:
                        break
                    electrodes.pop(electrode_idx)
                    self.log(f"Deleted electrode '{elec_name}'")
                else:
                    break

                # Rebuild shaft indices
                shafts_dict.clear()
                for idx, electrode in enumerate(electrodes):
                    shaft = electrode.shaft
                    if shaft not in shafts_dict:
                        shafts_dict[shaft] = []
                    shafts_dict[shaft].append(idx)

                # Update UI
                self.before_electrode_count_label.setText(f"Electrodes: {len(self.before_electrodes)}")
                self._update_shaft_list_with_colors(shaft_list, shafts_dict)

                selected_shafts = shaft_list.selectedItems()
                if selected_shafts:
                    self._on_before_shaft_selected()

                # Redraw
                if self.before_mri_data is not None:
                    self._plot_slices(self.before_mri_data, self.before_ct_data, self.before_canvas,
                                    "Unaligned CT Overlaid on MRI",
                                    self.before_slice_x, self.before_slice_y, self.before_slice_z)
                break

    def _assign_selected_electrodes(self):
        """Assign selected electrodes to an existing or new shaft."""
        electrode_list = self.before_electrode_list
        shaft_list = self.before_shaft_list
        electrodes = self.before_electrodes
        shafts = self.before_shafts
        shaft_visibility = self.before_shaft_visibility

        if electrode_list is None or shaft_list is None:
            return

        selected_items = electrode_list.selectedItems()
        if not selected_items:
            self.log("No electrodes selected")
            QMessageBox.information(self, "Assign to Shaft", "Select one or more electrodes first.")
            return

        selected_indices = [it.data(Qt.UserRole) for it in selected_items if it.data(Qt.UserRole) is not None]
        if not selected_indices:
            self.log("No electrodes selected")
            QMessageBox.information(self, "Assign to Shaft", "Select one or more electrodes first.")
            return

        selected_shafts = {
            electrodes[idx].shaft
            for idx in selected_indices
            if idx is not None and idx < len(electrodes)
        }
        current_shaft = None
        if len(selected_shafts) == 1:
            only_shaft = next(iter(selected_shafts))
            if only_shaft != "Unassigned":
                current_shaft = only_shaft

        choice, ok = self._prompt_shaft_choice("Assign to Shaft", current_shaft, shafts, include_new=True)
        if not ok or not choice:
            return

        target_shaft = choice
        if choice == "<New Shaft>":
            default_name = f"S_{self.shaft_counter}"
            new_name, ok = QInputDialog.getText(
                self, "New Shaft", "Enter new shaft name:", text=default_name
            )
            if not ok or not new_name:
                return
            if new_name in shafts:
                self.log(f"Shaft name '{new_name}' already exists!")
                return
            target_shaft = new_name
            self.shaft_counter += 1
        target_trajectory = self._resolve_target_planning_trajectory(
            target_shaft,
            shafts.get(target_shaft, []),
        )

        for idx in selected_indices:
            if idx is not None and idx < len(electrodes):
                electrodes[idx].shaft = target_shaft

        shafts.clear()
        for idx, electrode in enumerate(electrodes):
            shaft = electrode.shaft
            shafts.setdefault(shaft, []).append(idx)

        for name in list(shaft_visibility.keys()):
            if name not in shafts:
                del shaft_visibility[name]
        for name in shafts:
            if name not in shaft_visibility:
                shaft_visibility[name] = True

        target_indices = shafts.get(target_shaft, selected_indices)
        if target_trajectory is not None:
            self._apply_planning_trajectory_to_indices(
                target_shaft,
                target_trajectory,
                target_indices,
                "manual_assignment",
            )
        else:
            self._clear_planning_metadata_for_indices(selected_indices)

        self._update_shaft_list_with_colors(shaft_list, shafts)

        for i in range(shaft_list.count()):
            item = shaft_list.item(i)
            item_name = item.data(Qt.UserRole)
            if not item_name:
                text = item.text()
                if text:
                    item_name = text.split(' (')[0]
            if item_name == target_shaft:
                shaft_list.setCurrentItem(item)
                self._on_before_shaft_selected()
                break

        self._select_electrodes_in_list(selected_indices)

    def _toggle_all_electrodes(self):
        """Show or hide all electrodes for the current view."""
        electrodes = self.before_electrodes
        list_widget = self.before_electrode_list
        button = self.before_showhide_all_btn
        visible = not getattr(self, 'before_all_electrodes_visible', True)

        for elec in electrodes:
            elec.visible = visible

        self.before_all_electrodes_visible = visible
        self.before_detected_visible = visible

        try:
            if list_widget is not None:
                for i in range(list_widget.count()):
                    item = list_widget.item(i)
                    widget = list_widget.itemWidget(item)
                    if widget is not None:
                        widget.set_hidden_state(not visible)
        except Exception:
            pass

        try:
            if button:
                button.setText('Hide All Electrodes' if visible else 'Show All Electrodes')
        except Exception:
            pass

        if self.before_mri_data is not None:
            self._plot_slices(self.before_mri_data, self.before_ct_data, self.before_canvas,
                              "Unaligned CT Overlaid on MRI",
                              self.before_slice_x, self.before_slice_y, self.before_slice_z)
        self._schedule_detected_sync(electrodes)
    
    def _clear_before_electrodes(self):
        """Clear all electrode markers from before view."""
        self.before_electrodes = []
        self.before_shafts = {}
        self.before_electrode_count_label.setText("Electrodes: 0")
        self.before_shaft_list.clear()
        self.before_all_electrodes_visible = True
        self.before_detected_visible = True
        self.before_detected_positions = None
        self.electrode_space = self._current_electrode_space()
        self.electrode_display_space = self.electrode_space
        self.before_selected_parcellation_label = None
        self.before_selected_parcellation_name = None
        try:
            self.before_group_shafts_btn.setEnabled(False)
        except Exception:
            pass
        try:
            self.before_showhide_all_btn.setText("Hide All Electrodes")
        except Exception:
            pass
        if self.before_mri_data is not None:
            self._plot_slices(self.before_mri_data, self.before_ct_data, self.before_canvas,
                            "Unaligned CT Overlaid on MRI",
                            self.before_slice_x, self.before_slice_y, self.before_slice_z)
        self._sync_detected_to_controller([])
        self._update_mni_button_state()
        self.log("Cleared all electrode markers from Before view")
    def _auto_detect_before_electrodes(self):
        """Automatically detect electrodes using MIP-based 2D detection + slice scanning."""
        if self.before_ct_data is None:
            self.log("Error: No CT data loaded for auto-detection")
            return

        self.electrode_space = self._current_electrode_space()
        self.electrode_display_space = self.electrode_space
        
        self.log("Starting MIP-based electrode detection...")
        disable_on_exit = True
        try:
            # Get threshold values from UI
            threshold_min = self.before_threshold_lower.value()
            threshold_max = self.before_threshold_upper.value()
            
            # Get intensity distribution
            ct_max = np.max(self.before_ct_data)
            ct_min = np.min(self.before_ct_data)
            ct_mean = np.mean(self.before_ct_data)
            ct_999 = np.percentile(self.before_ct_data, 99.9)
            ct_9999 = np.percentile(self.before_ct_data, 99.99)
            
            self.log(f"CT intensity: min={ct_min:.1f}, mean={ct_mean:.1f}, 99.9%={ct_999:.1f}, 99.99%={ct_9999:.1f}, max={ct_max:.1f}")
            self.log(f"Using threshold range: {threshold_min:.0f} - {threshold_max:.0f} HU")
            
            # Build a conservative brain mask from CT to suppress extracranial detections
            mask_dilate = self.before_mask_dilate_spinbox.value() if hasattr(self, 'before_mask_dilate_spinbox') else DEFAULT_MASK_DILATE
            if getattr(self, 'before_mask_filter_checkbox', None) and not self.before_mask_filter_checkbox.isChecked():
                brain_mask = np.ones_like(self.before_ct_data, dtype=bool)
                self.log("Brain mask filtering disabled.")
            else:
                brain_mask = self._build_brain_mask(self.before_ct_data, dilate_iters=int(mask_dilate))

            # Create thresholded data for MIP computation (do not mask electrodes out)
            ct_thresh = np.where(
                (self.before_ct_data >= threshold_min) & (self.before_ct_data <= threshold_max),
                self.before_ct_data,
                np.nan
            )
            
            # Step 1: Compute MIPs (Maximum Intensity Projections)
            self.log("Computing axial and coronal MIPs...")
            ct_mip_axial = self._safe_nanmax(ct_thresh, axis=2)    # X-Y view (max along Z)
            ct_mip_coronal = self._safe_nanmax(ct_thresh, axis=1)  # X-Z view (max along Y)
            
            # Step 2: Find blobs in axial MIP (X-Y positions)
            self.log("Finding electrode blobs in Axial MIP...")
            blob_min = self.before_blob_min_spinbox.value() if hasattr(self, 'before_blob_min_spinbox') else DEFAULT_BLOB_MIN
            axial_positions = self._find_mip_blob_centroids(
                ct_mip_axial, threshold_min, threshold_max, min_size=int(blob_min)
            )
            self.log(f"Found {len(axial_positions)} blobs in Axial MIP")
            self.before_axial_mip_positions = axial_positions
            
            # Step 3: Find blobs in coronal MIP (X-Z positions)
            self.log("Finding electrode blobs in Coronal MIP...")
            coronal_positions = self._find_mip_blob_centroids(
                ct_mip_coronal, threshold_min, threshold_max, min_size=int(blob_min)
            )
            self.log(f"Found {len(coronal_positions)} blobs in Coronal MIP")
            
            if len(axial_positions) == 0:
                self.log("No electrodes detected in Axial MIP. Try adjusting threshold.")
                return

            # Step 4: Expand axial candidates across Z to get all 3D hits
            self.log("Expanding axial MIP candidates across Z slices (ROI radius=2)...")
            electrode_positions_3d, axial_failed = self._expand_axial_candidates_to_3d(
                self.before_ct_data, axial_positions, threshold_min, threshold_max, roi_radius=2
            )
            self._save_mip_debug(
                ct_mip_axial,
                ct_mip_coronal,
                axial_positions,
                coronal_positions,
                threshold_min,
                threshold_max,
                mri_data=self.before_mri_data,
                slice_y=getattr(self, "before_slice_y", None),
                slice_z=getattr(self, "before_slice_z", None),
                ct_alpha=self.before_ct_alpha,
                axial_failed_positions=axial_failed
            )
            
            if len(electrode_positions_3d) == 0:
                self.log("No electrodes verified from axial MIP candidates.")
                return
                
            self.log(f"Verified {len(electrode_positions_3d)} electrodes using MIP-based detection")
            
            # Filter out electrodes outside the brain using CT-derived mask
            self.log("Filtering electrodes using brain mask from CT...")
            
            # filtered_electrodes = []
            # for pos_3d in electrode_positions_3d:
            #     x_idx, y_idx, z_idx = pos_3d.astype(int)
            #     # Check if this position is within brain tissue
            #     if brain_mask[x_idx, y_idx, z_idx]:
            #         filtered_electrodes.append(pos_3d)
                    
            # if len(filtered_electrodes) == 0:
            #     self.log("Brain mask removed all electrodes; using unfiltered candidates.")
            # else:
            #     electrode_positions_3d = np.array(filtered_electrodes)
            
            # self.log(f"Kept {len(electrode_positions_3d)} electrodes within brain tissue")
            
            self.before_detected_positions = electrode_positions_3d
            self._populate_detected_electrodes(electrode_positions_3d)

            # Redraw
            if self.before_mri_data is not None:
                self._plot_slices(self.before_mri_data, self.before_ct_data, self.before_canvas,
                                "Unaligned CT Overlaid on MRI",
                                self.before_slice_x, self.before_slice_y, self.before_slice_z)
            
            self.log(f"Auto-detected {len(self.before_shafts)} shafts with {len(self.before_electrodes)} electrodes")
            
        except ImportError as e:
            disable_on_exit = False
            self.log(f"Error: Missing required library for auto-detection: {e}")
            self.log("Please install scipy and scikit-learn: pip install scipy scikit-learn")
        except Exception as e:
            disable_on_exit = False
            self.log(f"Error during auto-detection: {e}")
        finally:
            if disable_on_exit:
                try:
                    self.before_auto_detect_btn.setEnabled(False)
                except Exception:
                    pass
    
    def _rename_before_shaft(self):
        """Rename the selected shaft in before view."""
        current_item = self.before_shaft_list.currentItem()
        if current_item is None:
            self.log("Please select a shaft to rename")
            return
        current_name = current_item.data(Qt.UserRole)
        if not current_name:
            current_text = current_item.text()
            current_name = current_text.split(' (')[0]
        edit = self._prompt_shaft_edit(current_name)
        if edit is not None:
            self._apply_shaft_edit(
                current_name,
                edit["name"],
                edit["trajectory_key"],
                edit["current_trajectory_key"],
                edit.get("name_edited", False),
            )
    
    def _delete_before_shaft(self):
        """Delete the selected shaft and all its electrodes from before view."""
        current_item = self.before_shaft_list.currentItem()
        if current_item is None:
            self.log("Please select a shaft to delete")
            return
        
        current_text = current_item.text()
        shaft_name = current_text.split(' (')[0]
        
        # Confirm deletion
        from qtpy.QtWidgets import QMessageBox
        reply = QMessageBox.question(
            self,
            "Delete Shaft?",
            f"Are you sure you want to delete '{shaft_name}' and all its electrodes?",
            QMessageBox.Yes | QMessageBox.No,
            QMessageBox.No
        )
        
        if reply == QMessageBox.Yes:
            if shaft_name in self.before_shafts:
                # Get electrode indices to delete
                electrode_indices = self.before_shafts[shaft_name]
                num_electrodes = len(electrode_indices)
                
                # Remove electrodes (in reverse order to maintain indices)
                for idx in sorted(electrode_indices, reverse=True):
                    if idx < len(self.before_electrodes):
                        self.before_electrodes.pop(idx)
                
                # Rebuild shaft dictionary with updated indices
                old_shafts = self.before_shafts.copy()
                self.before_shafts = {}
                
                for old_shaft_name, old_indices in old_shafts.items():
                    if old_shaft_name == shaft_name:
                        continue  # Skip deleted shaft
                    
                    # Recalculate indices after deletion
                    new_indices = []
                    for old_idx in old_indices:
                        # Count how many electrodes were removed before this one
                        removed_before = sum(1 for del_idx in electrode_indices if del_idx < old_idx)
                        new_idx = old_idx - removed_before
                        new_indices.append(new_idx)
                    
                    self.before_shafts[old_shaft_name] = new_indices
                
                # Update UI with colored text
                self._update_shaft_list_with_colors(self.before_shaft_list, self.before_shafts)
                
                self.before_electrode_count_label.setText(f"Electrodes: {len(self.before_electrodes)}")
                
                # Redraw
                if self.before_mri_data is not None:
                    self._plot_slices(self.before_mri_data, self.before_ct_data, self.before_canvas,
                                    "Unaligned CT Overlaid on MRI",
                                    self.before_slice_x, self.before_slice_y, self.before_slice_z)
                
                self.log(f"Deleted shaft '{shaft_name}' with {num_electrodes} electrodes")
    
    def _delete_before_last_electrode(self):
        """Delete the most recently added electrode from before view."""
        if len(self.before_electrodes) == 0:
            self.log("No electrodes to delete")
            return
        
        # Remove last electrode
        last_electrode = self.before_electrodes.pop()
        shaft_name = last_electrode.shaft
        
        # Update shaft dictionary
        if shaft_name in self.before_shafts:
            # Remove the last index from this shaft
            self.before_shafts[shaft_name] = [idx for idx in self.before_shafts[shaft_name] 
                                              if idx < len(self.before_electrodes)]
            
            # If shaft is now empty, remove it
            if len(self.before_shafts[shaft_name]) == 0:
                del self.before_shafts[shaft_name]
        
        # Update UI
        self.before_shaft_list.clear()
        for name in sorted(self.before_shafts.keys()):
            count = len(self.before_shafts[name])
            self.before_shaft_list.addItem(f"{name} ({count} electrodes)")
        
        self.before_electrode_count_label.setText(f"Electrodes: {len(self.before_electrodes)}")
        
        # Redraw
        if self.before_mri_data is not None:
            self._plot_slices(self.before_mri_data, self.before_ct_data, self.before_canvas,
                            "Unaligned CT Overlaid on MRI",
                            self.before_slice_x, self.before_slice_y, self.before_slice_z)
        
        self.log(f"Deleted last electrode from shaft '{shaft_name}'")
    
    def _on_before_shaft_selected(self, *, redraw_2d=True, sync_3d=True):
        """Handle shaft selection in before view - populate electrode list and optionally refresh displays."""
        def _coord_text(value):
            try:
                value = float(value)
            except (TypeError, ValueError):
                return "?"
            if not np.isfinite(value):
                return "?"
            if abs(value - round(value)) < 1e-6:
                return str(int(round(value)))
            return f"{value:.1f}"

        # Populate electrode list for selected shaft
        self.before_electrode_list.setUpdatesEnabled(False)
        blocker = QtCore.QSignalBlocker(self.before_electrode_list)
        self.before_electrode_list.clear()
        selected_items = self.before_shaft_list.selectedItems()
        
        if selected_items:
            # Get shaft name from item data (not text, since we're using custom widgets)
            shaft_name = selected_items[0].data(Qt.UserRole)
            
            if shaft_name in self.before_shafts:
                electrode_indices = self._order_electrode_indices_distal_first(
                    self.before_electrodes,
                    self.before_shafts[shaft_name],
                )
                self.before_shafts[shaft_name] = list(electrode_indices)
                self._refresh_contact_numbers_for_shaft(shaft_name, electrode_indices)
                electrodes_with_pos = []
                for idx in electrode_indices:
                    if idx < len(self.before_electrodes):
                        elec = self.before_electrodes[idx]
                        pos_3d = (elec.x, elec.y, elec.z)
                        electrodes_with_pos.append((idx, pos_3d))
                
                # Add to list with custom widget items
                for elec_num, (idx, pos_3d) in enumerate(electrodes_with_pos, 1):
                    elec = self.before_electrodes[idx]
                    # Use custom name if available, otherwise use shaft + number
                    contact_name = elec.name or f"{shaft_name}{elec_num}"
                    parcellation = self._electrode_parcellation_name(elec)
                    display_name = (
                        f"{contact_name} "
                        f"[{_coord_text(elec.x)}, {_coord_text(elec.y)}, {_coord_text(elec.z)}]"
                    )
                    if parcellation:
                        display_name = f"{display_name} [{parcellation}]"
                    
                    # Create list item
                    item = QtWidgets.QListWidgetItem()
                    item.setFlags(Qt.ItemIsSelectable | Qt.ItemIsEnabled)
                    item.setData(Qt.UserRole, idx)  # Store original electrode index
                    
                    # Create custom widget with text and icons
                    widget = ListItemWidget(display_name)
                    widget.text_label.setToolTip(display_name)
                    widget.edit_btn.setToolTip("Rename/Assign")
                    widget.set_hidden_state(not elec.visible)
                    
                    # Connect widget signals to handlers
                    widget.edit_clicked.connect(lambda w, el=self.before_electrode_list: self._edit_electrode_from_widget(w, el))
                    widget.hide_clicked.connect(lambda w, el=self.before_electrode_list: self._toggle_electrode_visibility_from_widget(w, el))
                    widget.delete_clicked.connect(lambda w, el=self.before_electrode_list: self._delete_electrode_from_widget(w, el))
                    
                    # Add item and set custom widget
                    self.before_electrode_list.addItem(item)
                    self.before_electrode_list.setItemWidget(item, widget)

        del blocker
        self.before_electrode_list.setUpdatesEnabled(True)
        
        # Redraw to highlight
        if redraw_2d and self.before_mri_data is not None:
            if self.before_pg_views:
                self._update_pg_markers()
            else:
                self._plot_slices(self.before_mri_data, self.before_ct_data, self.before_canvas,
                                "Unaligned CT Overlaid on MRI",
                                self.before_slice_x, self.before_slice_y, self.before_slice_z)
        selected_items = self.before_shaft_list.selectedItems()
        selected_shaft = selected_items[0].data(Qt.UserRole) if selected_items else None
        if sync_3d:
            self._sync_selected_shaft_to_controller(selected_shaft)
    
    def _on_before_shaft_pressed(self, item):
        if self._consume_shaft_action_flag(self.before_shaft_list):
            return
        mods = QtWidgets.QApplication.keyboardModifiers()
        if mods & (Qt.ControlModifier | Qt.ShiftModifier):
            self._before_shaft_last_clicked = None

    def _on_before_shaft_clicked(self, item):
        if self._consume_shaft_action_flag(self.before_shaft_list):
            return
        mods = QtWidgets.QApplication.keyboardModifiers()
        if mods & (Qt.ControlModifier | Qt.ShiftModifier):
            self._before_shaft_last_clicked = None
            return
        if item is None:
            self._before_shaft_last_clicked = None
            return
        if self._before_shaft_last_clicked is item and item.isSelected():
            blocker = QtCore.QSignalBlocker(self.before_shaft_list)
            item.setSelected(False)
            del blocker
            self._before_shaft_last_clicked = None
            self._on_before_shaft_selected()
            return
        self._before_shaft_last_clicked = item

    def _edit_before_shaft(self):
        """Open edit dialog for selected shaft in before view."""
        selected_items = self.before_shaft_list.selectedItems()
        if not selected_items:
            self.log("No shaft selected to edit")
            return
        old_name = selected_items[0].data(Qt.UserRole)
        if not old_name:
            shaft_text = selected_items[0].text()
            old_name = shaft_text.split(' (')[0]
        edit = self._prompt_shaft_edit(old_name)
        if edit is not None:
            self._apply_shaft_edit(
                old_name,
                edit["name"],
                edit["trajectory_key"],
                edit["current_trajectory_key"],
                edit.get("name_edited", False),
            )
    
    def _toggle_before_shaft_visibility(self):
        """Toggle visibility of selected shaft in before view."""
        selected_items = self.before_shaft_list.selectedItems()
        if not selected_items:
            self.log("No shaft selected")
            return
        
        shaft_name = selected_items[0].data(Qt.UserRole)
        if not shaft_name:
            shaft_text = selected_items[0].text()
            shaft_name = shaft_text.split(' (')[0]
        
        # Toggle visibility
        is_hidden = self.before_shaft_visibility.get(shaft_name, True)
        self.before_shaft_visibility[shaft_name] = not is_hidden
        
        # Update button appearance
        hide_btn = getattr(self, "before_hide_shaft_btn", None)
        if hide_btn is not None:
            hide_btn.setChecked(not self.before_shaft_visibility[shaft_name])
        
        # Redraw
        if self.before_mri_data is not None:
            self._plot_slices(self.before_mri_data, self.before_ct_data, self.before_canvas,
                            "Unaligned CT Overlaid on MRI",
                            self.before_slice_x, self.before_slice_y, self.before_slice_z)
        self._schedule_detected_sync(self.before_electrodes)
        
        status = "hidden" if not self.before_shaft_visibility.get(shaft_name, True) else "shown"
        self.log(f"Shaft '{shaft_name}' {status}")
    
    def _on_before_shaft_renamed(self, item):
        """Handle shaft renaming via double-click edit in before view."""
        new_text = item.text()
        new_name = new_text.split(' (')[0]  # Remove electrode count
        old_name = item.data(Qt.UserRole)
        
        if old_name and old_name != new_name and old_name in self.before_shafts:
            # Check if new name already exists
            if new_name in self.before_shafts:
                self.log(f"Shaft name '{new_name}' already exists!")
                # Revert the change
                self._update_shaft_list_with_colors(self.before_shaft_list, self.before_shafts)
                return
            
            # Rename the shaft
            self.before_shafts[new_name] = self.before_shafts.pop(old_name)
            
            # Update shaft visibility dict
            if old_name in self.before_shaft_visibility:
                self.before_shaft_visibility[new_name] = self.before_shaft_visibility.pop(old_name)
            
            # Update all electrodes in this shaft
            for idx in self.before_shafts[new_name]:
                if idx < len(self.before_electrodes):
                    self.before_electrodes[idx].shaft = new_name
            
            self.log(f"Renamed shaft from '{old_name}' to '{new_name}'")
        
        # Update the shaft list with proper formatting
        self._update_shaft_list_with_colors(self.before_shaft_list, self.before_shafts)
    
    def _edit_before_electrode(self):
        """Open rename dialog for selected electrode in before view."""
        selected_items = self.before_electrode_list.selectedItems()
        if not selected_items:
            self.log("No electrode selected to rename")
            return
        
        electrode_idx = selected_items[0].data(Qt.UserRole)
        if electrode_idx is None or electrode_idx >= len(self.before_electrodes):
            return
        
        old_name = self.before_electrodes[electrode_idx].name or f"Electrode {electrode_idx}"
        
        new_name, ok = QInputDialog.getText(
            self, "Rename Electrode", f"Enter new name for electrode:", 
            text=old_name
        )
        
        if ok and new_name and new_name != old_name:
            self.before_electrodes[electrode_idx].name = new_name
            self.log(f"Renamed electrode to '{new_name}'")
            self._on_before_shaft_selected()  # Refresh electrode list
            
            # Redraw
            if self.before_mri_data is not None:
                self._plot_slices(self.before_mri_data, self.before_ct_data, self.before_canvas,
                                "Unaligned CT Overlaid on MRI",
                                self.before_slice_x, self.before_slice_y, self.before_slice_z)
    
    def _toggle_before_electrode_visibility(self):
        """Toggle visibility of selected electrode in before view."""
        selected_items = self.before_electrode_list.selectedItems()
        if not selected_items:
            self.log("No electrode selected")
            return
        
        electrode_idx = selected_items[0].data(Qt.UserRole)
        if electrode_idx is None or electrode_idx >= len(self.before_electrodes):
            return
        
        # Toggle visibility in electrode data
        is_visible = self.before_electrodes[electrode_idx].visible
        self.before_electrodes[electrode_idx].visible = not is_visible
        
        # Update button appearance
        self.before_hide_electrode_btn.setChecked(not self.before_electrodes[electrode_idx].visible)
        
        # Redraw
        if self.before_mri_data is not None:
            self._plot_slices(self.before_mri_data, self.before_ct_data, self.before_canvas,
                            "Unaligned CT Overlaid on MRI",
                            self.before_slice_x, self.before_slice_y, self.before_slice_z)
        
        status = "hidden" if not self.before_electrodes[electrode_idx].visible else "shown"
        self.log(f"Electrode {status}")
    
    def _on_before_electrode_selected(self):
        """Handle electrode selection in before view - navigate to electrode position."""
        selected_items = self.before_electrode_list.selectedItems()
        if not selected_items:
            if self.before_selected_parcellation_label is not None:
                self.before_selected_parcellation_label = None
                self.before_selected_parcellation_name = None
                if self.before_mri_data is not None:
                    self._plot_slices(self.before_mri_data, self.before_ct_data, self.before_canvas,
                                    "Unaligned CT Overlaid on MRI",
                                    self.before_slice_x, self.before_slice_y, self.before_slice_z)
            return
        if selected_items:
            item = selected_items[0]
            electrode_idx = item.data(Qt.UserRole)
            
            if electrode_idx is not None and electrode_idx < len(self.before_electrodes):
                elec = self.before_electrodes[electrode_idx]
                pos_3d = (elec.x, elec.y, elec.z)
                elec_name = elec.name or f"Electrode {electrode_idx+1}"
                self._set_selected_parcellation_from_electrode(electrode_idx, elec_name)
                
                if pos_3d is not None:
                    # Update slice positions to center on electrode
                    self.before_slice_x = int(pos_3d[0])
                    self.before_slice_y = int(pos_3d[1])
                    self.before_slice_z = int(pos_3d[2])
                    
                    # Update sliders
                    self.before_sag_slider.setValue(self.before_slice_x)
                    self.before_cor_slider.setValue(self.before_slice_y)
                    self.before_ax_slider.setValue(self.before_slice_z)
                    
                    # Redraw will happen automatically from slider valueChanged signal
                    if self.before_mri_data is not None:
                        self._plot_slices(self.before_mri_data, self.before_ct_data, self.before_canvas,
                                        "Unaligned CT Overlaid on MRI",
                                        self.before_slice_x, self.before_slice_y, self.before_slice_z)
                    self.log(f"Navigated to '{elec_name}' at position {pos_3d}")
    
    def _on_before_electrode_renamed(self, item):
        """Handle electrode renaming via double-click edit in before view."""
        new_name = item.text()
        electrode_idx = item.data(Qt.UserRole)
        
        if electrode_idx is not None and electrode_idx < len(self.before_electrodes):
            self.before_electrodes[electrode_idx].name = new_name
            self.log(f"Renamed electrode to '{new_name}'")
            
            # Redraw to update
            if self.before_mri_data is not None:
                self._plot_slices(self.before_mri_data, self.before_ct_data, self.before_canvas,
                                "Unaligned CT Overlaid on MRI",
                                self.before_slice_x, self.before_slice_y, self.before_slice_z)
    
    def _delete_before_electrode(self):
        """Delete selected electrode from before view."""
        selected_items = self.before_electrode_list.selectedItems()
        if not selected_items:
            self.log("No electrode selected")
            return
        
        electrode_idx = selected_items[0].data(Qt.UserRole)
        if electrode_idx is None or electrode_idx >= len(self.before_electrodes):
            return
        
        # Find which shaft this electrode belongs to
        shaft_name = self.before_electrodes[electrode_idx].shaft
        
        # Remove electrode
        del self.before_electrodes[electrode_idx]
        
        # Update shaft's electrode list (need to rebuild with new indices)
        if shaft_name in self.before_shafts:
            # Rebuild all shafts with updated indices
            for shaft in self.before_shafts:
                self.before_shafts[shaft] = [i for i, e in enumerate(self.before_electrodes) if e.shaft == shaft]
            
            # If shaft is now empty, remove it
            if len(self.before_shafts[shaft_name]) == 0:
                del self.before_shafts[shaft_name]
                if shaft_name in self.before_shaft_visibility:
                    del self.before_shaft_visibility[shaft_name]
        
        # Update UI
        self._update_shaft_list_with_colors(self.before_shaft_list, self.before_shafts)
        self.before_electrode_count_label.setText(f"Electrodes: {len(self.before_electrodes)}")
        self._on_before_shaft_selected()  # Refresh electrode list
        
        # Redraw
        if self.before_mri_data is not None:
            self._plot_slices(self.before_mri_data, self.before_ct_data, self.before_canvas,
                            "Unaligned CT Overlaid on MRI",
                            self.before_slice_x, self.before_slice_y, self.before_slice_z)
        
        self.log(f"Deleted electrode from shaft '{shaft_name}'")
    
    def _toggle_before_mip(self, state):
        """Toggle Maximum Intensity Projection for before view."""
        self.before_show_mip = self.before_mip_checkbox.isChecked()
        
        if self.before_mri_data is not None:
            self._plot_slices(self.before_mri_data, self.before_ct_data, self.before_canvas,
                            "Unaligned CT Overlaid on MRI",
                            self.before_slice_x, self.before_slice_y, self.before_slice_z)
        
        status = "enabled" if self.before_show_mip else "disabled"
        self.log(f"Maximum Intensity Projection {status}")
    
    def log(self, message):
        """Add message to log output."""
        self.log_text.append(message)
        self.log_text.verticalScrollBar().setValue(
            self.log_text.verticalScrollBar().maximum()
        )
    
    def _format_matrix(self, matrix):
        """Format a 4x4 transformation matrix for display."""
        lines = []
        for i in range(4):
            row = "  [" + " ".join([f"{matrix[i, j]:.2f}" for j in range(4)]) + "]"
            lines.append(row)
        return "\n".join(lines)
    
    def _update_transform_display(self, matrix, label="Transform Matrix"):
        """Update the transform matrix display."""
        matrix_str = f"{label}:\n"
        if matrix is None:
            matrix_str += "Matrix unavailable for this registration method."
        else:
            matrix_str += self._format_matrix(matrix)
        self.transform_text.setPlainText(matrix_str)
        
    def load_and_display_before(self):
        """Load and display unaligned CT and MRI."""
        try:
            # Guard against missing paths (dialog may be opened without files)
            if not self.mri_path or not self.ct_path:
                self.log("MRI or CT not selected. Use the Browse buttons to choose files.")
                # Ensure UI shows current (possibly empty) paths
                try:
                    if hasattr(self, 'mri_path_display'):
                        self.mri_path_display.setText(self.mri_path or "")
                    if hasattr(self, 'ct_path_display'):
                        self.ct_path_display.setText(self.ct_path or "")
                except Exception:
                    pass
                # Disable registration until both files are present
                try:
                    if hasattr(self, 'register_btn'):
                        self.register_btn.setEnabled(False)
                except Exception:
                    pass
                return

            self.log("Loading MRI and CT images...")

            # Load images
            mri_img = nib.load(self.mri_path)
            ct_img = nib.load(self.ct_path)
            self.before_mri_img = mri_img
            self.before_ct_img = ct_img
            
            # Get data
            mri_data = mri_img.get_fdata()
            ct_data = ct_img.get_fdata()
            
            self.log(f"MRI shape: {mri_data.shape}")
            self.log(f"CT shape: {ct_data.shape}")
            
            # Resample CT to MRI space for comparison
            from dipy.align import resample
            ct_resampled = resample(
                moving=ct_data,
                static=mri_data,
                moving_affine=ct_img.affine,
                static_affine=mri_img.affine,
            )
            
            # Extract data if resample returns an image object
            if hasattr(ct_resampled, 'get_fdata'):
                ct_resampled = ct_resampled.get_fdata()
            elif not isinstance(ct_resampled, np.ndarray):
                ct_resampled = np.asarray(ct_resampled)
            
            # Store data for slice browsing
            self.before_mri_data = np.asarray(mri_data)
            self.before_ct_data = np.asarray(ct_resampled)
            
            # Initialize slice indices to middle
            self.before_slice_x = self.before_mri_data.shape[0] // 2
            self.before_slice_y = self.before_mri_data.shape[1] // 2
            self.before_slice_z = self.before_mri_data.shape[2] // 2
            
            # Setup sliders
            self.before_sag_slider.setMaximum(self.before_mri_data.shape[0] - 1)
            self.before_sag_slider.setValue(self.before_slice_x)
            self.before_sag_slider.setEnabled(True)
            
            self.before_cor_slider.setMaximum(self.before_mri_data.shape[1] - 1)
            self.before_cor_slider.setValue(self.before_slice_y)
            self.before_cor_slider.setEnabled(True)
            
            self.before_ax_slider.setMaximum(self.before_mri_data.shape[2] - 1)
            self.before_ax_slider.setValue(self.before_slice_z)
            self.before_ax_slider.setEnabled(True)
            
            # Update threshold spinbox ranges based on actual CT data
            ct_min = np.min(self.before_ct_data)
            ct_max = np.max(self.before_ct_data)
            ct_mid = min(max(ct_max * 0.4, ct_min), ct_max)
            self.before_threshold_lower.setRange(ct_min, ct_max)
            self.before_threshold_upper.setRange(ct_min, ct_max)
            self.before_threshold_lower.setValue(ct_mid)
            self.before_threshold_upper.setValue(ct_max)  # Set upper to actual max
            self.log(f"CT intensity range: {ct_min:.0f} to {ct_max:.0f} HU")
            
            # Update labels
            self.before_sag_label.setText(f"Sagittal Slice: {self.before_slice_x}/{self.before_mri_data.shape[0]-1}")
            self.before_cor_label.setText(f"Coronal Slice: {self.before_slice_y}/{self.before_mri_data.shape[1]-1}")
            self.before_ax_label.setText(f"Axial Slice: {self.before_slice_z}/{self.before_mri_data.shape[2]-1}")
            
            # Load parcellation overlay volume if a FreeSurfer subject folder is already selected
            if self.fs_subject_path:
                self._load_before_parcellation_volume()

            # Plot overlay
            self._plot_slices(self.before_mri_data, self.before_ct_data, self.before_canvas,
                            "Unaligned CT Overlaid on MRI",
                            self.before_slice_x, self.before_slice_y, self.before_slice_z)
            
            self.log("✓ Before visualization complete")
            self.log("Click 'Register' to start coregistration")
            
        except Exception as e:
            self.log(f"Error loading images: {e}")
            import traceback
            self.log(traceback.format_exc())
    
    def _plot_slices_pg(self, mri_data, ct_data, title, slice_x, slice_y, slice_z):
        """Plot slices using PyQtGraph."""
        if not self.before_pg_views:
            return
        saved_ranges = self._capture_before_pg_view_ranges()

        mri_data = np.asarray(mri_data)
        ct_data = np.asarray(ct_data)

        show_mri = self.before_show_mri
        show_ct = self.before_show_ct
        ct_alpha = self.before_ct_alpha
        show_mip = self.before_show_mip
        electrodes = self.before_electrodes
        shaft_visibility = self.before_shaft_visibility
        show_shaft_labels = self.before_show_shaft_labels
        marking_active = bool(
            getattr(self, "before_electrode_checkbox", None)
            and self.before_electrode_checkbox.isChecked()
        )

        # Threshold CT only for MIP rendering and detection-style views.
        ct_thresh = ct_data.copy()
        threshold = np.quantile(ct_data, 0.95)
        ct_thresh[ct_thresh < threshold] = np.nan
        ct_slice_overlay = self._prepare_ct_slice_overlay(ct_data)

        if show_mip:
            if getattr(self, "before_mip_threshold_checkbox", None) and self.before_mip_threshold_checkbox.isChecked():
                threshold_min = self.before_threshold_lower.value()
                threshold_max = self.before_threshold_upper.value()
                ct_mip_source = np.where(
                    (ct_data >= threshold_min) & (ct_data <= threshold_max),
                    ct_data,
                    np.nan
                )
            else:
                ct_mip_source = ct_thresh
            ct_mip_axial = self._safe_nanmax(ct_mip_source, axis=2).T
            ct_mip_coronal = self._safe_nanmax(ct_mip_source, axis=1).T
            ct_mip_sagittal = self._safe_nanmax(ct_mip_source, axis=0).T

            slices_data = {
                "axial": (mri_data[:, :, slice_z].T, ct_mip_axial),
                "coronal": (mri_data[:, slice_y, :].T, ct_mip_coronal),
                "sagittal": (mri_data[slice_x, :, :].T, ct_mip_sagittal),
            }
        else:
            slices_data = {
                "axial": (mri_data[:, :, slice_z].T, ct_slice_overlay[:, :, slice_z].T),
                "coronal": (mri_data[:, slice_y, :].T, ct_slice_overlay[:, slice_y, :].T),
                "sagittal": (mri_data[slice_x, :, :].T, ct_slice_overlay[slice_x, :, :].T),
            }

        shaft_label_stats = {}
        visible_label_coords = []
        if show_mip and show_shaft_labels and not marking_active:
            label_positions = {}
            for electrode in electrodes:
                if not electrode.visible:
                    continue
                shaft_name = electrode.shaft
                if not shaft_visibility.get(shaft_name, True):
                    continue
                coord = (electrode.x, electrode.y, electrode.z)
                visible_label_coords.append(coord)
                if shaft_name != "Unassigned":
                    label_positions.setdefault(shaft_name, []).append(coord)
            for name, coords in label_positions.items():
                if coords:
                    arr = np.asarray(coords, dtype=float)
                    shaft_label_stats[name] = {
                        "mean": np.mean(arr, axis=0),
                        "coords": arr
                    }
        shaft_color_map = self._get_shaft_color_map(self.before_shafts) if show_shaft_labels else {}

        view_flags = {
            "axial": self.before_show_axial,
            "coronal": self.before_show_coronal,
            "sagittal": self.before_show_sagittal,
        }

        for view_key, view in self.before_pg_views.items():
            view_box = view["view"]
            show_view = view_flags.get(view_key, True)
            container = self.before_pg_containers.get(view_key)
            widget = self.before_pg_widgets.get(view_key)
            if container is not None:
                container.setVisible(show_view)
            if widget is not None:
                widget.setVisible(show_view)
            view_box.setVisible(show_view)
            if not show_view:
                continue

            mri_slice, ct_slice = slices_data[view_key]
            mri_slice = np.asarray(mri_slice)
            if mri_slice.ndim != 2:
                mri_slice = np.squeeze(mri_slice)
            if mri_slice.ndim != 2:
                mri_slice = np.zeros((1, 1), dtype=np.float32)
            self.before_pg_slices[view_key] = (mri_slice, ct_slice)
            height, width = mri_slice.shape
            self._set_before_pg_view_range(view_key, view_box, width, height, saved_ranges)

            mri_item = view["mri"]
            ct_item = view["ct"]
            mri_item.setVisible(show_mri)
            ct_item.setVisible(show_ct)

            if show_mri:
                mri_slice = np.asarray(mri_slice, dtype=np.float32)
                mri_item.setImage(mri_slice, autoLevels=False)
                mri_item.setLevels(self._safe_mri_levels(mri_slice))

            if show_ct:
                levels = None if show_mip else self._ct_slice_overlay_levels()
                ct_rgba = self._make_ct_rgba(ct_slice, ct_alpha, levels=levels)
                if ct_rgba is not None:
                    ct_item.setImage(ct_rgba, autoLevels=False)
                else:
                    ct_item.clear()

            parcellation_item = view.get("parcellation")
            if parcellation_item is not None:
                if self.before_parcellation_overlay and self.before_parcellation_available:
                    parcellation_slice = self._get_before_parcellation_slice(
                        view_key, slice_x, slice_y, slice_z
                    )
                    parcellation_rgba = self._make_parcellation_rgba(parcellation_slice)
                    if parcellation_rgba is not None:
                        parcellation_item.setImage(parcellation_rgba, autoLevels=False)
                    else:
                        parcellation_item.clear()
                else:
                    parcellation_item.clear()

            colorbar = view.get("colorbar")
            if colorbar is not None:
                if show_ct and show_mip:
                    ct_valid = ct_slice[np.isfinite(ct_slice)]
                    if ct_valid.size:
                        vmin = float(np.nanmin(ct_valid))
                        vmax = float(np.nanmax(ct_valid))
                        colorbar.setLevels((vmin, vmax), update_items=False)
                        colorbar.setVisible(True)
                    else:
                        colorbar.setVisible(False)
                else:
                    colorbar.setVisible(False)

            for item in view.get("shaft_labels", []):
                view_box.removeItem(item)
            view["shaft_labels"] = []

            # Crosshair positions
            if view_key == "axial":
                view["vline"].setPos(slice_x)
                view["hline"].setPos(slice_y)
                left_label = "L"
                right_label = "R"
                top_label = "A"
                bottom_label = "P"
            elif view_key == "coronal":
                view["vline"].setPos(slice_x)
                view["hline"].setPos(slice_z)
                left_label = "L"
                right_label = "R"
                top_label = "S"
                bottom_label = "I"
            else:
                view["vline"].setPos(slice_y)
                view["hline"].setPos(slice_z)
                left_label = "P"
                right_label = "A"
                top_label = "S"
                bottom_label = "I"

            height, width = mri_slice.shape
            labels = view["labels"]
            labels["left"].setText(left_label)
            labels["right"].setText(right_label)
            labels["top"].setText(top_label)
            labels["bottom"].setText(bottom_label)
            labels["left"].setPos(1, height / 2.0)
            labels["right"].setPos(width - 1, height / 2.0)
            y_inverted = bool(view_box.state.get("yInverted", False))
            top_y = 1 if y_inverted else height - 1
            bottom_y = height - 1 if y_inverted else 1
            labels["top"].setPos(width / 2.0, top_y)
            labels["bottom"].setPos(width / 2.0, bottom_y)

            if show_mip and shaft_label_stats and show_shaft_labels and not marking_active:
                label_offset = DEFAULT_SHAFT_LABEL_OFFSET
                view_avoid_points = self._project_shaft_coords_to_view(
                    visible_label_coords,
                    view_key,
                )
                placed_label_boxes = []
                for shaft_name, stats in shaft_label_stats.items():
                    color = shaft_color_map.get(shaft_name, "#333333")
                    coords = stats["coords"]
                    label_text = self._shaft_display_label(
                        shaft_name,
                        (getattr(self, "before_shafts", {}) or {}).get(shaft_name, []),
                    )
                    label_pos = self._compute_shaft_label_position(
                        coords,
                        view_key,
                        mri_data.shape,
                        label_offset,
                        avoid_points=view_avoid_points,
                        placed_labels=placed_label_boxes,
                        label_text=label_text,
                    )
                    if label_pos is None:
                        continue
                    text_color = QtGui.QColor(color)
                    text_color.setAlpha(210)
                    _, anchor_x = self._shaft_label_horizontal_alignment(
                        label_pos,
                        view_key,
                        mri_slice.shape,
                    )
                    label_item = pg.TextItem(label_text, color=text_color, anchor=(anchor_x, 0.5))
                    label_item.setZValue(12)
                    label_item.setPos(label_pos[0], label_pos[1])
                    view_box.addItem(label_item)
                    view["shaft_labels"].append(label_item)
                    half_width, half_height = self._estimate_label_half_size(
                        label_text,
                        mri_data.shape,
                    )
                    placed_label_boxes.append({
                        "pos": (float(label_pos[0]), float(label_pos[1])),
                        "half_width": half_width,
                        "half_height": half_height,
                        "anchor_x": anchor_x,
                    })

        self._update_pg_markers()

    def _plot_slices(self, mri_data, ct_data, canvas, title, slice_x, slice_y, slice_z):
        """Plot MRI with CT overlay at specified slice positions with crosshairs.
        
        Args:
            mri_data: MRI volume data
            ct_data: CT volume data
            canvas: Matplotlib canvas to plot on
            title: Title for the plot
            slice_x: Sagittal slice index
            slice_y: Coronal slice index
            slice_z: Axial slice index
        """
        if canvas == self.before_canvas:
            self._plot_slices_pg(mri_data, ct_data, title, slice_x, slice_y, slice_z)
            return
        # Ensure we have numpy arrays
        mri_data = np.asarray(mri_data)
        ct_data = np.asarray(ct_data)
        
        show_mri = self.before_show_mri
        show_ct = self.before_show_ct
        ct_alpha = self.before_ct_alpha
        electrodes = self.before_electrodes
        show_mip = self.before_show_mip
        show_shaft_labels = self.before_show_shaft_labels
        shaft_visibility = self.before_shaft_visibility
        mip_mask_checkbox = getattr(self, "before_mip_threshold_checkbox", None)
        mip_mask_enabled = bool(mip_mask_checkbox and mip_mask_checkbox.isChecked())
        
        # Threshold CT only for MIP rendering and detection-style views.
        ct_thresh = ct_data.copy()
        threshold = np.quantile(ct_data, 0.95)
        ct_thresh[ct_thresh < threshold] = np.nan
        ct_slice_overlay = self._prepare_ct_slice_overlay(ct_data)
        
        # Clear previous plot
        selectors = getattr(canvas, "_rect_selectors", None)
        if selectors:
            for selector in selectors:
                try:
                    selector.set_active(False)
                    selector.disconnect_events()
                except Exception:
                    pass
        canvas.axes.clear()
        
        # Create 3-panel view (filter visible axes on demand)
        fig = canvas.figure
        fig.clear()
        fig.suptitle(title, fontsize=12, fontweight='bold')
        
        # Get slice data for each view
        # Axial: slice through Z, shows X-Y plane
        # Coronal: slice through Y, shows X-Z plane  
        # Sagittal: slice through X, shows Y-Z plane
        
        # If MIP is enabled, compute maximum intensity projection instead of single slice
        if show_mip:
            if mip_mask_enabled:
                threshold_min = self.before_threshold_lower.value()
                threshold_max = self.before_threshold_upper.value()
                ct_mip_source = np.where(
                    (ct_data >= threshold_min) & (ct_data <= threshold_max),
                    ct_data,
                    np.nan
                )
            else:
                ct_mip_source = ct_thresh
            # Compute MIP across the entire volume for each view
            ct_mip_axial = self._safe_nanmax(ct_mip_source, axis=2).T  # Max along Z axis
            ct_mip_coronal = self._safe_nanmax(ct_mip_source, axis=1).T  # Max along Y axis
            ct_mip_sagittal = self._safe_nanmax(ct_mip_source, axis=0).T  # Max along X axis
            
            slices_data = [
                (mri_data[:, :, slice_z].T, ct_mip_axial, f'Axial MIP'),
                (mri_data[:, slice_y, :].T, ct_mip_coronal, f'Coronal MIP'),
                (mri_data[slice_x, :, :].T, ct_mip_sagittal, f'Sagittal MIP')
            ]
        else:
            slices_data = [
                (mri_data[:, :, slice_z].T, ct_slice_overlay[:, :, slice_z].T, f'Axial (Z={slice_z})'),
                (mri_data[:, slice_y, :].T, ct_slice_overlay[:, slice_y, :].T, f'Coronal (Y={slice_y})'),
                (mri_data[slice_x, :, :].T, ct_slice_overlay[slice_x, :, :].T, f'Sagittal (X={slice_x})')
            ]
        
        shaft_label_stats = {}
        visible_label_coords = []
        if show_mip:
            label_positions = {}
            for electrode in electrodes:
                if not electrode.visible:
                    continue
                shaft_name = electrode.shaft
                if not shaft_visibility.get(shaft_name, True):
                    continue
                coord = (electrode.x, electrode.y, electrode.z)
                visible_label_coords.append(coord)
                if shaft_name != "Unassigned":
                    label_positions.setdefault(shaft_name, []).append(coord)
            for name, coords in label_positions.items():
                if coords:
                    arr = np.asarray(coords, dtype=float)
                    shaft_label_stats[name] = {
                        "mean": np.mean(arr, axis=0),
                        "coords": arr
                    }

        view_flags = {
            "axial": self.before_show_axial,
            "coronal": self.before_show_coronal,
            "sagittal": self.before_show_sagittal,
        }
        view_order = ["axial", "coronal", "sagittal"]
        view_data = []
        for idx, view_key in enumerate(view_order):
            if view_flags.get(view_key, True):
                mri_slice, ct_slice, view_title = slices_data[idx]
                view_data.append((idx, mri_slice, ct_slice, view_title))

        if not view_data:
            view_data = [
                (idx, mri_slice, ct_slice, view_title)
                for idx, (mri_slice, ct_slice, view_title) in enumerate(slices_data)
            ]

        axes = [fig.add_subplot(1, len(view_data), i + 1) for i in range(len(view_data))]
        canvas._view_axes = axes
        selectors = []
        for ax, (view_idx, _, _, _) in zip(axes, view_data):
            selector = RectangleSelector(
                ax,
                lambda eclick, erelease, v=view_idx: self._on_before_rect_select(eclick, erelease, v),
                useblit=True,
                button=[1],  # left-click drag
                minspanx=4,
                minspany=4,
                spancoords='data',
                interactive=False
            )
            selectors.append(selector)
        canvas._rect_selectors = selectors
        fig.subplots_adjust(top=0.88, wspace=0.08)

        # Plot each view
        for ax, (view_idx, mri_slice, ct_slice, view_title) in zip(axes, view_data):
            # Plot MRI if enabled
            if show_mri:
                ax.imshow(mri_slice, cmap='gray', origin='lower')
            
            # Plot CT if enabled
            if show_ct:
                if show_mip:
                    im = ax.imshow(ct_slice, cmap='hot', alpha=ct_alpha, origin='lower')
                else:
                    im = ax.imshow(
                        np.ma.masked_invalid(ct_slice),
                        cmap='hot',
                        alpha=ct_alpha,
                        origin='lower',
                        vmin=-1000,
                        vmax=3000,
                    )

                # Add colorbar for MIP views to show intensity values
                if show_mip:
                    # Get the valid (non-NaN) intensity range for better colorbar scaling
                    valid_data = ct_slice[~np.isnan(ct_slice)]
                    if len(valid_data) > 0:
                        vmin = np.nanmin(ct_slice)
                        vmax = np.nanmax(ct_slice)
                        # Update the image with explicit vmin/vmax to ensure max is at top
                        im.set_clim(vmin, vmax)
                        # Add colorbar to the side of each subplot
                        cbar = fig.colorbar(im, ax=ax, fraction=0.046, pad=0.04)
                        cbar.set_label('Intensity (HU)', rotation=270, labelpad=15, fontsize=8)
                        cbar.ax.tick_params(labelsize=7)
                        # Ensure the colorbar shows the maximum value at the top
                        # Set explicit ticks to show min, middle, and max values
                        tick_values = [vmin, (vmin + vmax) / 2, vmax]
                        cbar.set_ticks(tick_values)
                        cbar.set_ticklabels([f'{vmin:.0f}', f'{(vmin+vmax)/2:.0f}', f'{vmax:.0f}'])

            if self.before_parcellation_overlay and self.before_parcellation_available:
                parcellation_slice = self._get_before_parcellation_slice(
                    ['axial', 'coronal', 'sagittal'][view_idx],
                    slice_x,
                    slice_y,
                    slice_z,
                )
                parcellation_rgba = self._make_parcellation_rgba(parcellation_slice)
                if parcellation_rgba is not None:
                    ax.imshow(parcellation_rgba, origin='lower')

            ax.set_title(view_title, fontsize=12, fontweight='bold')
            ax.axis('off')
            
            # Draw crosshairs to show slice intersections
            # Axial view (Z slice): crosshairs at X and Y positions
            if view_idx == 0:  # Axial
                ax.axvline(slice_x, color='cyan', linestyle='--', linewidth=1, alpha=0.6)
                ax.axhline(slice_y, color='cyan', linestyle='--', linewidth=1, alpha=0.6)
            # Coronal view (Y slice): crosshairs at X and Z positions
            elif view_idx == 1:  # Coronal
                ax.axvline(slice_x, color='cyan', linestyle='--', linewidth=1, alpha=0.6)
                ax.axhline(slice_z, color='cyan', linestyle='--', linewidth=1, alpha=0.6)
            # Sagittal view (X slice): crosshairs at Y and Z positions
            elif view_idx == 2:  # Sagittal
                ax.axvline(slice_y, color='cyan', linestyle='--', linewidth=1, alpha=0.6)
                ax.axhline(slice_z, color='cyan', linestyle='--', linewidth=1, alpha=0.6)

            # Orientation labels: L/R on X-axis views, S/I on Z-axis views.
            label_style = dict(
                color='white',
                fontsize=10,
                fontweight='bold',
                bbox=dict(boxstyle='round,pad=0.2', facecolor='black', alpha=0.5, edgecolor='none')
            )
            if view_idx in (0, 1):  # Axial, Coronal: X axis -> Left/Right
                ax.text(0.02, 0.5, 'L', transform=ax.transAxes,
                             ha='left', va='center', **label_style)
                ax.text(0.98, 0.5, 'R', transform=ax.transAxes,
                             ha='right', va='center', **label_style)
            if view_idx == 2:  # Sagittal: X axis -> Posterior/Anterior
                ax.text(0.02, 0.5, 'P', transform=ax.transAxes,
                             ha='left', va='center', **label_style)
                ax.text(0.98, 0.5, 'A', transform=ax.transAxes,
                             ha='right', va='center', **label_style)
            if view_idx == 0:  # Axial: Y axis -> Anterior/Posterior
                ax.text(0.5, 0.98, 'A', transform=ax.transAxes,
                             ha='center', va='top', **label_style)
                ax.text(0.5, 0.02, 'P', transform=ax.transAxes,
                             ha='center', va='bottom', **label_style)
            if view_idx in (1, 2):  # Coronal, Sagittal: Z axis -> Superior/Inferior
                ax.text(0.5, 0.98, 'S', transform=ax.transAxes,
                             ha='center', va='top', **label_style)
                ax.text(0.5, 0.02, 'I', transform=ax.transAxes,
                             ha='center', va='bottom', **label_style)
            
            # Draw electrode markers for this view with shaft-based coloring
            # Only show electrodes that are on the current displayed slice
            shaft_colors = ['red', 'blue', 'green', 'orange', 'purple', 'cyan', 'magenta', 'yellow']
            shafts = self.before_shafts
            
            # Create shaft name to color mapping
            shaft_names = sorted(shafts.keys())
            shaft_color_map = {}
            if "Unassigned" in shaft_names:
                shaft_color_map["Unassigned"] = "#777777"
            colored_names = [name for name in shaft_names if name != "Unassigned"]
            for idx, name in enumerate(colored_names):
                shaft_color_map[name] = shaft_colors[idx % len(shaft_colors)]
            
            hover_idx = self.before_hover_elec_idx
            for elec_idx, electrode in enumerate(electrodes):
                if not electrode.visible:
                    continue
                # Check if electrode should be displayed on this view's current slice
                show_electrode = False
                x_display, y_display = 0, 0
                
                if hasattr(electrode, "slice_axial"):  # Auto-detected electrode with slice info
                    if view_idx == 0:  # Axial view
                        # Show on MIP or if within 1 slice tolerance of current slice
                        if show_mip or abs(electrode.slice_axial - slice_z) <= 1:
                            show_electrode = True
                            x_display = electrode.x
                            y_display = electrode.y
                    elif view_idx == 1:  # Coronal view
                        if show_mip or abs(electrode.slice_coronal - slice_y) <= 1:
                            show_electrode = True
                            x_display = electrode.x
                            y_display = electrode.z
                    elif view_idx == 2:  # Sagittal view
                        if show_mip or abs(electrode.slice_sagittal - slice_x) <= 1:
                            show_electrode = True
                            x_display = electrode.y
                            y_display = electrode.z
                else:  # Manually marked electrode (old format)
                    if electrode.view == view_idx:
                        show_electrode = True
                        x_display = electrode.x
                        y_display = electrode.y
                
                if show_electrode:
                    shaft_name = electrode.shaft
                    
                    # Check if this shaft is visible (default to visible if not in dict)
                    if not shaft_visibility.get(shaft_name, True):
                        continue  # Skip this electrode if shaft is hidden
                    
                    if shaft_name == "Detected":
                        color = '#00aa00'
                    else:
                        color = shaft_color_map.get(shaft_name, 'red')
                    
                    # Check if this shaft is selected
                    shaft_list = self.before_shaft_list
                    selected_items = shaft_list.selectedItems()
                    is_selected = False
                    if selected_items:
                        selected_item = selected_items[0]
                        selected_shaft_name = selected_item.data(Qt.UserRole)
                        if not selected_shaft_name:
                            selected_shaft_text = selected_item.text()
                            if selected_shaft_text:
                                selected_shaft_name = selected_shaft_text.split(' (')[0]
                        is_selected = (shaft_name == selected_shaft_name)
                    
                    marker_style = 'D' if getattr(electrode, "tissue", "unknown") == "gray" else 'o'
                    # Draw electrode marker - hover/selection highlight
                    if hover_idx is not None and elec_idx == hover_idx:
                        ax.plot(x_display, y_display, marker=marker_style, linestyle='None',
                                   markersize=DEFAULT_MARKER_SIZE + 4,
                                   markerfacecolor='yellow',
                                   markeredgewidth=2.0,
                                   markeredgecolor='black',
                                   alpha=1.0)
                    elif is_selected:
                        # Highlighted marker: larger and more prominent
                        ax.plot(x_display, y_display, marker=marker_style, linestyle='None',
                                   markersize=DEFAULT_MARKER_SIZE + 3,  # Larger marker
                                   markerfacecolor='none',  # Transparent center
                                   markeredgewidth=2.5,  # Thicker edge
                                   markeredgecolor=color,  # Colored edge
                                   alpha=1.0)  # Fully opaque
                    else:
                        # Normal marker: small and unobtrusive
                        ax.plot(x_display, y_display, marker=marker_style, linestyle='None',
                                   markersize=DEFAULT_MARKER_SIZE,  # Small marker
                                   markerfacecolor='none',  # Transparent center
                                   markeredgewidth=1.5,  # Thin edge
                                   markeredgecolor=color,  # Colored edge
                                   alpha=0.8)  # Slightly transparent

            marking_active = bool(
                getattr(self, "before_electrode_checkbox", None)
                and self.before_electrode_checkbox.isChecked()
            )

            if show_mip and shaft_label_stats and show_shaft_labels and not marking_active:
                shaft_label_style = dict(
                    fontsize=8,
                    fontweight='bold',
                    va='center',
                    alpha=0.7,
                    zorder=1,
                    bbox=dict(boxstyle='round,pad=0.2', facecolor='black',
                              alpha=0.18, edgecolor='none')
                )
                label_offset = DEFAULT_SHAFT_LABEL_OFFSET
                view_key = {0: "axial", 1: "coronal", 2: "sagittal"}.get(view_idx, "axial")
                view_avoid_points = self._project_shaft_coords_to_view(
                    visible_label_coords,
                    view_key,
                )
                placed_label_boxes = []
                for shaft_name, stats in shaft_label_stats.items():
                    color = shaft_color_map.get(shaft_name, "#333333")
                    coords = stats["coords"]
                    label_text = self._shaft_display_label(
                        shaft_name,
                        (getattr(self, "before_shafts", {}) or {}).get(shaft_name, []),
                    )
                    label_pos = self._compute_shaft_label_position(
                        coords,
                        view_key,
                        mri_data.shape,
                        label_offset,
                        avoid_points=view_avoid_points,
                        placed_labels=placed_label_boxes,
                        label_text=label_text,
                    )
                    if label_pos is None:
                        continue
                    x_label, y_label = float(label_pos[0]), float(label_pos[1])
                    height, width = mri_slice.shape
                    x_label = max(0, min(width - 1, x_label))
                    y_label = max(0, min(height - 1, y_label))
                    ha, anchor_x = self._shaft_label_horizontal_alignment(
                        (x_label, y_label),
                        view_key,
                        mri_slice.shape,
                    )
                    ax.text(x_label, y_label, label_text, color=color,
                                 ha=ha, **shaft_label_style)
                    half_width, half_height = self._estimate_label_half_size(
                        label_text,
                        mri_data.shape,
                    )
                    placed_label_boxes.append({
                        "pos": (x_label, y_label),
                        "half_width": half_width,
                        "half_height": half_height,
                        "anchor_x": anchor_x,
                    })
        
        fig.tight_layout()
        canvas.draw()
    
    def start_registration(self):
        """Start the registration process in a background thread."""
        self.registration_complete = False
        self._update_accept_button()
        self.register_btn.setEnabled(False)
        self.registration_running = True
        self._update_status_chips()
        self.log("\n" + "="*50)
        self.log("Starting registration...")
        self.log("="*50)
        self.log("NOTE: Registration typically takes 3-7 minutes.")
        self.log("The progress bar will update as each phase completes.")
        self.log("The dialog may appear frozen during computation - this is normal.")
        
        # Create and start worker thread
        method = "mne_rigid"
        if getattr(self, "registration_method_combo", None) is not None:
            method = self.registration_method_combo.currentData() or "mne_rigid"
        self.worker = RegistrationWorker(self.mri_path, self.ct_path, self.output_dir, method=method)
        self.worker.progress.connect(self.update_progress)
        self.worker.finished.connect(self.registration_finished)
        self.worker.start()
        
        # Start a timer to show "still working" messages
        self._progress_timer_count = 0
        self._progress_timer = QtCore.QTimer()
        self._progress_timer.timeout.connect(self._show_still_working)
        self._progress_timer.start(15000)  # Update every 15 seconds
    
    def _show_still_working(self):
        """Periodically show that registration is still working."""
        self._progress_timer_count += 1
        elapsed_mins = self._progress_timer_count * 0.25  # 15 seconds = 0.25 minutes
        self.log(f"⏳ Still working... ({elapsed_mins:.1f} minutes elapsed)")
        QtWidgets.QApplication.processEvents()

    def _update_accept_button(self):
        """Enable accept based on registration state."""
        self._update_action_buttons()

    def _on_register_clicked(self):
        """Route the register button to either register or reset."""
        if not hasattr(self, "register_btn"):
            return
        if getattr(self, "registration_complete", False) or self.trans is not None or self.ct_registered_path:
            self.reset_registration()
            return
        self.start_registration()

    def _apply_results(self):
        """Emit results and close or hide the panel."""
        self.applied.emit(self)
        if self.as_panel:
            self.hide()
            self.panel_hidden.emit()
        else:
            self.accept()

    def _on_cancel_clicked(self):
        """Handle cancel action based on panel mode."""
        if self.as_panel:
            self.hide()
            self.panel_hidden.emit()
        else:
            self.reject()
    
    def _set_progress_message(self, message):
        """Update the progress message in the UI."""
        if hasattr(self, "progress_label") and self.progress_label is not None:
            try:
                self.progress_label.setText(message)
            except Exception:
                pass
        if hasattr(self, "progress_bar") and self.progress_bar is not None:
            try:
                self.progress_bar.setFormat(message)
            except Exception:
                pass

    def _refresh_subject_view_combo(self):
        """Enable/disable subject view selection based on loaded FreeSurfer subject."""
        combo = getattr(self, "subject_view_combo", None)
        if combo is None:
            self._update_mni_button_state()
            return
        has_subject = bool(
            self.controller
            and getattr(self.controller, "custom_subject", None)
            and getattr(self.controller, "custom_subjects_dir", None)
        )
        cohort_active = bool(
            self.controller and getattr(self.controller, "mni_cohort_active", False)
        )
        combo.blockSignals(True)
        combo.setEnabled(has_subject and not cohort_active)
        mode = "fsaverage"
        if has_subject and not getattr(self.controller, "force_fsaverage", False):
            mode = "subject"
        if cohort_active:
            mode = "fsaverage"
            combo.setToolTip("MNI cohort mode uses the fsaverage surface")
        else:
            combo.setToolTip("Switch between subject surfaces and fsaverage")
        idx = combo.findData(mode)
        if idx >= 0:
            combo.setCurrentIndex(idx)
        combo.blockSignals(False)
        self._update_mni_button_state()

    def _on_subject_view_changed(self, _index):
        """Handle switching between subject and fsaverage surfaces."""
        combo = getattr(self, "subject_view_combo", None)
        if combo is None or self.controller is None:
            return
        mode = combo.currentData() or "fsaverage"
        if hasattr(self.controller, "set_subject_view_mode"):
            self.controller.set_subject_view_mode(mode)

    def update_progress(self, percentage, message):
        """Update progress bar and status."""
        self.progress_bar.setValue(percentage)
        self._set_progress_message(message)
        self.log(f"[{percentage}%] {message}")
        QtWidgets.QApplication.processEvents()
    
    def registration_finished(self, success, message, trans, ct_registered_path):
        """Handle registration completion."""
        # Stop the progress timer
        if hasattr(self, '_progress_timer'):
            self._progress_timer.stop()

        self.registration_running = False
        
        self.log("\n" + message)
        
        if success:
            self.trans = trans
            self.ct_registered_path = ct_registered_path
            self.registration_complete = True
            self.electrode_display_space = self._current_electrode_space()
            
            # Update transform matrix display
            matrix = None
            if isinstance(trans, dict) and "trans" in trans:
                matrix = trans.get("trans")
            else:
                try:
                    matrix = np.array(trans)
                except Exception:
                    matrix = None
            if matrix is not None:
                self._update_transform_display(matrix, "Registration Transform Matrix")
            else:
                self._update_transform_display(None, "Registration Transform Matrix (ANTs)")

            # Auto-save transform alongside CT for future reuse.
            self._save_transform_for_ct()
            
            # Update the view with registered CT
            self._update_view_with_registered_ct()
            
            # Change Register button to Reset
            self.register_btn.setText("Reset")
            self.register_btn.setStyleSheet(
                "QPushButton { padding: 8px 12px; font-size: 13px; background-color: #ff9800; "
                "color: white; border-radius: 6px; }"
            )
            self.register_btn.setEnabled(True)
            self.register_btn.update()
            
            # Enable accept button if registration completed or skip is on
            self._update_accept_button()
            self._set_progress_message("Registration complete! CT is now aligned to MRI")
            if self.before_electrodes:
                self._schedule_detected_sync(self.before_electrodes)
            
        else:
            self._set_progress_message("Registration failed. See log for details.")
            self.register_btn.setEnabled(True)
            self.registration_complete = False
            self._update_accept_button()
        self._update_status_chips()
    
    def _update_view_with_registered_ct(self):
        """Update the current view to show registered CT."""
        try:
            self.log("\nUpdating view with registered CT...")
            
            # Load registered CT
            ct_aligned_img = nib.load(self.ct_registered_path)
            ct_aligned_data = ct_aligned_img.get_fdata()
            
            # Update the CT data shown in the view
            self.before_ct_data = np.asarray(ct_aligned_data)
            
            # Update threshold spinbox ranges based on registered CT data
            ct_min = np.min(self.before_ct_data)
            ct_max = np.max(self.before_ct_data)
            ct_mid = min(max(ct_max * 0.4, ct_min), ct_max)
            self.before_threshold_lower.setRange(ct_min, ct_max)
            self.before_threshold_upper.setRange(ct_min, ct_max)
            self.before_threshold_lower.setValue(ct_mid)
            self.before_threshold_upper.setValue(ct_max)
            self.log(f"Registered CT intensity range: {ct_min:.0f} to {ct_max:.0f} HU")
            
            # Redraw with registered CT
            self._plot_slices(self.before_mri_data, self.before_ct_data, self.before_canvas,
                            "Registered CT Overlaid on MRI",
                            self.before_slice_x, self.before_slice_y, self.before_slice_z)
            
            self.log("✓ View updated with registered CT")
            
        except Exception as e:
            self.log(f"Error updating view: {e}")
            import traceback
            self.log(traceback.format_exc())
    
    def reset_registration(self):
        """Reset to original unregistered state."""
        self.log("Reset requested.")
        reply = QtWidgets.QMessageBox.question(
            self,
            "Reset Registration?",
            "This will reload the original unregistered CT.\n"
            "Are you sure you want to reset?",
            QtWidgets.QMessageBox.Yes | QtWidgets.QMessageBox.No,
            QtWidgets.QMessageBox.No
        )
        
        if reply == QtWidgets.QMessageBox.Yes:
            self.log("\n" + "="*50)
            self.log("Resetting to original unregistered CT...")
            self.log("="*50)
            self.registration_running = False

            # Remove saved transform and registered CT files if present.
            try:
                transform_path = self._transform_path(self.ct_path) if self.ct_path else None
                if transform_path and os.path.exists(transform_path):
                    os.remove(transform_path)
                    self.log(f"Deleted transform file: {transform_path}")
            except Exception as e:
                self.log(f"Failed to delete transform file: {e}")
            try:
                reg_path = self.ct_registered_path
                if reg_path:
                    reg_path_str = str(reg_path)
                    if self.ct_path and os.path.abspath(reg_path_str) == os.path.abspath(self.ct_path):
                        reg_path_str = None
                    if reg_path_str and os.path.exists(reg_path_str):
                        os.remove(reg_path_str)
                        self.log(f"Deleted registered CT: {reg_path_str}")
            except Exception as e:
                self.log(f"Failed to delete registered CT: {e}")
            
            # Reset transform to identity
            identity_affine = np.eye(4)
            self.trans = None
            self.ct_registered_path = None
            self.electrode_display_space = self._current_electrode_space()
            self._update_transform_display(identity_affine, "Identity Matrix (No Registration)")
            self.registration_complete = False
            self._update_accept_button()
            
            # Reload original CT using the standard load path.
            try:
                self._suppress_transform_load = True
                if self.mri_path:
                    self._load_mri_and_display(self.mri_path)
                elif self.ct_path:
                    self._load_ct_and_display(self.ct_path)
                else:
                    self.log("Reset skipped: MRI/CT paths missing.")
                
                # Reset button back to Register
                self.register_btn.setText("Register")
                self.register_btn.setStyleSheet(
                    "QPushButton { padding: 8px 12px; font-size: 13px; background-color: #4CAF50; "
                    "color: white; border-radius: 6px; }"
                )
                self.register_btn.update()
                
                # Disable accept button unless skip is on
                self._update_accept_button()
                
                # Reset progress
                self.progress_bar.setValue(0)
                self._set_progress_message("Ready to register")
                self._update_status_chips()
                if self.before_electrodes:
                    self._schedule_detected_sync(self.before_electrodes)
                
                self.log("✓ Reset complete")
                
            except Exception as e:
                self.log(f"Error during reset: {e}")
                import traceback
                self.log(traceback.format_exc())
            finally:
                self._suppress_transform_load = False
    
    def get_results(self):
        """Return registration results including electrode positions."""
        return {
            'trans': self.trans,
            'ct_registered_path': self.ct_registered_path,
            'electrodes': self.before_electrodes,
            'electrode_space': self._current_display_electrode_space(),
            'electrode_native_space': self._current_electrode_space(),
        }
